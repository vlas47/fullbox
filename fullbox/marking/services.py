from __future__ import annotations

from django.db import IntegrityError, transaction
from django.db.models import Count, Q
from django.http import HttpResponseBadRequest, JsonResponse
from django.utils import timezone
from openpyxl import load_workbook

from sku.models import SKUBarcode

from .models import MarkingCode
from .utils import extract_processing_items


def _views():
    from . import views as marking_views

    return marking_views


def processing_marking_summary_response(*, request, order_id: str):
    ok, response = _views()._require_processing_role(request)
    if not ok:
        return response
    latest, _payload, _agency = _views()._get_processing_order(order_id)
    if not latest:
        return HttpResponseBadRequest("Заявка не найдена")
    rows = (
        MarkingCode.objects.filter(order_type="processing", order_id=order_id, used_at__isnull=False)
        .values("sku_code", "size")
        .annotate(count=Count("id"))
    )
    items = [
        {"sku_code": row["sku_code"], "size": row["size"] or "", "count": row["count"]}
        for row in rows
    ]
    total_count = sum(row["count"] for row in rows)
    return JsonResponse({"ok": True, "items": items, "total_count": total_count})


def processing_marking_scan_response(*, request, order_id: str):
    ok, response = _views()._require_processing_role(request)
    if not ok:
        return response
    latest, payload, agency = _views()._get_processing_order(order_id)
    if not latest:
        return HttpResponseBadRequest("Заявка не найдена")
    data = _views()._parse_json_body(request)
    if data is None:
        return HttpResponseBadRequest("Некорректный JSON")
    code = (data.get("code") or "").strip()
    sku_code = (data.get("sku_code") or "").strip()
    size = (data.get("size") or "").strip()
    barcode = (data.get("barcode") or "").strip()
    box_barcode = (data.get("box_barcode") or "").strip()
    if not code:
        return JsonResponse({"ok": False, "error": "Код ЧЗ не указан"}, status=400)
    if not sku_code:
        return JsonResponse({"ok": False, "error": "Артикул не указан"}, status=400)
    items = extract_processing_items(payload)
    allowed_pairs = {(item["sku_code"], item["size"]) for item in items}
    if not size:
        matched = [pair for pair in allowed_pairs if pair[0] == sku_code]
        if len(matched) == 1:
            size = matched[0][1]
    if (sku_code, size) not in allowed_pairs:
        if size and (sku_code, "") in allowed_pairs:
            size = ""
        else:
            return JsonResponse(
                {"ok": False, "error": "Позиция не найдена в заявке."},
                status=400,
            )
    now = timezone.localtime()
    existing = (
        MarkingCode.objects.select_related("agency", "sku")
        .filter(code=code)
        .first()
    )
    if existing:
        if existing.used_at:
            return JsonResponse({"ok": False, "error": "Код уже использован."}, status=409)
        if agency and existing.agency_id and existing.agency_id != agency.id:
            return JsonResponse({"ok": False, "error": "Код принадлежит другому клиенту."}, status=409)
        if existing.order_type and existing.order_type != "processing":
            return JsonResponse({"ok": False, "error": "Код закреплен в другом процессе."}, status=409)
        if existing.order_id and existing.order_id != order_id:
            return JsonResponse({"ok": False, "error": "Код закреплен за другой заявкой."}, status=409)
        if existing.sku_code and existing.sku_code != sku_code:
            return JsonResponse({"ok": False, "error": "Код относится к другому артикулу."}, status=409)
        if existing.size and size and existing.size != size:
            return JsonResponse({"ok": False, "error": "Код относится к другому размеру."}, status=409)
        if existing.box_barcode and box_barcode and existing.box_barcode != box_barcode:
            return JsonResponse({"ok": False, "error": "Код закреплен за другим коробом."}, status=409)
        update_fields = []
        if not existing.order_id:
            existing.order_id = order_id
            update_fields.append("order_id")
        if not existing.order_type:
            existing.order_type = "processing"
            update_fields.append("order_type")
        if not existing.size and size:
            existing.size = size
            update_fields.append("size")
        if not existing.barcode and barcode:
            existing.barcode = barcode
            update_fields.append("barcode")
        if box_barcode and not existing.box_barcode:
            existing.box_barcode = box_barcode
            update_fields.append("box_barcode")
        if not existing.sku:
            sku = _views()._resolve_sku(agency, sku_code)
            existing.sku = sku
            update_fields.append("sku")
        existing.used_at = now
        existing.used_by = request.user if request.user.is_authenticated else None
        update_fields.extend(["used_at", "used_by"])
        existing.save(update_fields=update_fields)
    else:
        sku = _views()._resolve_sku(agency, sku_code)
        try:
            MarkingCode.objects.create(
                order_type="processing",
                order_id=order_id,
                agency=agency,
                sku=sku,
                sku_code=sku_code,
                size=size,
                barcode=barcode,
                box_barcode=box_barcode,
                code=code,
                source="scan",
                created_by=request.user if request.user.is_authenticated else None,
                used_at=now,
                used_by=request.user if request.user.is_authenticated else None,
            )
        except IntegrityError:
            return JsonResponse({"ok": False, "error": "Код уже учтен."}, status=409)
    count = MarkingCode.objects.filter(
        order_type="processing",
        order_id=order_id,
        sku_code=sku_code,
        size=size,
        used_at__isnull=False,
    ).count()
    total_count = MarkingCode.objects.filter(
        order_type="processing",
        order_id=order_id,
        used_at__isnull=False,
    ).count()
    return JsonResponse(
        {
            "ok": True,
            "sku_code": sku_code,
            "size": size,
            "count": count,
            "total_count": total_count,
        }
    )


def receiving_marking_scan_response(*, request, order_id: str):
    ok, response = _views()._require_receiving_role(request)
    if not ok:
        return response
    latest, payload, agency = _views()._get_receiving_order(order_id)
    if not latest:
        return HttpResponseBadRequest("Заявка не найдена")
    data = _views()._parse_json_body(request)
    if data is None:
        return HttpResponseBadRequest("Некорректный JSON")
    code = (data.get("code") or "").strip()
    sku_code = (data.get("sku_code") or "").strip()
    size = (data.get("size") or "").strip()
    barcode = (data.get("barcode") or "").strip()
    box_barcode = (data.get("box_barcode") or "").strip()
    if not code:
        return JsonResponse({"ok": False, "error": "Код ЧЗ не указан"}, status=400)
    if not sku_code:
        return JsonResponse({"ok": False, "error": "Артикул не указан"}, status=400)
    if not box_barcode:
        return JsonResponse({"ok": False, "error": "Откройте короб перед сканированием ЧЗ."}, status=400)

    items = _views()._extract_receiving_items(payload)
    allowed_pairs = {(item["sku_code"], item["size"]) for item in items}
    if not size:
        matched = [pair for pair in allowed_pairs if pair[0] == sku_code]
        if len(matched) == 1:
            size = matched[0][1]
    if (sku_code, size) not in allowed_pairs:
        if size and (sku_code, "") in allowed_pairs:
            size = ""
        else:
            return JsonResponse(
                {"ok": False, "error": "Позиция не найдена в заявке."},
                status=400,
            )

    sku = _views()._resolve_sku(agency, sku_code)
    order_cz_mode = str(payload.get("receiving_mode") or "").strip().lower() == "cz"
    if not order_cz_mode and (not sku or not sku.honest_sign):
        return JsonResponse(
            {"ok": False, "error": "Для этой позиции не включен режим Честного знака."},
            status=400,
        )

    now = timezone.localtime()
    existing = (
        MarkingCode.objects.select_related("agency", "sku")
        .filter(code=code)
        .first()
    )
    if existing:
        if existing.used_at:
            return JsonResponse({"ok": False, "error": "Код уже использован."}, status=409)
        if agency and existing.agency_id and existing.agency_id != agency.id:
            return JsonResponse({"ok": False, "error": "Код принадлежит другому клиенту."}, status=409)
        if existing.order_type and existing.order_type != "receiving":
            return JsonResponse({"ok": False, "error": "Код закреплен в другом процессе."}, status=409)
        if existing.order_id and existing.order_id != order_id:
            return JsonResponse({"ok": False, "error": "Код закреплен за другой заявкой."}, status=409)
        if existing.sku_code and existing.sku_code != sku_code:
            return JsonResponse({"ok": False, "error": "Код относится к другому артикулу."}, status=409)
        if existing.size and size and existing.size != size:
            return JsonResponse({"ok": False, "error": "Код относится к другому размеру."}, status=409)
        update_fields = []
        if not existing.order_id:
            existing.order_id = order_id
            update_fields.append("order_id")
        if not existing.order_type:
            existing.order_type = "receiving"
            update_fields.append("order_type")
        if not existing.size and size:
            existing.size = size
            update_fields.append("size")
        if not existing.barcode and barcode:
            existing.barcode = barcode
            update_fields.append("barcode")
        if existing.box_barcode != box_barcode:
            existing.box_barcode = box_barcode
            update_fields.append("box_barcode")
        if not existing.sku:
            existing.sku = sku
            update_fields.append("sku")
        existing.used_at = now
        existing.used_by = request.user if request.user.is_authenticated else None
        update_fields.extend(["used_at", "used_by"])
        existing.save(update_fields=update_fields)
    else:
        try:
            MarkingCode.objects.create(
                order_type="receiving",
                order_id=order_id,
                agency=agency,
                sku=sku,
                sku_code=sku_code,
                size=size,
                barcode=barcode,
                box_barcode=box_barcode,
                code=code,
                source="scan",
                created_by=request.user if request.user.is_authenticated else None,
                used_at=now,
                used_by=request.user if request.user.is_authenticated else None,
            )
        except IntegrityError:
            return JsonResponse({"ok": False, "error": "Код уже учтен."}, status=409)

    count = MarkingCode.objects.filter(
        order_type="receiving",
        order_id=order_id,
        sku_code=sku_code,
        size=size,
        used_at__isnull=False,
    ).count()
    total_count = MarkingCode.objects.filter(
        order_type="receiving",
        order_id=order_id,
        used_at__isnull=False,
    ).count()
    return JsonResponse(
        {
            "ok": True,
            "sku_code": sku_code,
            "size": size,
            "count": count,
            "total_count": total_count,
        }
    )


def processing_marking_print_response(*, request, order_id: str):
    ok, response = _views()._require_processing_role(request)
    if not ok:
        return response
    latest, _payload, agency = _views()._get_processing_order(order_id)
    if not latest:
        return HttpResponseBadRequest("Заявка не найдена")
    data = _views()._parse_json_body(request)
    if data is None:
        return HttpResponseBadRequest("Некорректный JSON")
    qty = data.get("qty")
    try:
        qty_value = int(qty)
    except (TypeError, ValueError):
        qty_value = 0
    if qty_value <= 0:
        return JsonResponse({"ok": False, "error": "Количество для печати не указано."}, status=400)
    barcode = (data.get("barcode") or "").strip()
    size = (data.get("size") or "").strip()
    if not barcode:
        return JsonResponse({"ok": False, "error": "ШК не указан."}, status=400)
    now = timezone.now()
    with transaction.atomic():
        base_qs = MarkingCode.objects.select_for_update().filter(
            order_type="processing",
            used_at__isnull=True,
            printed_at__isnull=True,
            barcode=barcode,
        )
        if agency:
            base_qs = base_qs.filter(agency=agency)
        if size:
            base_qs = base_qs.filter(size=size)
        reserved = list(
            base_qs.filter(order_id=order_id)
            .order_by("created_at")
            .values_list("id", "code")[:qty_value]
        )
        remaining = qty_value - len(reserved)
        extra = []
        if remaining > 0:
            extra = list(
                base_qs.filter(Q(order_id__isnull=True) | Q(order_id=""))
                .order_by("created_at")
                .values_list("id", "code")[:remaining]
            )
        codes = reserved + extra
        if len(codes) < qty_value:
            return JsonResponse(
                {"ok": False, "error": "Недостаточно кодов ЧЗ.", "available": len(codes)},
                status=409,
            )
        ids = [item[0] for item in codes]
        MarkingCode.objects.filter(id__in=ids).update(
            printed_at=now,
            printed_by=request.user,
            order_id=order_id,
        )
    return JsonResponse({"ok": True, "codes": [item[1] for item in codes], "count": len(codes)})


def processing_marking_reset_printed_response(*, request, order_id: str):
    ok, response = _views()._require_processing_role(request)
    if not ok:
        return response
    latest, _payload, agency = _views()._get_processing_order(order_id)
    if not latest:
        return HttpResponseBadRequest("Заявка не найдена")
    data = _views()._parse_json_body(request)
    if data is None:
        return HttpResponseBadRequest("Некорректный JSON")
    barcode = (data.get("barcode") or "").strip()
    size = (data.get("size") or "").strip()
    qty = data.get("qty")
    try:
        qty_value = int(qty)
    except (TypeError, ValueError):
        qty_value = 0
    if not barcode:
        return JsonResponse({"ok": False, "error": "ШК не указан."}, status=400)
    with transaction.atomic():
        qs = MarkingCode.objects.select_for_update().filter(
            order_type="processing",
            order_id=order_id,
            used_at__isnull=True,
            printed_at__isnull=False,
            barcode=barcode,
        )
        if agency:
            qs = qs.filter(agency=agency)
        if size:
            qs = qs.filter(size=size)
        if qty_value > 0:
            ids = list(
                qs.order_by("-printed_at", "-created_at")
                .values_list("id", flat=True)[:qty_value]
            )
        else:
            ids = list(qs.values_list("id", flat=True))
        if ids:
            MarkingCode.objects.filter(id__in=ids).update(
                printed_at=None,
                printed_by=None,
            )
    return JsonResponse({"ok": True, "count": len(ids)})


def processing_marking_import_response(*, request, order_id: str):
    ok, response = _views()._require_processing_role(request)
    if not ok:
        return response
    latest, payload, agency = _views()._get_processing_order(order_id)
    if not latest:
        return HttpResponseBadRequest("Заявка не найдена")
    file = request.FILES.get("file")
    if not file:
        return JsonResponse({"ok": False, "error": "Файл не выбран."}, status=400)
    try:
        workbook = load_workbook(file, read_only=True, data_only=True)
    except Exception:
        return JsonResponse({"ok": False, "error": "Не удалось прочитать .xlsx файл."}, status=400)
    sheet = workbook.active

    rows = []
    barcodes = set()
    invalid_rows = 0
    for idx, row in enumerate(sheet.iter_rows(values_only=True), start=1):
        barcode = _views()._normalize_cell(row[0]) if row and len(row) > 0 else ""
        code = _views()._normalize_cell(row[1]) if row and len(row) > 1 else ""
        if idx == 1:
            header = barcode.lower()
            if "штрих" in header or "barcode" in header:
                continue
        if not barcode or not code:
            if barcode or code:
                invalid_rows += 1
            continue
        rows.append((barcode, code))
        barcodes.add(barcode)

    if not rows:
        return JsonResponse(
            {"ok": False, "error": "В файле нет данных для импорта."},
            status=400,
        )

    barcode_qs = SKUBarcode.objects.select_related("sku").filter(value__in=barcodes)
    barcode_map = {item.value: item for item in barcode_qs}
    existing_codes = set(
        MarkingCode.objects.filter(code__in=[code for _, code in rows]).values_list("code", flat=True)
    )

    items = extract_processing_items(payload)
    allowed_pairs = {(item["sku_code"], item["size"]) for item in items}
    added = 0
    duplicates = 0
    unknown_barcodes = 0
    mismatched_barcodes = 0
    seen_codes = set()
    to_create = []

    for barcode, code in rows:
        if code in seen_codes:
            duplicates += 1
            continue
        seen_codes.add(code)
        if code in existing_codes:
            duplicates += 1
            continue
        barcode_obj = barcode_map.get(barcode)
        if not barcode_obj or not barcode_obj.sku:
            unknown_barcodes += 1
            continue
        sku_obj = barcode_obj.sku
        if agency and sku_obj.agency and sku_obj.agency_id != agency.id:
            mismatched_barcodes += 1
            continue
        sku_code = sku_obj.sku_code
        size = (barcode_obj.size or sku_obj.size or "").strip()
        if (sku_code, size) not in allowed_pairs:
            if (sku_code, "") in allowed_pairs:
                size = ""
            else:
                matched = [pair for pair in allowed_pairs if pair[0] == sku_code]
                if len(matched) == 1:
                    size = matched[0][1]
                else:
                    unknown_barcodes += 1
                    continue
        to_create.append(
            MarkingCode(
                order_type="processing",
                order_id=order_id,
                agency=agency,
                sku=sku_obj,
                sku_code=sku_code,
                size=size,
                barcode=barcode,
                code=code,
                source="import",
                created_by=request.user if request.user.is_authenticated else None,
            )
        )

    if to_create:
        with transaction.atomic():
            MarkingCode.objects.bulk_create(to_create, batch_size=500)
        added = len(to_create)

    return JsonResponse(
        {
            "ok": True,
            "added": added,
            "duplicates": duplicates,
            "unknown_barcodes": unknown_barcodes,
            "mismatched_barcodes": mismatched_barcodes,
            "invalid_rows": invalid_rows,
        }
    )
