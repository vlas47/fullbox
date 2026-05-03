from __future__ import annotations

from io import BytesIO
from pathlib import Path

from docx import Document
from docx.shared import Pt
from django.conf import settings
from django.utils import timezone

from head_manager.models import OwnCompany

from .models import ShippingOrder


RETURN_ACT_TEMPLATE_DOCX = Path(settings.BASE_DIR) / "static" / "shipping" / "return_act_template.docx"
FULLBOX_NAME = "FullBox"
FULLBOX_WAREHOUSE_ADDRESS = "142450, МО, г. Старая Купавна, ул. Магистральная, д. 59"
FULLBOX_PHONE = "+7 499 450-35-55"
UNIT_NAME = "шт"
UNIT_OKEI_CODE = "796"
ITEM_ROWS_PRIMARY_START = 25
ITEM_ROWS_PRIMARY_END = 58
ITEM_ROWS_SECONDARY_START = 64
ITEM_ROWS_SECONDARY_END = 70
SERVICE_ROWS_START = 79
SERVICE_ROWS_END = 84


def _text(value) -> str:
    return str(value or "").strip()


def _first_nonempty(*values) -> str:
    for value in values:
        prepared = _text(value)
        if prepared:
            return prepared
    return ""


def _join_nonempty(*values, sep: str = ", ") -> str:
    prepared = [_text(value) for value in values if _text(value)]
    return sep.join(prepared)


def _format_date(value) -> str:
    if not value:
        return ""
    if hasattr(value, "strftime"):
        return value.strftime("%d.%m.%Y")
    return _text(value)


def _default_own_company() -> OwnCompany | None:
    return OwnCompany.objects.filter(is_active=True).order_by("-is_default", "name").first()


def _customer_name(order: ShippingOrder) -> str:
    return (
        _text(getattr(order.agency, "short_name", ""))
        or _text(getattr(order.agency, "agn_name", ""))
        or _text(getattr(order.agency, "fio_agn", ""))
    )


def _holder_name(company: OwnCompany | None) -> str:
    return _text(getattr(company, "short_name", "")) or _text(getattr(company, "name", "")) or FULLBOX_NAME


def _holder_address(company: OwnCompany | None) -> str:
    return (
        _text(getattr(company, "postal_address", ""))
        or _text(getattr(company, "address", ""))
        or FULLBOX_WAREHOUSE_ADDRESS
    )


def _holder_phone(company: OwnCompany | None) -> str:
    return _text(getattr(company, "phone", "")) or FULLBOX_PHONE


def _holder_identity(company: OwnCompany | None) -> str:
    return _join_nonempty(_holder_name(company), _holder_address(company), _holder_phone(company))


def _depositor_identity(order: ShippingOrder) -> str:
    return _join_nonempty(
        _customer_name(order),
        _first_nonempty(getattr(order.agency, "adres", ""), getattr(order.agency, "fakt_adres", "")),
        _text(getattr(order.agency, "phone", "")),
    )


def _document_number(order: ShippingOrder) -> str:
    return _text(order.number) or str(order.pk)


def _document_date(order: ShippingOrder):
    return timezone.localdate(order.shipped_at) if order.shipped_at else timezone.localdate()


def _item_characteristic(item) -> str:
    return _join_nonempty(
        f"размер {item.size}" if _text(item.size) else "",
        f"ШК {item.barcode}" if _text(item.barcode) else "",
        sep=", ",
    )


def _item_name(item) -> str:
    return _text(item.name) or _text(item.sku_code) or "Товар"


def _item_code(item) -> str:
    return _text(item.sku_code)


def _item_qty(item) -> str:
    qty = int(item.qty_shipped or 0) or int(item.qty_requested or 0)
    return str(qty) if qty > 0 else ""


def _item_rows(order: ShippingOrder) -> list[dict]:
    rows: list[dict] = []
    for index, item in enumerate(order.items.order_by("id"), start=1):
        rows.append(
            {
                "index": str(index),
                "name": _item_name(item),
                "code": _item_code(item),
                "characteristic": _item_characteristic(item),
                "unit_name": UNIT_NAME,
                "unit_code": UNIT_OKEI_CODE,
                "qty": _item_qty(item),
                "comment": _text(item.comment),
            }
        )
    return rows


def _set_cell_text(cell, text: str) -> None:
    prepared = _text(text)
    if not cell.paragraphs:
        cell.text = prepared
        return
    paragraph = cell.paragraphs[0]
    if paragraph.runs:
        primary_run = paragraph.runs[0]
        primary_run.text = prepared
        primary_run.font.size = Pt(10)
        for run in paragraph.runs[1:]:
            run.text = ""
            run.font.size = Pt(10)
    else:
        run = paragraph.add_run(prepared)
        run.font.size = Pt(10)
    for extra_paragraph in cell.paragraphs[1:]:
        extra_paragraph._element.getparent().remove(extra_paragraph._element)


def _table_cells(table) -> list:
    return [row.cells for row in table.rows]


def _set_table_text(cells, row: int, col: int, text: str) -> None:
    _set_cell_text(cells[row][col], text)


def _fill_item_range(cells, start_row: int, end_row: int, rows: list[dict]) -> None:
    for current_row in range(start_row, end_row + 1):
        source_index = current_row - start_row
        row_data = rows[source_index] if source_index < len(rows) else {}
        _set_table_text(cells, current_row, 0, row_data.get("index", ""))
        _set_table_text(cells, current_row, 1, row_data.get("name", ""))
        _set_table_text(cells, current_row, 4, row_data.get("code", ""))
        _set_table_text(cells, current_row, 7, row_data.get("characteristic", ""))
        _set_table_text(cells, current_row, 11, row_data.get("unit_name", ""))
        _set_table_text(cells, current_row, 13, row_data.get("unit_code", ""))
        _set_table_text(cells, current_row, 19, row_data.get("qty", ""))
        _set_table_text(cells, current_row, 29, row_data.get("comment", ""))


def render_return_act_doc(order: ShippingOrder) -> bytes:
    company = _default_own_company()
    document = Document(str(RETURN_ACT_TEMPLATE_DOCX))
    cells = _table_cells(document.tables[0])
    document_number = _document_number(order)
    document_date = _format_date(_document_date(order))
    contract_number = _text(getattr(order.agency, "contract_numb", ""))
    item_rows = _item_rows(order)
    total_qty = sum(int(row.get("qty") or 0) for row in item_rows)
    special_marks = _text(order.comment)
    holder_signatory = _text(getattr(company, "director_name", ""))
    depositor_signatory = _text(getattr(order.agency, "fio_agn", "")) or _customer_name(order)

    _set_table_text(cells, 5, 0, _holder_identity(company))
    _set_table_text(cells, 10, 0, _depositor_identity(order))
    _set_table_text(cells, 14, 13, contract_number)
    _set_table_text(cells, 14, 28, document_date if contract_number else "")
    _set_table_text(cells, 17, 13, document_number)
    _set_table_text(cells, 17, 21, document_date)
    _fill_item_range(cells, ITEM_ROWS_PRIMARY_START, ITEM_ROWS_PRIMARY_END, item_rows)
    _set_table_text(cells, 59, 19, str(total_qty) if total_qty > 0 else "")
    _fill_item_range(cells, ITEM_ROWS_SECONDARY_START, ITEM_ROWS_SECONDARY_END, item_rows)
    _set_table_text(cells, 71, 19, str(total_qty) if total_qty > 0 else "")
    _set_table_text(cells, 72, 19, str(total_qty) if total_qty > 0 else "")
    for current_row in range(SERVICE_ROWS_START, SERVICE_ROWS_END + 1):
        _set_table_text(cells, current_row, 0, "")
        _set_table_text(cells, current_row, 3, "")
        _set_table_text(cells, current_row, 5, "")
        _set_table_text(cells, current_row, 9, "")
        _set_table_text(cells, current_row, 12, "")
        _set_table_text(cells, current_row, 18, "")
        _set_table_text(cells, current_row, 22, "")
        _set_table_text(cells, current_row, 26, "")
        _set_table_text(cells, current_row, 29, "")
    _set_table_text(cells, 92, 4, "Генеральный директор")
    _set_table_text(cells, 92, 24, holder_signatory)
    _set_table_text(cells, 93, 4, "Представитель поклажедателя")
    _set_table_text(cells, 93, 24, depositor_signatory)
    _set_table_text(cells, 94, 0, f"Особые отметки {special_marks}" if special_marks else "Особые отметки")
    _set_table_text(cells, 101, 17, depositor_signatory)
    _set_table_text(cells, 103, 6, "Представитель хранителя")
    _set_table_text(cells, 103, 25, holder_signatory)

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def return_act_doc_filename(order: ShippingOrder) -> str:
    number = _document_number(order)
    safe_number = "".join(ch if ch.isalnum() or ch in ("-", "_") else "_" for ch in number)
    return f"return-act-{safe_number}.docx"
