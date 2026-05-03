from __future__ import annotations

from io import BytesIO
from pathlib import Path
from decimal import Decimal
from textwrap import wrap
import re

import fitz
from docx import Document
from docx.shared import Pt
from django.conf import settings
from django.utils import timezone

from audit.models import OrderAuditEntry
from head_manager.models import Carrier, OwnCompany
from logistics.models import LogisticsTrip

from .dispatch import shipping_dispatch_trip_link
from .marketplace_warehouses import load_marketplace_warehouse_catalog, marketplace_key_for_market
from .models import ShippingOrder, ShippingTransportNote
from .packing import _shipping_packing_summary, shipping_packing_slips_data


TRANSPORT_NOTE_ROLES = {
    "storekeeper",
    "head_manager",
    "director",
    "admin",
    "developer",
}

FULLBOX_NAME = "FullBox"
FULLBOX_WAREHOUSE_ADDRESS = "142450, МО, г. Старая Купавна, ул. Магистральная, д. 59"
FULLBOX_PHONE = "+7 499 450-35-55"
TRANSPORT_NOTE_TEMPLATE_PDF = Path(settings.BASE_DIR) / "static" / "shipping" / "transport_note_blank.pdf"
TRANSPORT_NOTE_TEMPLATE_DOCX = Path(settings.BASE_DIR) / "static" / "shipping" / "transport_note_template.docx"
_TRIP_LOADING_AUDIT_TYPE = "logistics_trip"
_TRIP_LOADING_AUDIT_ACT = "trip_loading_progress"


def can_manage_transport_note(scope: str | None, role: str | None) -> bool:
    return scope == "staff" and role in TRANSPORT_NOTE_ROLES


def _normalize_loading_scan(value: str | None) -> str:
    return re.sub(r"\s+", "", str(value or "").strip()).upper()


def _transport_note_trip_loaded(order: ShippingOrder, trip: LogisticsTrip) -> bool:
    if trip.status in {LogisticsTrip.STATUS_DEPARTED, LogisticsTrip.STATUS_COMPLETED}:
        return True
    if trip.status != LogisticsTrip.STATUS_LOADING:
        return False

    packing_summary = _shipping_packing_summary(order) or {}
    slips = shipping_packing_slips_data(order, packing_summary)
    load_keys = {
        _normalize_loading_scan(
            str(slip.get("pallet_code") or slip.get("qr_value") or slip.get("slip_key") or "").strip()
        )
        for slip in slips
    }
    load_keys.discard("")
    load_keys.discard("-")
    if not load_keys:
        return False

    entry = (
        OrderAuditEntry.objects.filter(
            order_id=str(trip.pk),
            order_type=_TRIP_LOADING_AUDIT_TYPE,
            payload__act=_TRIP_LOADING_AUDIT_ACT,
        )
        .order_by("-created_at")
        .first()
    )
    payload = entry.payload if entry and isinstance(entry.payload, dict) else {}
    raw_loaded_keys = payload.get("loaded_pallet_keys") if isinstance(payload, dict) else []
    if not isinstance(raw_loaded_keys, list):
        raw_loaded_keys = []
    loaded_keys = {
        _normalize_loading_scan(value)
        for value in raw_loaded_keys
        if _normalize_loading_scan(value)
    }
    return bool(loaded_keys) and load_keys.issubset(loaded_keys)


def can_access_transport_note(order: ShippingOrder, scope: str | None, role: str | None) -> bool:
    if not can_manage_transport_note(scope, role):
        return False
    if order.status in {ShippingOrder.STATUS_SHIPPED, ShippingOrder.STATUS_PARTIAL}:
        return True
    if order.status != ShippingOrder.STATUS_PACKED:
        return False

    trip_link = shipping_dispatch_trip_link(order)
    trip = trip_link.trip if trip_link else None
    vehicle_type = str(order.vehicle_type or getattr(trip, "vehicle_type", "") or "").strip()
    if vehicle_type == ShippingOrder.VEHICLE_FULFILLMENT:
        return bool(trip) and _transport_note_trip_loaded(order, trip)
    return True


def _text(value) -> str:
    return str(value or "").strip()


def _first_nonempty(*values) -> str:
    for value in values:
        prepared = _text(value)
        if prepared:
            return prepared
    return ""


def _join_nonempty(*values, sep: str = ", ") -> str:
    prepared = [_text(value) for value in values if _text(value)]
    return sep.join(prepared)


def _format_party_identity(*, name: str, inn: str = "", address: str = "", phone: str = "") -> str:
    parts: list[str] = []
    prepared_name = _text(name)
    if prepared_name:
        parts.append(prepared_name)
    prepared_inn = _text(inn)
    if prepared_inn:
        parts.append(f"ИНН {prepared_inn}")
    prepared_address = _text(address)
    if prepared_address:
        parts.append(prepared_address)
    prepared_phone = _text(phone)
    if prepared_phone:
        parts.append(f"тел. {prepared_phone}")
    return ", ".join(parts)


def _format_party_identity_docx(*, name: str, inn: str = "", address: str = "", phone: str = "") -> str:
    return _join_nonempty(
        _text(name),
        f"ИНН {inn}" if _text(inn) else "",
        _short_address(address, limit=180),
        f"тел. {phone}" if _text(phone) else "",
        sep=", ",
    )


def _short_address(value: str, limit: int = 96) -> str:
    prepared = _text(value).replace("\n", ", ")
    if len(prepared) <= limit:
        return prepared
    return prepared[: limit - 1].rstrip(", ") + "…"


def _format_party_identity_pdf(*, name: str, inn: str = "", address: str = "", phone: str = "") -> str:
    lines: list[str] = []
    prepared_name = _text(name)
    if prepared_name:
        lines.append(prepared_name)
    second_line = _join_nonempty(
        f"ИНН {inn}" if _text(inn) else "",
        f"тел. {phone}" if _text(phone) else "",
        sep=", ",
    )
    if second_line:
        lines.append(second_line)
    prepared_address = _short_address(address)
    if prepared_address:
        lines.append(prepared_address)
    return "\n".join(lines)


def _format_currency(value) -> str:
    if value in (None, ""):
        return ""
    try:
        number = Decimal(str(value))
    except Exception:
        return _text(value)
    return f"{number.quantize(Decimal('0.01'))} руб."


def _format_weight(value) -> str:
    if value in (None, ""):
        return ""
    try:
        number = Decimal(str(value))
    except Exception:
        return _text(value)
    return f"{number.quantize(Decimal('0.001'))} кг."


def _format_date(value) -> str:
    if not value:
        return ""
    if hasattr(value, "strftime"):
        return value.strftime("%d.%m.%Y")
    return _text(value)


def _default_own_company() -> OwnCompany | None:
    return (
        OwnCompany.objects.filter(is_active=True)
        .order_by("-is_default", "name")
        .first()
    )


def _default_carrier() -> Carrier | None:
    return Carrier.objects.filter(is_active=True).order_by("name").first()


def _own_company_name(company: OwnCompany | None) -> str:
    return _text(getattr(company, "short_name", "")) or _text(getattr(company, "name", "")) or FULLBOX_NAME


def _own_company_address(company: OwnCompany | None) -> str:
    return (
        _text(getattr(company, "postal_address", ""))
        or _text(getattr(company, "address", ""))
        or FULLBOX_WAREHOUSE_ADDRESS
    )


def _own_company_phone(company: OwnCompany | None) -> str:
    return _text(getattr(company, "phone", "")) or FULLBOX_PHONE


def _own_company_inn(company: OwnCompany | None) -> str:
    return _text(getattr(company, "inn", ""))


def _carrier_name_from_directory(carrier: Carrier | None) -> str:
    return _text(getattr(carrier, "short_name", "")) or _text(getattr(carrier, "name", ""))


def _carrier_address_from_directory(carrier: Carrier | None) -> str:
    return _text(getattr(carrier, "postal_address", "")) or _text(getattr(carrier, "address", ""))


def _customer_name(order: ShippingOrder) -> str:
    return (
        _text(getattr(order.agency, "short_name", ""))
        or _text(order.agency.agn_name)
        or _text(order.agency.fio_agn)
    )


def _marketplace_name(order: ShippingOrder) -> str:
    return _text(getattr(order.marketplace, "name", ""))


def _prefixed_marketplace_name(order: ShippingOrder, name: str) -> str:
    marketplace = _marketplace_name(order)
    prepared_name = _text(name)
    if not marketplace:
        return prepared_name
    if not prepared_name:
        return marketplace
    lowered_marketplace = marketplace.lower()
    lowered_name = prepared_name.lower()
    if lowered_name.startswith(lowered_marketplace):
        return prepared_name
    return f"{marketplace} / {prepared_name}"


def _normalize_lookup_token(value: str) -> str:
    return "".join(ch.lower() for ch in _text(value) if ch.isalnum())


def _marketplace_warehouse_match(order: ShippingOrder) -> dict | None:
    market_key = marketplace_key_for_market(getattr(order, "marketplace", None))
    if not market_key:
        return None
    destination = _text(order.destination_warehouse)
    destination_address = _text(order.destination_address)
    transit_address = _text(order.transit_address)
    order_tokens = {
        _normalize_lookup_token(destination),
        _normalize_lookup_token(destination_address),
        _normalize_lookup_token(transit_address),
    }
    order_tokens.discard("")
    for row in load_marketplace_warehouse_catalog().get(market_key, []):
        row_name = _text(row.get("name"))
        row_address = _text(row.get("address"))
        row_type = _text(row.get("type"))
        row_display = " · ".join(part for part in (row_type, row_name) if part)
        row_tokens = {
            _normalize_lookup_token(row_name),
            _normalize_lookup_token(row_address),
            _normalize_lookup_token(row_display),
            _normalize_lookup_token(f"{row_display} {row_address}"),
        }
        row_tokens.discard("")
        if order_tokens & row_tokens:
            return row
        if row_name and row_name.lower() in destination.lower():
            return row
    return None


def _format_eta_window(order: ShippingOrder) -> str:
    if not order.eta_at:
        return ""
    local_eta = timezone.localtime(order.eta_at)
    return f"{local_eta.strftime('%d.%m.%Y')} г. с {local_eta.strftime('%H:%M')}"


def _format_slot_window(order: ShippingOrder) -> str:
    if not order.slot_date and not order.slot_time:
        return ""
    parts = []
    if order.slot_date:
        parts.append(_format_date(order.slot_date))
    if order.slot_time:
        parts.append(f"с {order.slot_time.strftime('%H:%M')}")
    return " ".join(parts).strip()


def _agency_customer_identity(order: ShippingOrder) -> str:
    return _format_party_identity(
        name=_customer_name(order),
        inn=_text(order.agency.inn),
        address=_first_nonempty(order.agency.adres, order.agency.fakt_adres),
        phone=_text(order.agency.phone),
    )


def _slot(
    left_mm: float,
    top_mm: float,
    width_mm: float,
    text: str,
    *,
    css_class: str = "",
    height_mm: float = 6.0,
) -> dict:
    return {
        "left_mm": left_mm,
        "top_mm": top_mm,
        "width_mm": width_mm,
        "height_mm": height_mm,
        "text": _text(text),
        "css_class": css_class,
    }


def _mm_to_pt(value_mm: float) -> float:
    return float(value_mm) * 72 / 25.4


def _slot_font_size(slot: dict) -> float:
    css_class = slot.get("css_class", "")
    if "slot-xxs" in css_class:
        return 4.2
    if "slot-md" in css_class:
        return 7.2
    if "slot-sm" in css_class:
        return 6.4
    return 5.2


def _slot_align(slot: dict) -> int:
    css_class = slot.get("css_class", "")
    if "slot-center" in css_class:
        return 1
    if "slot-right" in css_class:
        return 2
    return 0


def _transport_note_font_path() -> str | None:
    candidates = [
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
        Path("/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf"),
        Path(r"C:\Windows\Fonts\arial.ttf"),
        Path(r"C:\Windows\Fonts\calibri.ttf"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return str(candidate)
    return None


def _draw_pdf_slots(page: fitz.Page, slots: list[dict]) -> None:
    font_path = _transport_note_font_path()
    font_name = "tnfont"
    if font_path:
        page.insert_font(fontname=font_name, fontfile=font_path)
    for slot in slots:
        text = _text(slot.get("text"))
        if not text:
            continue
        rect = fitz.Rect(
            _mm_to_pt(slot["left_mm"]),
            _mm_to_pt(slot["top_mm"]),
            _mm_to_pt(slot["left_mm"] + slot["width_mm"]),
            _mm_to_pt(slot["top_mm"] + slot["height_mm"]),
        )
        page.insert_textbox(
            rect,
            text,
            fontname=font_name if font_path else "helv",
            fontsize=_slot_font_size(slot),
            align=_slot_align(slot),
            lineheight=1.05,
            color=(0, 0, 0),
        )


def _set_cell_text(cell, text: str) -> None:
    text = _text(text)
    if not cell.paragraphs:
        cell.text = text
    first = cell.paragraphs[0]
    if first.runs:
        primary_run = first.runs[0]
        primary_run.text = text
        primary_run.font.size = Pt(10)
        for run in first.runs[1:]:
            run.text = ""
            run.font.size = Pt(10)
    else:
        primary_run = first.add_run(text)
        primary_run.font.size = Pt(10)
    for paragraph in cell.paragraphs[1:]:
        paragraph._element.getparent().remove(paragraph._element)


def _set_table_text(table, row: int, col: int, text: str) -> None:
    _set_cell_text(table.cell(row, col), text)


def _split_long_text(value: str, width: int, rows: int) -> list[str]:
    prepared = _text(value)
    if not prepared:
        return [""] * rows
    chunks = wrap(prepared, width=width, break_long_words=False, break_on_hyphens=False)
    chunks = chunks[:rows]
    while len(chunks) < rows:
        chunks.append("")
    return chunks


def _cargo_name(order: ShippingOrder) -> str:
    parts: list[str] = []
    total = order.items.count()
    for item in order.items.order_by("id")[:4]:
        name = _text(item.name) or "Товар"
        sku_code = _text(item.sku_code)
        qty = int(item.qty_requested or 0)
        chunk = name
        if sku_code:
            chunk = f"{chunk} (арт. {sku_code})"
        if qty > 0:
            chunk = f"{chunk} x {qty} шт."
        parts.append(chunk)
    if total > 4:
        parts.append(f"и еще {total - 4} поз.")
    return "; ".join(parts)


def _places_count(order: ShippingOrder) -> int:
    packing_summary = _shipping_packing_summary(order) or {}
    if int(packing_summary.get("box_count") or 0) > 0:
        return int(packing_summary["box_count"])
    if int(order.expected_boxes or 0) > 0:
        return int(order.expected_boxes)
    return 0


def _place_type_label(order: ShippingOrder) -> str:
    if getattr(order, "get_place_type_display", None) and _text(order.place_type):
        return _text(order.get_place_type_display())
    if order.supply_type == ShippingOrder.SUPPLY_MONOPALLET:
        return "Паллета"
    if order.supply_type == ShippingOrder.SUPPLY_SUPERSAFE:
        return "Супер сейф"
    return "Короб"


def _cargo_weight(order: ShippingOrder) -> Decimal | None:
    total_weight = Decimal("0")
    has_weight = False
    for item in order.items.select_related("sku"):
        sku = getattr(item, "sku", None)
        if sku is None or sku.weight_kg is None:
            continue
        qty = Decimal(str(int(item.qty_requested or 0)))
        total_weight += Decimal(str(sku.weight_kg)) * qty
        has_weight = True
    if not has_weight:
        return None
    return total_weight.quantize(Decimal("0.001"))


def _consignee_name(order: ShippingOrder) -> str:
    warehouse = _marketplace_warehouse_match(order)
    if warehouse:
        return _prefixed_marketplace_name(order, _text(warehouse.get("name")))
    destination = _text(order.destination_warehouse)
    marketplace = _marketplace_name(order)
    if marketplace and destination:
        return f"{marketplace} / {destination}"
    return destination or marketplace


def _consignee_address(order: ShippingOrder) -> str:
    warehouse = _marketplace_warehouse_match(order)
    return _first_nonempty(order.transit_address, order.destination_address, (warehouse or {}).get("address", ""), order.destination_warehouse)


def _legacy_consignee_address_candidates(order: ShippingOrder) -> set[str]:
    warehouse = _marketplace_warehouse_match(order) or {}
    candidates = {
        _text(order.transit_address),
        _text(order.destination_address),
        _text(warehouse.get("address", "")),
        _text(order.destination_warehouse),
        _consignee_address(order),
    }
    return {value for value in candidates if value}


def _legacy_consignee_name(order: ShippingOrder) -> str:
    destination = _text(order.destination_warehouse)
    marketplace = _text(getattr(order.marketplace, "name", ""))
    if marketplace and destination:
        return f"{marketplace} / {destination}"
    return destination or marketplace


def _legacy_shipper_name(order: ShippingOrder) -> str:
    return _text(order.agency.agn_name) or _text(order.agency.fio_agn)


def _legacy_shipper_address(order: ShippingOrder) -> str:
    return _first_nonempty(order.agency.adres, order.agency.fakt_adres)


def _legacy_customer_name(order: ShippingOrder) -> str:
    return _text(order.agency.agn_name) or _text(order.agency.fio_agn)


def _carrier_defaults(order: ShippingOrder) -> dict:
    trip_link = shipping_dispatch_trip_link(order)
    trip = trip_link.trip if trip_link else None
    own_company = _default_own_company()
    carrier = _default_carrier()
    if order.vehicle_type == ShippingOrder.VEHICLE_CLIENT:
        return {
            "carrier_name": _customer_name(order),
            "carrier_inn": _text(order.agency.inn),
            "carrier_address": _first_nonempty(order.agency.adres, order.agency.fakt_adres),
            "carrier_phone": _text(order.agency.phone),
            "driver_name": _text(getattr(trip, "driver_name", "")),
            "driver_phone": _first_nonempty(order.driver_phone, getattr(trip, "driver_phone", "")),
            "vehicle_number": _first_nonempty(order.vehicle_number, getattr(trip, "vehicle_number", "")),
        }
    if carrier is not None:
        return {
            "carrier_name": _carrier_name_from_directory(carrier),
            "carrier_inn": _text(carrier.inn),
            "carrier_address": _carrier_address_from_directory(carrier),
            "carrier_phone": _text(carrier.phone),
            "driver_name": _text(getattr(trip, "driver_name", "")),
            "driver_phone": _first_nonempty(order.driver_phone, getattr(trip, "driver_phone", "")),
            "vehicle_number": _first_nonempty(order.vehicle_number, getattr(trip, "vehicle_number", "")),
        }
    return {
        "carrier_name": _own_company_name(own_company),
        "carrier_inn": _own_company_inn(own_company),
        "carrier_address": _own_company_address(own_company),
        "carrier_phone": _own_company_phone(own_company),
        "driver_name": _text(getattr(trip, "driver_name", "")),
        "driver_phone": _first_nonempty(order.driver_phone, getattr(trip, "driver_phone", "")),
        "vehicle_number": _first_nonempty(order.vehicle_number, getattr(trip, "vehicle_number", "")),
    }


def default_transport_note_data(order: ShippingOrder) -> dict:
    own_company = _default_own_company()
    shipper_address = _own_company_address(own_company)
    shipper_name = _own_company_name(own_company)
    carrier_defaults = _carrier_defaults(order)
    doc_date = timezone.localdate(order.shipped_at) if order.shipped_at else timezone.localdate()
    return {
        "document_number": _text(order.number),
        "document_date": doc_date,
        "shipper_name": shipper_name,
        "shipper_inn": _own_company_inn(own_company),
        "shipper_address": shipper_address,
        "shipper_phone": _own_company_phone(own_company),
        "consignee_name": _consignee_name(order),
        "consignee_inn": "",
        "consignee_address": _consignee_address(order),
        "consignee_phone": "",
        "loading_address": _own_company_address(own_company),
        "unloading_address": _consignee_address(order),
        "cargo_name": _cargo_name(order),
        "cargo_package_count": _places_count(order),
        "cargo_package_type": _place_type_label(order),
        "cargo_weight_kg": _cargo_weight(order),
        "cargo_declared_value": None,
        "accompanying_documents": f"Заявка на отгрузку №{order.number} от {timezone.localtime(order.created_at).strftime('%d.%m.%Y')}",
        "special_instructions": _text(order.comment),
        "transportation_conditions": "",
        "delivery_notes": "",
        "trailer_number": "",
        "service_cost": None,
        **carrier_defaults,
    }


def get_or_create_transport_note(order: ShippingOrder) -> ShippingTransportNote:
    defaults = default_transport_note_data(order)
    note, created = ShippingTransportNote.objects.get_or_create(order=order, defaults=defaults)
    if created:
        return note

    own_company = _default_own_company()
    carrier = _default_carrier()
    warehouse = _marketplace_warehouse_match(order)
    legacy_values = {
        "shipper_name": {_legacy_shipper_name(order)},
        "shipper_inn": {_text(order.agency.inn)},
        "shipper_address": {_legacy_shipper_address(order)},
        "shipper_phone": {_text(order.agency.phone)},
        "consignee_name": {_legacy_consignee_name(order), _text((warehouse or {}).get("name", ""))},
        "consignee_address": _legacy_consignee_address_candidates(order),
        "unloading_address": _legacy_consignee_address_candidates(order),
        "carrier_name": {FULLBOX_NAME},
        "carrier_inn": {""},
        "carrier_address": {FULLBOX_WAREHOUSE_ADDRESS},
        "carrier_phone": {FULLBOX_PHONE},
    }
    if own_company is not None:
        legacy_values["carrier_name"].add(_own_company_name(own_company))
        legacy_values["carrier_inn"].add(_own_company_inn(own_company))
        legacy_values["carrier_address"].add(_own_company_address(own_company))
        legacy_values["carrier_phone"].add(_own_company_phone(own_company))
    if carrier is not None:
        legacy_values["carrier_name"].add(_carrier_name_from_directory(carrier))
        legacy_values["carrier_inn"].add(_text(carrier.inn))
        legacy_values["carrier_address"].add(_carrier_address_from_directory(carrier))
        legacy_values["carrier_phone"].add(_text(carrier.phone))

    update_fields: list[str] = []
    for field, value in defaults.items():
        current = getattr(note, field)
        current_text = _text(current)
        value_text = _text(value)
        if current not in ("", None, 0):
            legacy_match = field in legacy_values and current_text in {_text(item) for item in legacy_values[field] if _text(item)}
            if not (legacy_match and value_text and current_text != value_text):
                continue
        if value in ("", None, 0):
            continue
        setattr(note, field, value)
        update_fields.append(field)
    if update_fields:
        note.save(update_fields=[*update_fields, "updated_at"])
    return note


def transport_note_item_rows(order: ShippingOrder) -> list[dict]:
    rows: list[dict] = []
    for index, item in enumerate(order.items.order_by("id"), start=1):
        qty = int(item.qty_shipped or 0) or int(item.qty_requested or 0)
        rows.append(
            {
                "index": index,
                "sku_code": _text(item.sku_code) or "-",
                "name": _text(item.name) or "-",
                "size": _text(item.size) or "-",
                "barcode": _text(item.barcode) or "-",
                "qty": qty,
            }
        )
    return rows


def build_transport_note_preview_context(order: ShippingOrder, note: ShippingTransportNote) -> dict:
    rows = transport_note_item_rows(order)
    trip_link = shipping_dispatch_trip_link(order)
    trip = trip_link.trip if trip_link else None
    own_company = _default_own_company()
    customer_identity = _agency_customer_identity(order)
    customer_identity_pdf = _format_party_identity_pdf(
        name=_customer_name(order),
        inn=_text(order.agency.inn),
        address=_first_nonempty(order.agency.adres, order.agency.fakt_adres),
        phone=_text(order.agency.phone),
    )
    shipper_identity = _format_party_identity(
        name=note.shipper_name,
        inn=note.shipper_inn,
        address=note.shipper_address,
        phone=note.shipper_phone,
    )
    shipper_identity_pdf = _format_party_identity_pdf(
        name=note.shipper_name,
        inn=note.shipper_inn,
        address=note.shipper_address,
        phone=note.shipper_phone,
    )
    consignee_identity = _format_party_identity(
        name=note.consignee_name,
        inn=note.consignee_inn,
        address=note.consignee_address,
        phone=note.consignee_phone,
    )
    consignee_identity_pdf = _format_party_identity_pdf(
        name=note.consignee_name,
        inn=note.consignee_inn,
        address=note.consignee_address,
        phone=note.consignee_phone,
    )
    carrier_identity = _format_party_identity(
        name=note.carrier_name,
        inn=note.carrier_inn,
        address=note.carrier_address,
        phone=note.carrier_phone,
    )
    carrier_identity_pdf = _format_party_identity_pdf(
        name=note.carrier_name,
        inn=note.carrier_inn,
        address=note.carrier_address,
        phone=note.carrier_phone,
    )
    shipper_meta_pdf = _join_nonempty(
        f"ИНН {note.shipper_inn}" if _text(note.shipper_inn) else "",
        _short_address(note.shipper_address, limit=64),
        sep=", ",
    )
    customer_name_pdf = _customer_name(order)
    customer_meta_pdf = _join_nonempty(
        f"ИНН {order.agency.inn}" if _text(order.agency.inn) else "",
        _short_address(_first_nonempty(order.agency.adres, order.agency.fakt_adres), limit=68),
        sep=", ",
    )
    consignee_name_pdf = _text(note.consignee_name)
    consignee_meta_pdf = _short_address(note.consignee_address, limit=120)
    driver_identity = _join_nonempty(note.driver_name, note.driver_phone, sep=", ")
    loading_window = _format_eta_window(order)
    unloading_window = _format_slot_window(order) or _format_eta_window(order)
    vehicle_docs = _join_nonempty(
        f"Тип авто: {order.get_vehicle_type_display()}" if _text(order.vehicle_type) else "",
        f"Транспорт рейса: {trip.vehicle_name}" if trip and _text(trip.vehicle_name) else "",
        sep="; ",
    )
    cargo_places = _join_nonempty(
        f"{int(note.cargo_package_count or 0)} {note.cargo_package_type}".strip() if int(note.cargo_package_count or 0) else "",
        sep="",
    )
    page1_slots = [
        _slot(72.0, 59.2, 30.0, note.document_number or order.number, css_class="slot-sm slot-center", height_mm=3.2),
        _slot(118.5, 59.2, 30.0, _format_date(note.document_date), css_class="slot-sm slot-center", height_mm=3.2),
        _slot(31.0, 64.2, 77.0, note.shipper_name, css_class="slot-xxs", height_mm=3.0),
        _slot(31.0, 71.8, 77.0, shipper_meta_pdf, css_class="slot-xxs", height_mm=4.1),
        _slot(120.0, 70.0, 81.0, customer_name_pdf, css_class="slot-xxs", height_mm=3.0),
        _slot(120.0, 77.2, 81.0, customer_meta_pdf, css_class="slot-xxs", height_mm=4.2),
        _slot(148.4, 84.2, 33.2, _text(order.agency.contract_numb), css_class="slot-xs slot-center", height_mm=3.2),
        _slot(78.0, 98.6, 104.0, consignee_name_pdf, css_class="slot-xxs", height_mm=3.0),
        _slot(78.0, 105.2, 104.0, consignee_meta_pdf, css_class="slot-xxs", height_mm=4.2),
        _slot(73.4, 111.0, 70.0, note.unloading_address, css_class="slot-xxs", height_mm=3.8),
        _slot(6.1, 125.2, 97.0, note.cargo_name, css_class="slot-xxs", height_mm=7.2),
        _slot(148.0, 125.4, 17.0, cargo_places, css_class="slot-sm slot-center", height_mm=3.2),
        _slot(101.5, 107.0, 21.2, _format_weight(note.cargo_weight_kg), css_class="slot-xs slot-center", height_mm=4.5),
        _slot(162.3, 107.0, 28.2, _format_currency(note.cargo_declared_value), css_class="slot-xs slot-center", height_mm=4.5),
        _slot(10.6, 161.8, 187.0, note.accompanying_documents, css_class="slot-xxs", height_mm=5.0),
        _slot(8.8, 183.4, 90.0, note.special_instructions, css_class="slot-xs", height_mm=14.0),
        _slot(105.8, 183.4, 90.0, note.transportation_conditions, css_class="slot-xs", height_mm=14.0),
        _slot(52.9, 213.1, 63.5, carrier_identity_pdf, css_class="slot-xs", height_mm=6.5),
        _slot(151.7, 213.1, 45.9, driver_identity, css_class="slot-xs", height_mm=7.0),
        _slot(
            8.8,
            221.9,
            88.2,
            _join_nonempty(note.vehicle_number, note.trailer_number, sep=", "),
            css_class="slot-xs",
            height_mm=7.0,
        ),
        _slot(105.8, 221.9, 88.2, vehicle_docs, css_class="slot-xs", height_mm=7.0),
    ]
    page2_slots = [
        _slot(70.8, 10.3, 68.3, shipper_identity_pdf or customer_identity_pdf, css_class="slot-xs", height_mm=6.2),
        _slot(70.6, 18.7, 67.0, _own_company_name(own_company), css_class="slot-xs", height_mm=4.5),
        _slot(7.1, 27.2, 90.0, note.loading_address, css_class="slot-xs", height_mm=7.0),
        _slot(132.6, 27.0, 44.8, loading_window, css_class="slot-xs slot-center", height_mm=6.0),
        _slot(10.6, 36.0, 77.6, loading_window, css_class="slot-xs", height_mm=6.0),
        _slot(123.5, 36.0, 63.5, loading_window, css_class="slot-xs", height_mm=6.0),
        _slot(60.0, 48.7, 42.3, _format_weight(note.cargo_weight_kg), css_class="slot-xs slot-center", height_mm=4.5),
        _slot(148.2, 48.7, 28.2, cargo_places, css_class="slot-xs slot-center", height_mm=4.5),
        _slot(8.8, 64.6, 190.5, note.delivery_notes, css_class="slot-xs", height_mm=11.0),
        _slot(7.1, 117.1, 88.2, note.unloading_address, css_class="slot-xs", height_mm=7.0),
        _slot(134.1, 117.1, 52.9, unloading_window, css_class="slot-xs slot-center", height_mm=6.0),
        _slot(8.8, 129.1, 77.6, unloading_window, css_class="slot-xs", height_mm=6.0),
        _slot(123.5, 129.1, 63.5, unloading_window, css_class="slot-xs", height_mm=6.0),
        _slot(7.1, 137.6, 91.7, note.delivery_notes, css_class="slot-xs", height_mm=6.0),
        _slot(151.7, 137.6, 28.2, cargo_places, css_class="slot-xs slot-center", height_mm=4.5),
        _slot(7.1, 146.8, 77.6, _format_weight(note.cargo_weight_kg), css_class="slot-xs", height_mm=4.5),
        _slot(123.5, 146.8, 70.6, note.delivery_notes, css_class="slot-xs", height_mm=4.5),
        _slot(7.1, 177.8, 116.4, note.special_instructions or note.transportation_conditions, css_class="slot-xs", height_mm=12.0),
        _slot(176.4, 177.8, 21.2, note.driver_name, css_class="slot-xs slot-center", height_mm=4.5),
        _slot(7.1, 193.3, 42.3, _format_currency(note.service_cost), css_class="slot-xs slot-center", height_mm=4.5),
        _slot(65.3, 193.3, 21.2, "без НДС" if note.service_cost else "", css_class="slot-xs slot-center", height_mm=4.5),
        _slot(105.8, 193.3, 28.2, "", css_class="slot-xs slot-center", height_mm=4.5),
        _slot(165.8, 193.3, 31.7, _format_currency(note.service_cost), css_class="slot-xs slot-center", height_mm=4.5),
        _slot(52.9, 205.3, 102.3, "безналичный расчет" if note.service_cost else "", css_class="slot-xs slot-center", height_mm=4.5),
        _slot(106.1, 212.7, 97.7, carrier_identity_pdf, css_class="slot-xs", height_mm=7.2),
        _slot(
            11.6,
            234.6,
            77.6,
            _join_nonempty(f"Заявка {order.number}", _format_date(note.document_date), sep=" от "),
            css_class="slot-xs",
            height_mm=5.0,
        ),
        _slot(
            111.5,
            234.6,
            77.6,
            _join_nonempty(_text(order.agency.contract_numb), _format_date(note.document_date), sep=" от "),
            css_class="slot-xs",
            height_mm=5.0,
        ),
        _slot(54.7, 246.9, 116.4, customer_identity_pdf, css_class="slot-xs", height_mm=6.8),
    ]
    return {
        "item_rows": rows,
        "total_qty": sum(int(row["qty"] or 0) for row in rows),
        "places_count": int(note.cargo_package_count or 0),
        "has_weight": note.cargo_weight_kg is not None,
        "return_url": f"/shipping/{order.pk}/",
        "print_title": note.document_number or order.number,
        "document_date_label": _format_date(note.document_date),
        "customer_name_label": _customer_name(order),
        "customer_identity": customer_identity,
        "contract_label": _text(order.agency.contract_numb),
        "shipper_identity": shipper_identity,
        "consignee_identity": consignee_identity,
        "carrier_identity": carrier_identity,
        "driver_identity": driver_identity,
        "cargo_places_label": cargo_places,
        "cargo_weight_label": _format_weight(note.cargo_weight_kg),
        "cargo_declared_value_label": _format_currency(note.cargo_declared_value),
        "service_cost_label": _format_currency(note.service_cost),
        "loading_window_label": loading_window,
        "unloading_window_label": unloading_window,
        "vehicle_docs_label": vehicle_docs,
        "loading_party_label": shipper_identity or customer_identity,
        "loading_point_owner_label": _own_company_name(own_company),
        "delivery_party_label": consignee_identity,
        "page1_slots": [slot for slot in page1_slots if slot["text"]],
        "page2_slots": [slot for slot in page2_slots if slot["text"]],
    }


def render_transport_note_pdf(order: ShippingOrder, note: ShippingTransportNote) -> bytes:
    preview = build_transport_note_preview_context(order, note)
    document = fitz.open(str(TRANSPORT_NOTE_TEMPLATE_PDF))
    try:
        _draw_pdf_slots(document[0], preview["page1_slots"])
        _draw_pdf_slots(document[1], preview["page2_slots"])
        return document.tobytes(garbage=4, deflate=True)
    finally:
        document.close()


def render_transport_note_docx(order: ShippingOrder, note: ShippingTransportNote) -> bytes:
    document = Document(str(TRANSPORT_NOTE_TEMPLATE_DOCX))
    customer_name = _customer_name(order)
    shipper_identity_docx = _format_party_identity_docx(
        name=note.shipper_name,
        inn=note.shipper_inn,
        address=note.shipper_address,
        phone=note.shipper_phone,
    )
    customer_identity_docx = _format_party_identity_docx(
        name=customer_name,
        inn=_text(order.agency.inn),
        address=_first_nonempty(order.agency.adres, order.agency.fakt_adres),
        phone=_text(order.agency.phone),
    )
    consignee_identity_docx = _format_party_identity_docx(
        name=note.consignee_name,
        inn=note.consignee_inn,
        address=note.consignee_address,
        phone=note.consignee_phone,
    )
    carrier_identity_docx = _format_party_identity_docx(
        name=note.carrier_name,
        inn=note.carrier_inn,
        address=note.carrier_address,
        phone=note.carrier_phone,
    )
    cargo_places = _join_nonempty(
        str(int(note.cargo_package_count or 0)) if int(note.cargo_package_count or 0) else "",
        note.cargo_package_type,
        sep=" ",
    )
    loading_window = _format_eta_window(order)
    unloading_window = _format_slot_window(order) or _format_eta_window(order)
    cargo_lines = _split_long_text(note.cargo_name, width=110, rows=2)
    docs_lines = _split_long_text(note.accompanying_documents, width=180, rows=3)
    contract_label = _text(order.agency.contract_numb)
    trip_link = shipping_dispatch_trip_link(order)
    trip = trip_link.trip if trip_link else None
    vehicle_docs = _join_nonempty(
        f"Тип авто: {order.get_vehicle_type_display()}" if _text(order.vehicle_type) else "",
        f"Транспорт рейса: {trip.vehicle_name}" if trip and _text(trip.vehicle_name) else "",
        sep="; ",
    )
    order_date = _format_date(timezone.localtime(order.created_at).date() if order.created_at else note.document_date)
    vehicle_identity = _join_nonempty(note.vehicle_number, note.trailer_number, sep=", ")
    weight_label = _format_weight(note.cargo_weight_kg)
    declared_value_label = _format_currency(note.cargo_declared_value)
    service_cost_label = _format_currency(note.service_cost)
    package_count_label = str(int(note.cargo_package_count or 0)) if int(note.cargo_package_count or 0) else ""

    if len(document.tables) >= 2:
        table1 = document.tables[0]
        table2 = document.tables[1]

        _set_table_text(table1, 1, 1, _format_date(note.document_date))
        _set_table_text(table1, 1, 4, note.document_number or order.number)
        _set_table_text(table1, 1, 6, order_date)
        _set_table_text(table1, 1, 10, order.number)
        _set_table_text(table1, 2, 3, "1")
        _set_table_text(table1, 5, 0, shipper_identity_docx)
        _set_table_text(table1, 5, 5, customer_identity_docx)
        _set_table_text(table1, 7, 5, contract_label)
        _set_table_text(table1, 10, 0, note.consignee_name)
        _set_table_text(table1, 12, 0, note.unloading_address)
        _set_table_text(table1, 15, 0, cargo_lines[0])
        _set_table_text(table1, 15, 5, cargo_places)
        _set_table_text(table1, 17, 0, cargo_lines[1])
        _set_table_text(table1, 17, 5, weight_label)
        _set_table_text(table1, 19, 5, declared_value_label)
        _set_table_text(table1, 22, 0, docs_lines[0])
        _set_table_text(table1, 24, 0, docs_lines[1])
        _set_table_text(table1, 26, 0, docs_lines[2])
        _set_table_text(table1, 29, 0, note.special_instructions)
        _set_table_text(table1, 31, 5, note.transportation_conditions)
        _set_table_text(table1, 34, 0, carrier_identity_docx)
        _set_table_text(table1, 37, 0, vehicle_docs)
        _set_table_text(table1, 37, 5, vehicle_identity)

        _set_table_text(table2, 5, 0, note.loading_address)
        _set_table_text(table2, 5, 2, loading_window)
        _set_table_text(table2, 11, 0, package_count_label)
        _set_table_text(table2, 11, 2, note.cargo_package_type)
        _set_table_text(table2, 13, 0, note.delivery_notes)
        _set_table_text(table2, 24, 0, note.unloading_address)
        _set_table_text(table2, 24, 2, unloading_window)
        _set_table_text(table2, 28, 2, cargo_places)
        _set_table_text(table2, 38, 0, service_cost_label)
        _set_table_text(table2, 38, 1, "без НДС" if note.service_cost else "")
        _set_table_text(table2, 38, 4, service_cost_label)
        _set_table_text(table2, 42, 0, carrier_identity_docx)
        _set_table_text(table2, 44, 2, shipper_identity_docx)
    else:
        table = document.tables[0]
        _set_table_text(table, 1, 1, _format_date(note.document_date))
        _set_table_text(table, 1, 6, note.document_number or order.number)
        _set_table_text(table, 1, 8, order_date)
        _set_table_text(table, 1, 14, order.number)
        _set_table_text(table, 2, 5, "1")
        _set_table_text(table, 3, 2, note.shipper_name)
        _set_table_text(table, 6, 0, shipper_identity_docx)
        _set_table_text(table, 6, 7, customer_identity_docx)
        _set_table_text(table, 9, 0, contract_label)
        _set_table_text(table, 12, 0, note.consignee_name)
        _set_table_text(table, 14, 0, note.unloading_address)
        _set_table_text(table, 17, 0, cargo_lines[0])
        _set_table_text(table, 17, 7, cargo_places)
        _set_table_text(table, 19, 0, cargo_lines[1])
        _set_table_text(table, 19, 7, weight_label)
        _set_table_text(table, 21, 7, declared_value_label)
        _set_table_text(table, 24, 0, docs_lines[0])
        _set_table_text(table, 26, 0, docs_lines[1])
        _set_table_text(table, 28, 0, docs_lines[2])
        _set_table_text(table, 31, 0, note.special_instructions)
        _set_table_text(table, 33, 7, note.transportation_conditions)
        _set_table_text(table, 36, 0, carrier_identity_docx)
        _set_table_text(table, 39, 0, vehicle_identity)
        _set_table_text(table, 53, 0, note.loading_address)
        _set_table_text(table, 53, 7, loading_window)
        _set_table_text(table, 57, 0, weight_label)
        _set_table_text(table, 59, 0, package_count_label)
        _set_table_text(table, 59, 7, note.cargo_package_type)
        _set_table_text(table, 61, 0, note.delivery_notes)
        _set_table_text(table, 72, 0, note.unloading_address)
        _set_table_text(table, 72, 7, unloading_window)
        _set_table_text(table, 76, 7, cargo_places)
        _set_table_text(table, 86, 0, service_cost_label)
        _set_table_text(table, 86, 4, "без НДС" if note.service_cost else "")
        _set_table_text(table, 86, 12, service_cost_label)

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()

def transport_note_filename(order: ShippingOrder, note: ShippingTransportNote) -> str:
    number = _text(note.document_number) or _text(order.number) or str(order.pk)
    safe_number = "".join(ch if ch.isalnum() or ch in ("-", "_") else "_" for ch in number)
    return f"transport-note-{safe_number}.pdf"


def transport_note_docx_filename(order: ShippingOrder, note: ShippingTransportNote) -> str:
    number = _text(note.document_number) or _text(order.number) or str(order.pk)
    safe_number = "".join(ch if ch.isalnum() or ch in ("-", "_") else "_" for ch in number)
    return f"transport-note-{safe_number}.docx"
