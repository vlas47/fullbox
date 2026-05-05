"""Processing UI views and helper logic."""

import json
import base64
import io
import re
import uuid
import zipfile
from datetime import datetime, timedelta
from pathlib import Path
from urllib.parse import urlparse, urlencode

from django.conf import settings
from django.contrib.auth.decorators import login_required
from django.db import transaction, DatabaseError
from django.db.models import Count, Q, Sum
from django.shortcuts import redirect
from django.http import HttpResponse, HttpResponseBadRequest, HttpResponseForbidden, JsonResponse
from django.utils import timezone
from django.views.decorators.csrf import csrf_exempt
from django.views.decorators.http import require_GET, require_POST
from django.views.generic import TemplateView

from audit.models import OrderAuditEntry, log_order_action, log_staff_overaction, log_stock_move
from employees.access import (
    RoleRequiredMixin,
    get_employee_for_user,
    get_request_role,
    resolve_cabinet_url,
    is_staff_role,
)
from employees.models import Employee
from labels.utils import (
    LABEL_SIZES,
    SCANNER_EOLS,
    load_available_printers_data,
    load_label_settings,
    load_print_agent_status,
    load_scanner_settings,
    save_print_agent_status,
    set_print_agent_pause,
)
from marking.models import MarkingCode
from marking.utils import extract_processing_items
from orders.views import (
    OrdersDetailView,
    ProcessingPlacementActView as OrdersProcessingPlacementActView,
    ReceivingFlowView as OrdersReceivingFlowView,
    _OS_CELLS_PER_TIER,
    _OS_ROW_SECTIONS,
    _OS_TIERS,
    _current_status_entry,
    _flow_closed_from_entries,
    _item_key,
    _latest_payload_from_entries,
    _parse_int_value,
    _processing_receiving_items,
    _shorten_ip_name,
    _status_label_from_entry,
)
from sku.models import Agency, SKU, SKUBarcode
from sklad.services.warehouse_commands import WarehouseCommandService
from sklad.services.warehouse_policy import WarehouseActionPolicy
from sklad.services.stock_operations import OperationalStockService
from sklad.services.warehouse_state import WarehouseGoodsStateResolver
from sklad.services.warehouse_transitions import WarehouseStateCode
from sklad.services.warehouse_write_path import WarehouseWritePathService
from sklad.models import WarehouseReserve, WarehouseStockSnapshot
from sklad.services.stock_availability import StockAvailabilityService
from todo.models import Task
from reachtruck.services import (
    create_batch_move_tasks,
    sync_task_status_by_legacy_order_id,
)
from reachtruck.services.putaway_planner import (
    normalize_putaway_location,
    normalize_zone_code,
    parse_putaway_destinations,
    putaway_location_label,
    suggest_putaway_destinations,
)
from .models import ProcessingFlowSession, ProcessingPrintJob
from .services import ProcessingWorkflowService
from agent.models import AgentCommand, DeviceAgent
from openpyxl import load_workbook

GOODS_TYPE_LABELS = {
    "op": "Оптовый",
    "gv": "Готовый",
    "br": "Брак",
    "vz": "Возврат",
    "rh": "Расходный",
    "no": "Не обработанный",
}

_PROCESSING_WORK_WAREHOUSE_CODES = {
    WarehouseStateCode.RESERVED_FOR_PROCESSING,
    WarehouseStateCode.MOVING_TO_PROCESSING,
    WarehouseStateCode.IN_PROCESSING_ZONE,
    WarehouseStateCode.PROCESSING_IN_PROGRESS,
    WarehouseStateCode.PLACED_AFTER_PROCESSING,
    WarehouseStateCode.STORED,
}

_PROCESSING_CARD_WAREHOUSE_CODES = _PROCESSING_WORK_WAREHOUSE_CODES

PROCESSING_MARKING_LABELS = (
    {
        "field": "marking_5840_qty",
        "needed_field": "marking_5840_needed",
        "label": "Маркировка 58/40",
        "label_key": "item",
        "label_type": "58/40",
        "size_code": "58X40",
    },
    {
        "field": "marking_5860_qty",
        "needed_field": "marking_5860_needed",
        "label": "Маркировка 58/60",
        "label_key": "item_5860",
        "label_type": "58/60",
        "size_code": "58X60",
    },
    {
        "field": "marking_75120_qty",
        "needed_field": "marking_75120_needed",
        "label": "Маркировка 75/120",
        "label_key": "item_75120",
        "label_type": "75/120",
        "size_code": "75X120",
    },
)
PROCESSING_CZ_LABEL_KEY = "item_cz"
PROCESSING_CZ_LABEL_TYPE = "58/40 (шт/чз)"
PROCESSING_CZ_SIZE_CODE = "58X40"


def _format_payload_value(value):
    if value is None or value == "":
        return "-"
    return str(value).strip()


def _format_payload_list(value):
    if value is None or value == "":
        return "-"
    if isinstance(value, (list, tuple, set)):
        items = [str(item).strip() for item in value if str(item).strip()]
        return ", ".join(items) if items else "-"
    text = str(value).strip()
    return text if text else "-"


_ORG_FORM_RE = re.compile(r"\bобщество\s+с\s+ограниченной\s+ответственностью\b", re.IGNORECASE)


def _normalize_org_name(value) -> str:
    if value is None:
        return ""
    text = str(value)
    if not text.strip():
        return text.strip()
    text = _ORG_FORM_RE.sub("ООО", text)
    text = re.sub(r"\s{2,}", " ", text).strip()
    return text


def _non_empty_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Да" if value else ""
    if isinstance(value, (int, float)):
        return "" if value == 0 else str(value)
    text = str(value).strip()
    if not text or text in {"-", "0", "0.0"}:
        return ""
    if text.isdigit() and int(text) == 0:
        return ""
    return text


def _normalize_marking_size_code(value) -> str:
    return re.sub(r"[^0-9A-Z]", "", str(value or "").upper())


def _processing_marking_sizes_from_payload(payload: dict | None) -> list[str]:
    source = payload or {}
    raw_sizes = source.get("marking_sizes") or []
    if isinstance(raw_sizes, str):
        raw_sizes = [raw_sizes] if raw_sizes else []
    ordered: list[str] = []
    seen: set[str] = set()

    def add_size(raw_value) -> None:
        code = _normalize_marking_size_code(raw_value)
        if not code or code in seen:
            return
        seen.add(code)
        ordered.append(code)

    for raw_value in raw_sizes:
        add_size(raw_value)
    for option in PROCESSING_MARKING_LABELS:
        if (_parse_qty_value(source.get(option["field"])) or 0) > 0:
            add_size(option["size_code"])
    if (_parse_qty_value(source.get("marking_5840_each_qty")) or 0) > 0:
        add_size(PROCESSING_CZ_SIZE_CODE)
    return ordered


def _processing_marking_qty_by_label_key(payload: dict | None) -> dict[str, int]:
    source = payload or {}
    qty_map: dict[str, int] = {}
    for option in PROCESSING_MARKING_LABELS:
        qty_map[option["label_key"]] = _parse_qty_value(source.get(option["field"])) or 0
    qty_map[PROCESSING_CZ_LABEL_KEY] = _parse_qty_value(source.get("marking_5840_each_qty")) or 0
    return qty_map


def _processing_marking_label_key_from_dimensions(width_mm, height_mm) -> str:
    try:
        width = int(width_mm or 0)
        height = int(height_mm or 0)
    except (TypeError, ValueError):
        return ""
    dims = (width, height)
    if dims == (58, 40):
        return "item"
    if dims == (58, 60):
        return "item_5860"
    if dims == (75, 120):
        return "item_75120"
    return ""


def _format_list_value(value) -> str:
    if value is None:
        return ""
    if isinstance(value, (list, tuple, set)):
        items = [str(item).strip() for item in value if str(item).strip()]
        return ", ".join(items)
    return _non_empty_text(value)


def _processing_card_id(card: dict) -> str:
    if not isinstance(card, dict):
        return ""
    value = card.get("id") or card.get("article") or card.get("sku") or ""
    return str(value).strip()


def _processing_card_sets(payload: dict) -> tuple[set[str], set[str]]:
    payload = payload or {}
    processed = set()
    placed = set()
    for card in payload.get("cards") or []:
        card_id = _processing_card_id(card)
        if not card_id:
            continue
        if card.get("processed_at") or card.get("processed_done") or card.get("processed"):
            processed.add(card_id)
        if card.get("placed_at") or card.get("placed_done"):
            placed.add(card_id)
    for value in payload.get("processed_cards") or []:
        text = str(value).strip()
        if text:
            processed.add(text)
    for value in payload.get("placed_cards") or []:
        text = str(value).strip()
        if text:
            placed.add(text)
    return processed, placed


def _parse_json_value(raw, fallback):
    if raw is None:
        return fallback
    if isinstance(raw, (dict, list)):
        return raw
    text = str(raw).strip()
    if not text:
        return fallback
    try:
        return json.loads(text)
    except (TypeError, ValueError):
        return fallback


def _parse_json_body(request):
    if not request.body:
        return {}
    try:
        return json.loads(request.body.decode("utf-8"))
    except (UnicodeDecodeError, json.JSONDecodeError):
        return None


def _default_flow_state() -> dict:
    return {
        "boxes": [],
        "pallets": [],
        "activeBox": "",
        "activePallet": "",
    }


def _flow_session_for_request(order_id: str, agent_id: str, request, create: bool = False):
    agent_id = str(agent_id or "").strip()
    if not agent_id:
        return None
    user = request.user if request.user.is_authenticated else None
    employee = get_employee_for_user(request.user) if user else None
    qs = ProcessingFlowSession.objects.filter(
        order_id=order_id,
        order_type="processing",
        agent_id=agent_id,
        status=ProcessingFlowSession.STATUS_OPEN,
    )
    if user:
        qs = qs.filter(user=user)
    session = qs.order_by("-updated_at").first()
    if not session and create:
        session = ProcessingFlowSession.objects.create(
            order_id=order_id,
            order_type="processing",
            agent_id=agent_id,
            user=user,
            employee=employee,
            flow_state=_default_flow_state(),
            status=ProcessingFlowSession.STATUS_OPEN,
            last_seen=timezone.localtime(),
        )
        return session
    if session:
        update_fields = []
        now = timezone.localtime()
        if session.last_seen != now:
            session.last_seen = now
            update_fields.append("last_seen")
        if user and session.user_id != user.id:
            session.user = user
            update_fields.append("user")
        if employee and session.employee_id != employee.id:
            session.employee = employee
            update_fields.append("employee")
        normalized_state = _sanitize_flow_state_for_session(session.flow_state, session)
        if normalized_state != (session.flow_state if isinstance(session.flow_state, dict) else _default_flow_state()):
            session.flow_state = normalized_state
            update_fields.append("flow_state")
        if update_fields:
            save_fields = list(dict.fromkeys(update_fields))
            if "flow_state" in save_fields or "status" in save_fields:
                save_fields.append("updated_at")
            session.save(update_fields=save_fields)
    return session


def _flow_owner_label(session: ProcessingFlowSession | None) -> str:
    if not session:
        return ""
    employee = getattr(session, "employee", None)
    if employee and getattr(employee, "full_name", None):
        return employee.full_name
    user = getattr(session, "user", None)
    if user:
        full_name = user.get_full_name() or ""
        if full_name:
            return full_name
        return getattr(user, "username", "") or ""
    return ""


def _resolve_actor_label(label, agent_id, user_id, *, unknown: str = "") -> str:
    text = str(label or "").strip()
    if text:
        return text
    agent = str(agent_id or "").strip()
    if agent and agent != "__reopen_snapshot__":
        return agent
    user = str(user_id or "").strip()
    if user:
        return f"ID {user}"
    return str(unknown or "").strip()


def _entry_belongs_to_session(entry: dict, session: ProcessingFlowSession | None) -> bool:
    if not session:
        return True
    owner_agent = str(entry.get("owner_agent_id") or "").strip()
    owner_user_id = str(entry.get("owner_user_id") or "").strip()
    owner_label = str(entry.get("owner_user_label") or "").strip().casefold()
    session_agent = str(session.agent_id or "").strip()
    session_user_id = str(session.user_id or "").strip()
    session_labels = set()
    session_label = str(_flow_owner_label(session) or "").strip()
    if session_label:
        session_labels.add(session_label.casefold())
    employee = getattr(session, "employee", None)
    if employee and getattr(employee, "full_name", None):
        session_labels.add(str(employee.full_name).strip().casefold())
    user = getattr(session, "user", None)
    if user:
        full_name = str(user.get_full_name() or "").strip()
        username = str(getattr(user, "username", "") or "").strip()
        if full_name:
            session_labels.add(full_name.casefold())
        if username:
            session_labels.add(username.casefold())
    if owner_user_id:
        return bool(session_user_id and owner_user_id == session_user_id)
    if owner_agent:
        if not session_agent:
            return False
        if owner_agent != session_agent:
            return False
    if owner_label:
        if not session_labels:
            return False
        return owner_label in session_labels
    return True


def _apply_flow_owner(values: list, session: ProcessingFlowSession | None) -> list:
    if not session:
        return values
    owner_agent = session.agent_id or ""
    owner_user_id = session.user_id
    owner_label = _flow_owner_label(session)
    patched = []
    for raw in values or []:
        if not isinstance(raw, dict):
            continue
        entry = dict(raw)
        code = str(entry.get("code") or "").strip()
        if not code:
            continue
        entry["code"] = code
        if owner_agent and not entry.get("owner_agent_id"):
            entry["owner_agent_id"] = owner_agent
        if owner_user_id and not entry.get("owner_user_id"):
            entry["owner_user_id"] = owner_user_id
        if owner_label and not entry.get("owner_user_label"):
            entry["owner_user_label"] = owner_label
        patched.append(entry)
    return patched


def _filter_flow_values_for_session(values: list, session: ProcessingFlowSession | None) -> list:
    filtered = []
    for raw in values or []:
        if not isinstance(raw, dict):
            continue
        entry = dict(raw)
        code = str(entry.get("code") or "").strip()
        if not code:
            continue
        entry["code"] = code
        if not _entry_belongs_to_session(entry, session):
            continue
        filtered.append(entry)
    return filtered


def _sanitize_flow_state_for_session(flow_state: dict | None, session: ProcessingFlowSession | None) -> dict:
    source = flow_state if isinstance(flow_state, dict) else _default_flow_state()
    boxes_raw = source.get("boxes") or []
    if not isinstance(boxes_raw, list):
        boxes_raw = []
    boxes = []
    for raw in boxes_raw:
        if not isinstance(raw, dict):
            continue
        entry = dict(raw)
        code = str(entry.get("code") or "").strip()
        if not code:
            continue
        entry["code"] = code
        boxes.append(entry)
    boxes = _apply_flow_owner(boxes, session)
    box_codes = {str(box.get("code") or "").strip() for box in boxes if isinstance(box, dict)}
    pallets_raw = source.get("pallets") or []
    if not isinstance(pallets_raw, list):
        pallets_raw = []
    pallets = []
    for raw in pallets_raw:
        if not isinstance(raw, dict):
            continue
        entry = dict(raw)
        code = str(entry.get("code") or "").strip()
        if not code:
            continue
        entry["code"] = code
        pallets.append(entry)
    pallets = _apply_flow_owner(pallets, session)
    for pallet in pallets:
        if not isinstance(pallet, dict):
            continue
        pallet_boxes = [
            str(code).strip()
            for code in (pallet.get("boxes") or [])
            if str(code or "").strip()
        ]
        pallet["boxes"] = [code for code in pallet_boxes if code in box_codes]
    active_box = str(source.get("activeBox") or "").strip()
    active_pallet = str(source.get("activePallet") or "").strip()
    if active_box and active_box not in box_codes:
        active_box = ""
    pallet_codes = {str(pallet.get("code") or "").strip() for pallet in pallets if isinstance(pallet, dict)}
    if active_pallet and active_pallet not in pallet_codes:
        active_pallet = ""
    boxes, pallets = _dedupe_pallet_box_links(boxes, pallets)
    return {
        "boxes": boxes,
        "pallets": pallets,
        "activeBox": active_box,
        "activePallet": active_pallet,
    }


def _dedupe_pallet_box_links(boxes_data: list, pallets_data: list) -> tuple[list, list]:
    valid_box_codes = {
        str((box or {}).get("code") or "").strip()
        for box in (boxes_data or [])
        if isinstance(box, dict) and str((box or {}).get("code") or "").strip()
    }
    normalized_pallets = []
    for raw in pallets_data or []:
        if not isinstance(raw, dict):
            continue
        pallet = dict(raw)
        pallet_code = str(pallet.get("code") or "").strip()
        if not pallet_code:
            continue
        seen = set()
        cleaned_boxes = []
        for box_code in pallet.get("boxes") or []:
            code = str(box_code or "").strip()
            if not code or code not in valid_box_codes or code in seen:
                continue
            seen.add(code)
            cleaned_boxes.append(code)
        pallet["boxes"] = cleaned_boxes
        normalized_pallets.append(pallet)
    last_pallet_by_box = {}
    for pallet in normalized_pallets:
        pallet_code = str(pallet.get("code") or "").strip()
        for box_code in pallet.get("boxes") or []:
            last_pallet_by_box[str(box_code)] = pallet_code
    for pallet in normalized_pallets:
        pallet_code = str(pallet.get("code") or "").strip()
        pallet["boxes"] = [
            str(code)
            for code in (pallet.get("boxes") or [])
            if last_pallet_by_box.get(str(code)) == pallet_code
        ]
    return boxes_data, normalized_pallets


def _merge_flow_values_by_code(base_values, override_values):
    merged_map = {}
    for raw in base_values or []:
        if not isinstance(raw, dict):
            continue
        code = str(raw.get("code") or "").strip()
        if not code:
            continue
        item = dict(raw)
        item["code"] = code
        merged_map[code] = item
    for raw in override_values or []:
        if not isinstance(raw, dict):
            continue
        code = str(raw.get("code") or "").strip()
        if not code:
            continue
        item = dict(raw)
        item["code"] = code
        merged_map[code] = item
    return list(merged_map.values())


def _box_to_pallet_map(pallets_data: list[dict]) -> dict[str, str]:
    mapping = {}
    for pallet in pallets_data or []:
        if not isinstance(pallet, dict):
            continue
        pallet_code = str(pallet.get("code") or "").strip()
        if not pallet_code:
            continue
        for raw_box_code in pallet.get("boxes") or []:
            box_code = str(raw_box_code or "").strip()
            if not box_code:
                continue
            mapping[box_code] = pallet_code
    return mapping


def _has_box_reassignment_between_pallets(
    base_pallets: list[dict],
    new_pallets: list[dict],
    tracked_box_codes: set[str] | None = None,
) -> bool:
    base_map = _box_to_pallet_map(base_pallets or [])
    if not base_map:
        return False
    if tracked_box_codes:
        base_map = {
            str(box_code): str(pallet_code)
            for box_code, pallet_code in base_map.items()
            if str(box_code) in tracked_box_codes
        }
        if not base_map:
            return False
    new_map = _box_to_pallet_map(new_pallets or [])
    for box_code, base_pallet_code in base_map.items():
        new_pallet_code = str(new_map.get(box_code) or "").strip()
        if not new_pallet_code:
            continue
        if new_pallet_code != str(base_pallet_code or "").strip():
            return True
    return False


def _merge_flow_sessions(sessions: list[ProcessingFlowSession]) -> tuple[list, list]:
    def _session_sort_key(session: ProcessingFlowSession) -> float:
        ts = getattr(session, "updated_at", None) or getattr(session, "last_seen", None)
        if not ts:
            return 0.0
        try:
            return float(ts.timestamp())
        except Exception:
            return 0.0

    sessions_sorted = sorted(sessions, key=_session_sort_key)
    boxes_raw: list = []
    pallets_raw: list = []
    for session in sessions_sorted:
        state = _sanitize_flow_state_for_session(session.flow_state, session)
        if isinstance(state.get("boxes"), list):
            boxes_raw.extend(_apply_flow_owner(state.get("boxes") or [], session))
        if isinstance(state.get("pallets"), list):
            pallets_raw.extend(_apply_flow_owner(state.get("pallets") or [], session))
    boxes_map: dict[str, dict] = {}
    for box in boxes_raw:
        if not isinstance(box, dict):
            continue
        code = str(box.get("code") or "").strip()
        if not code:
            continue
        boxes_map[code] = box
    pallets_map: dict[str, dict] = {}
    for pallet in pallets_raw:
        if not isinstance(pallet, dict):
            continue
        code = str(pallet.get("code") or "").strip()
        if not code:
            continue
        pallets_map[code] = pallet
    boxes = list(boxes_map.values())
    pallets = list(pallets_map.values())
    box_codes = {box.get("code") for box in boxes}
    for pallet in pallets:
        if not isinstance(pallet, dict):
            continue
        pallet_boxes = [
            str(code).strip()
            for code in (pallet.get("boxes") or [])
            if str(code or "").strip()
        ]
        pallet["boxes"] = [code for code in pallet_boxes if code in box_codes]
    boxes, pallets = _dedupe_pallet_box_links(boxes, pallets)

    def _box_owner_key(box: dict) -> str:
        return (
            str(box.get("owner_user_id") or "").strip()
            or str(box.get("owner_user_label") or "").strip()
            or str(box.get("owner_agent_id") or "").strip()
        )

    def _owner_initials(label: str) -> str:
        parts = re.findall(r"[A-Za-zА-Яа-яЁё]+", str(label or ""))
        if not parts:
            return ""
        first = parts[0][0].upper() if parts[0] else ""
        second = parts[1][0].upper() if len(parts) > 1 and parts[1] else ""
        return f"{first}{second}"

    def _fixed_label_number(value) -> int:
        text = str(value or "").strip()
        if not text:
            return 0
        match = re.match(r"^(\d+)", text)
        if not match:
            return 0
        try:
            return int(match.group(1))
        except Exception:
            return 0

    sealed_by_owner: dict[str, list[dict]] = {}
    for box in boxes:
        if not isinstance(box, dict):
            continue
        if not box.get("sealed"):
            continue
        owner_key = _box_owner_key(box)
        if not owner_key:
            continue
        sealed_by_owner.setdefault(owner_key, []).append(box)

    for owner_key, owner_boxes in sealed_by_owner.items():
        owner_label = ""
        for box in owner_boxes:
            owner_label = str(box.get("owner_user_label") or box.get("owner_agent_id") or "").strip()
            if owner_label:
                break
        initials = _owner_initials(owner_label)
        used_numbers: set[int] = set()
        for box in owner_boxes:
            num = _fixed_label_number(box.get("fixed_label"))
            if num > 0:
                used_numbers.add(num)
        for box in sorted(owner_boxes, key=lambda item: str(item.get("code") or "")):
            current = str(box.get("fixed_label") or "").strip()
            if current:
                continue
            next_num = 1
            while next_num in used_numbers:
                next_num += 1
            used_numbers.add(next_num)
            box["fixed_label"] = f"{next_num}{initials}" if initials else str(next_num)

    def _owner_key(pallet: dict) -> str:
        return (
            str(pallet.get("owner_user_id") or "").strip()
            or str(pallet.get("owner_user_label") or "").strip()
            or str(pallet.get("owner_agent_id") or "").strip()
        )

    def _pallet_item_total(pallet: dict) -> int:
        total = 0
        for item in pallet.get("items") or []:
            if not isinstance(item, dict):
                continue
            total += _parse_qty_value(item.get("qty")) or 0
        for code in pallet.get("boxes") or []:
            box = boxes_map.get(str(code).strip())
            if not isinstance(box, dict):
                continue
            for item in box.get("items") or []:
                if not isinstance(item, dict):
                    continue
                total += _parse_qty_value(item.get("qty")) or 0
        return total

    open_by_owner: dict[str, list[dict]] = {}
    for pallet in pallets:
        if not isinstance(pallet, dict):
            continue
        if pallet.get("sealed"):
            continue
        owner = _owner_key(pallet)
        if not owner:
            continue
        open_by_owner.setdefault(owner, []).append(pallet)

    for owner, owner_pallets in open_by_owner.items():
        if len(owner_pallets) <= 1:
            continue
        keep = max(owner_pallets, key=_pallet_item_total)
        for pallet in owner_pallets:
            if pallet is keep:
                continue
            pallet["sealed"] = True
            if not pallet.get("closed_by_user_id"):
                pallet["closed_by_user_id"] = pallet.get("owner_user_id") or ""
            if not pallet.get("closed_by_user_label"):
                pallet["closed_by_user_label"] = pallet.get("owner_user_label") or ""
            if not pallet.get("closed_by_agent_id"):
                pallet["closed_by_agent_id"] = pallet.get("owner_agent_id") or ""

    for pallet in pallets:
        if not isinstance(pallet, dict):
            continue
        if not pallet.get("sealed"):
            continue
        owner_label = _resolve_actor_label(
            pallet.get("owner_user_label"),
            pallet.get("owner_agent_id"),
            pallet.get("owner_user_id"),
            unknown="",
        )
        if owner_label and not str(pallet.get("owner_user_label") or "").strip():
            pallet["owner_user_label"] = owner_label
        closer_label = _resolve_actor_label(
            pallet.get("closed_by_user_label"),
            pallet.get("closed_by_agent_id"),
            pallet.get("closed_by_user_id"),
            unknown="",
        )
        if not closer_label:
            closer_label = owner_label
        if closer_label and not str(pallet.get("closed_by_user_label") or "").strip():
            pallet["closed_by_user_label"] = closer_label

    return boxes, pallets


def _build_box_label_map(boxes_data: list[dict]) -> dict[str, str]:
    def _owner_key(box: dict) -> str:
        return (
            str(box.get("owner_user_id") or "").strip()
            or str(box.get("owner_user_label") or "").strip()
            or str(box.get("owner_agent_id") or "").strip()
            or "__global__"
        )

    def _owner_initials(label: str) -> str:
        parts = re.findall(r"[A-Za-zА-Яа-яЁё]+", str(label or ""))
        if not parts:
            return ""
        first = parts[0][0].upper() if parts[0] else ""
        second = parts[1][0].upper() if len(parts) > 1 and parts[1] else ""
        return f"{first}{second}"

    def _fixed_label_number(value) -> int:
        text = str(value or "").strip()
        if not text:
            return 0
        match = re.match(r"^(\d+)", text)
        if not match:
            return 0
        try:
            return int(match.group(1))
        except Exception:
            return 0

    labels: dict[str, str] = {}
    used_numbers: dict[str, set[int]] = {}
    owner_initials: dict[str, str] = {}
    missing_by_owner: dict[str, list[str]] = {}

    for box in boxes_data or []:
        if not isinstance(box, dict):
            continue
        code = str(box.get("code") or "").strip()
        if not code:
            continue
        owner = _owner_key(box)
        owner_label = str(box.get("owner_user_label") or box.get("owner_agent_id") or "").strip()
        owner_initials.setdefault(owner, _owner_initials(owner_label))
        used_numbers.setdefault(owner, set())
        fixed_label = str(box.get("fixed_label") or "").strip()
        if fixed_label:
            labels[code] = fixed_label
            num = _fixed_label_number(fixed_label)
            if num > 0:
                used_numbers[owner].add(num)
            continue
        missing_by_owner.setdefault(owner, []).append(code)

    for owner, codes in missing_by_owner.items():
        next_num = 1
        for code in sorted(codes):
            while next_num in used_numbers[owner]:
                next_num += 1
            used_numbers[owner].add(next_num)
            initials = owner_initials.get(owner) or ""
            labels[code] = f"{next_num}{initials}" if initials else str(next_num)

    return labels


def _state_has_open_box_with_items(flow_state: dict) -> bool:
    if not isinstance(flow_state, dict):
        return False
    boxes = flow_state.get("boxes") or []
    if not isinstance(boxes, list):
        return False
    for box in boxes:
        if not isinstance(box, dict):
            continue
        if box.get("sealed"):
            continue
        items = box.get("items") or []
        if not isinstance(items, list):
            continue
        total = sum(_parse_qty_value(item.get("qty")) or 0 for item in items if isinstance(item, dict))
        if total > 0:
            return True
    return False


def _order_has_open_boxes(order_id: str) -> bool:
    sessions = ProcessingFlowSession.objects.filter(
        order_id=order_id,
        order_type="processing",
        status=ProcessingFlowSession.STATUS_OPEN,
    )
    for session in sessions:
        if _state_has_open_box_with_items(session.flow_state or {}):
            return True
    return False


def _extract_zone_code(location_value, fallback_value=None) -> str:
    zone_raw = ""
    if isinstance(location_value, dict):
        zone_raw = location_value.get("zone") or ""
    if not zone_raw:
        zone_raw = fallback_value or ""
    text = str(zone_raw or "").strip().upper()
    if not text:
        return ""
    if text in {"OBR", "PR", "OTG", "MR", "OS"}:
        return text
    return text


def _has_processing_warehouse_move_task(order_id: str) -> bool:
    return bool(_processing_warehouse_move_progress(order_id).get("has_any_task"))


def _processing_placement_entry(entries: list[OrderAuditEntry]):
    closed_entry = next(
        (
            entry
            for entry in reversed(entries or [])
            if (entry.payload or {}).get("act") == "placement" and (entry.payload or {}).get("flow_closed")
        ),
        None,
    )
    if closed_entry:
        return closed_entry
    return next(
        (
            entry
            for entry in reversed(entries or [])
            if (entry.payload or {}).get("act") == "placement"
        ),
        None,
    )


def _processing_result_key(item: dict, fallback_card_id: str = "") -> tuple[str, str, str, str]:
    card_key = str(item.get("card_id") or fallback_card_id or "").strip().lower()
    article_key = str(item.get("article") or "").strip().lower()
    size_key = str(item.get("size") or "").strip().lower()
    dest_key = str(item.get("destination") or "").strip().lower() or "-"
    return card_key, article_key, size_key, dest_key


def _processing_results_are_ready(payload: dict, include_shipping: bool = True) -> bool:
    expected_results = _expected_processing_results(payload)
    if not expected_results:
        return True
    has_direction_distribution = any(
        str(dest or "").strip() and str(dest or "").strip() != "-"
        for _, _, _, dest in expected_results
    )
    result_requirements = _processing_result_requirements(
        payload,
        has_direction_distribution=has_direction_distribution,
    )
    results = payload.get("processing_results") or []
    results_map = {}
    results_by_triplet = {}
    if isinstance(results, list):
        for item in results:
            if not isinstance(item, dict):
                continue
            key = _processing_result_key(item)
            results_map[key] = item
            results_by_triplet[(key[1], key[2], key[3])] = item
    required_fields = ["processed"]
    if result_requirements.get("quality"):
        required_fields.extend(["defect", "shortage"])
    if result_requirements.get("labels"):
        required_fields.extend(["labels_printed", "labels_unboxed"])
    if include_shipping and result_requirements.get("shipping"):
        required_fields.append("shipped_qty")
    if result_requirements.get("tags"):
        required_fields.append("tags_replaced")

    def _is_field_ready(saved_row: dict, field_name: str) -> bool:
        if not isinstance(saved_row, dict):
            return False
        value = saved_row.get(field_name)
        if field_name == "tags_replaced" and _parse_qty_value(value) is None:
            # For tag replacement, default business behavior is "as processed".
            value = saved_row.get("processed")
        return _parse_qty_value(value) is not None

    def _field_score(saved_row: dict) -> int:
        if not isinstance(saved_row, dict):
            return -1
        return sum(1 for field_name in required_fields if _is_field_ready(saved_row, field_name))

    def _select_saved_row(expected_key: tuple[str, str, str, str]):
        candidates: list[dict] = []
        exact = results_map.get(expected_key)
        if isinstance(exact, dict):
            candidates.append(exact)
        if expected_key[0]:
            legacy = results_map.get(("", expected_key[1], expected_key[2], expected_key[3]))
            if isinstance(legacy, dict):
                candidates.append(legacy)
        triplet_row = results_by_triplet.get((expected_key[1], expected_key[2], expected_key[3]))
        if isinstance(triplet_row, dict):
            candidates.append(triplet_row)
        if not candidates:
            return None
        return max(
            candidates,
            key=lambda row: (
                _field_score(row),
                1 if str(row.get("card_id") or "").strip() else 0,
            ),
        )

    for key in expected_results:
        saved = _select_saved_row(key)
        if not saved:
            return False
        for field in required_fields:
            if not _is_field_ready(saved, field):
                return False
    return True


def _processing_results_ready_card_ids(payload: dict, include_shipping: bool = True) -> set[str]:
    expected_results = _expected_processing_results(payload)
    if not expected_results:
        return set()
    has_direction_distribution = any(
        str(dest or "").strip() and str(dest or "").strip() != "-"
        for _, _, _, dest in expected_results
    )
    result_requirements = _processing_result_requirements(
        payload,
        has_direction_distribution=has_direction_distribution,
    )
    results = payload.get("processing_results") or []
    results_map = {}
    results_by_triplet = {}
    if isinstance(results, list):
        for item in results:
            if not isinstance(item, dict):
                continue
            key = _processing_result_key(item)
            results_map[key] = item
            results_by_triplet[(key[1], key[2], key[3])] = item
    required_fields = ["processed"]
    if result_requirements.get("quality"):
        required_fields.extend(["defect", "shortage"])
    if result_requirements.get("labels"):
        required_fields.extend(["labels_printed", "labels_unboxed"])
    if include_shipping and result_requirements.get("shipping"):
        required_fields.append("shipped_qty")
    if result_requirements.get("tags"):
        required_fields.append("tags_replaced")

    def _is_field_ready(saved_row: dict, field_name: str) -> bool:
        if not isinstance(saved_row, dict):
            return False
        value = saved_row.get(field_name)
        if field_name == "tags_replaced" and _parse_qty_value(value) is None:
            value = saved_row.get("processed")
        return _parse_qty_value(value) is not None

    def _field_score(saved_row: dict) -> int:
        if not isinstance(saved_row, dict):
            return -1
        return sum(1 for field_name in required_fields if _is_field_ready(saved_row, field_name))

    def _select_saved_row(expected_key: tuple[str, str, str, str]):
        candidates: list[dict] = []
        exact = results_map.get(expected_key)
        if isinstance(exact, dict):
            candidates.append(exact)
        if expected_key[0]:
            legacy = results_map.get(("", expected_key[1], expected_key[2], expected_key[3]))
            if isinstance(legacy, dict):
                candidates.append(legacy)
        triplet_row = results_by_triplet.get((expected_key[1], expected_key[2], expected_key[3]))
        if isinstance(triplet_row, dict):
            candidates.append(triplet_row)
        if not candidates:
            return None
        return max(
            candidates,
            key=lambda row: (
                _field_score(row),
                1 if str(row.get("card_id") or "").strip() else 0,
            ),
        )

    card_state: dict[str, bool] = {}
    card_ids_by_article_size: dict[tuple[str, str], set[str]] = {}
    cards = payload.get("cards") or []
    if isinstance(cards, list):
        for card in cards:
            if not isinstance(card, dict):
                continue
            card_key = _processing_card_id(card).strip().lower()
            if not card_key:
                continue
            base_article = str(card.get("article") or "").strip().lower()
            for row in (card.get("rows") or []):
                if not isinstance(row, dict):
                    continue
                article_key = str(row.get("article") or base_article).strip().lower()
                size_key = str(row.get("size") or "").strip().lower()
                if not article_key:
                    continue
                card_ids_by_article_size.setdefault((article_key, size_key), set()).add(card_key)
    for key in expected_results:
        card_key = str(key[0] or "").strip().lower()
        saved = _select_saved_row(key)
        row_ready = bool(saved) and all(_is_field_ready(saved, field) for field in required_fields)
        if card_key:
            card_state[card_key] = bool(card_state.get(card_key, True) and row_ready)
            continue
        saved_card_key = str((saved or {}).get("card_id") or "").strip().lower()
        if saved_card_key:
            card_state[saved_card_key] = bool(card_state.get(saved_card_key, True) and row_ready)
        article_key = str(key[1] or "").strip().lower()
        size_key = str(key[2] or "").strip().lower()
        fallback_card_ids = card_ids_by_article_size.get((article_key, size_key)) or set()
        for fallback_card_id in fallback_card_ids:
            card_state[fallback_card_id] = bool(card_state.get(fallback_card_id, True) and row_ready)
    return {card_id for card_id, is_ready in card_state.items() if is_ready}


def _processing_cards_total(payload: dict | None) -> int:
    payload = payload if isinstance(payload, dict) else {}
    cards = payload.get("cards") or []
    if isinstance(cards, list):
        total = sum(1 for card in cards if isinstance(card, dict))
        if total > 0:
            return total
    stock_rows = payload.get("stock_rows") or []
    if isinstance(stock_rows, list) and stock_rows:
        return 1
    if str(payload.get("article") or "").strip() or str(payload.get("product_name") or "").strip():
        return 1
    return 0


def _processing_finish_checks(order_id: str, payload: dict, entries: list[OrderAuditEntry]) -> dict:
    blockers: list[str] = []

    placement_entry = _processing_placement_entry(entries)
    placement_payload = placement_entry.payload or {} if placement_entry else {}
    placement_state = str((placement_payload or {}).get("act_state") or "closed").strip().lower()
    placement_closed = bool(placement_entry and placement_state == "closed")
    placement_boxes = placement_payload.get("act_boxes") or []
    placement_pallets = placement_payload.get("act_pallets") or []
    if not isinstance(placement_boxes, list):
        placement_boxes = []
    if not isinstance(placement_pallets, list):
        placement_pallets = []
    has_boxes = any(
        str((box or {}).get("code") or "").strip()
        for box in placement_boxes
        if isinstance(box, dict)
    )
    has_pallets = any(
        str((pallet or {}).get("code") or "").strip()
        for pallet in placement_pallets
        if isinstance(pallet, dict)
    )
    warehouse_move_progress = _processing_warehouse_move_progress(
        str(order_id),
        placement_pallets,
        agency=entries[-1].agency if entries else None,
    )
    warehouse_move_created = bool(warehouse_move_progress.get("has_any_task"))
    warehouse_move_completed = bool(warehouse_move_progress.get("all_done"))
    warehouse_not_created = int(warehouse_move_progress.get("not_created_count") or 0)

    if _order_has_open_boxes(order_id):
        blockers.append("Закройте все открытые короба с товаром.")
    if not _processing_results_are_ready(payload):
        blockers.append("Заполните результаты обработки по всем строкам.")
    if not placement_closed:
        blockers.append("Закройте раскоробовку (акт размещения).")
    if placement_closed and (not has_boxes or not has_pallets):
        blockers.append("Разместите товар в короба и палеты.")
    if warehouse_not_created > 0:
        blockers.append("Сформируйте задание ричтракеру на перемещение палет в склад.")
    elif not warehouse_move_completed:
        blockers.append("Дождитесь выполнения задания ричтракером (доставка палет в склад).")

    return {
        "blockers": blockers,
        "placement_entry": placement_entry,
        "placement_payload": placement_payload,
        "placement_closed": placement_closed,
        "has_boxes": has_boxes,
        "has_pallets": has_pallets,
        "warehouse_move_created": warehouse_move_created,
        "warehouse_move_completed": warehouse_move_completed,
        "warehouse_move_progress": warehouse_move_progress,
    }


def _next_stock_move_number() -> str:
    order_ids = (
        OrderAuditEntry.objects.filter(order_type="stock_move")
        .values_list("order_id", flat=True)
        .distinct()
    )
    max_number = 0
    for order_id in order_ids:
        candidate = str(order_id).strip()
        if not candidate.isdigit():
            continue
        number = int(candidate)
        if number > max_number:
            max_number = number
    next_number = max_number + 1
    while OrderAuditEntry.objects.filter(order_type="stock_move", order_id=str(next_number)).exists():
        next_number += 1
    return str(next_number)


def _normalize_move_zone(raw: str | None) -> str:
    text = str(raw or "").strip()
    if not text:
        return ""
    zone = normalize_zone_code(text)
    return zone if zone in {"PR", "OTG", "MR", "OS", "OBR"} else ""


def _normalize_move_location(raw_location, fallback_payload=None) -> dict:
    return normalize_putaway_location(raw_location, fallback_payload)


def _move_location_label(location: dict | None) -> str:
    return putaway_location_label(location)


def _suggest_processing_warehouse_destinations(
    pallets,
    *,
    exclude_order_type: str,
    exclude_order_id: str,
    agency_id: int | None = None,
) -> dict[str, dict]:
    return suggest_putaway_destinations(
        pallets,
        exclude_order_type=exclude_order_type,
        exclude_order_id=exclude_order_id,
        agency_id=agency_id,
        row_sections=_OS_ROW_SECTIONS,
        tiers=_OS_TIERS,
        cells_per_tier=_OS_CELLS_PER_TIER,
    )


def _latest_stock_moves_by_pallet_for_processing(order_id: str) -> dict[str, dict]:
    target_id = str(order_id or "").strip()
    latest: dict[str, dict] = {}
    if not target_id:
        return latest
    entries = OrderAuditEntry.objects.filter(order_type="stock_move").order_by("-created_at")
    for entry in entries:
        payload = entry.payload or {}
        if not isinstance(payload, dict):
            continue
        processing_order_id = str(payload.get("processing_order_id") or "").strip()
        if processing_order_id != target_id:
            continue
        pallet_code = str(payload.get("pallet_code") or "").strip()
        if not pallet_code or pallet_code in latest:
            continue
        to_location = payload.get("to_location") or {}
        latest[pallet_code] = {
            "status": str(payload.get("status") or "").strip().lower(),
            "to_zone": _normalize_move_zone((to_location or {}).get("zone") or payload.get("to_zone")),
            "order_id": str(entry.order_id or "").strip(),
        }
    return latest


def _warehouse_processing_moves_by_pallet(
    order_id: str,
    *,
    pallet_codes: list[str] | None = None,
    agency=None,
) -> dict[str, dict]:
    target_id = str(order_id or "").strip()
    codes = {
        str(code or "").strip()
        for code in (pallet_codes or [])
        if str(code or "").strip()
    }
    latest: dict[str, dict] = {}
    if not target_id or not agency or not codes:
        return latest
    snapshots = (
        WarehouseStockSnapshot.objects.select_related("parent_container", "location")
        .filter(
            agency=agency,
            is_archived=False,
        )
        .filter(Q(container_code__in=codes) | Q(parent_container__container_code__in=codes))
        .order_by("id")
    )
    for snapshot in snapshots:
        pallet_code = str(
            getattr(getattr(snapshot, "parent_container", None), "container_code", "") or snapshot.container_code or ""
        ).strip()
        if not pallet_code or pallet_code not in codes:
            continue
        zone_code = str(snapshot.zone_code or getattr(snapshot.location, "zone_code", "") or "").strip().upper()
        if not zone_code or zone_code == "OBR":
            continue
        latest[pallet_code] = {
            "status": "done",
            "status_label": "Перемещено на склад",
            "to_zone": zone_code,
            "to_label": _move_location_label(
                {
                    "zone": zone_code,
                    "row": getattr(snapshot.location, "row_no", 0) or "",
                    "section": getattr(snapshot.location, "section_no", 0) or "",
                    "tier": getattr(snapshot.location, "tier_no", 0) or "",
                    "cell": getattr(snapshot.location, "cell_no", 0) or "",
                }
            ),
            "to_row": getattr(snapshot.location, "row_no", 0) or 0,
            "to_section": getattr(snapshot.location, "section_no", 0) or 0,
            "to_tier": getattr(snapshot.location, "tier_no", 0) or 0,
            "to_cell": getattr(snapshot.location, "cell_no", 0) or 0,
            "from_zone": "OBR",
            "order_id": target_id,
            "source": "warehouse",
        }
    return latest


def _latest_warehouse_moves_by_pallet_for_processing(
    order_id: str,
    *,
    pallet_codes: list[str] | None = None,
    agency=None,
) -> dict[str, dict]:
    target_id = str(order_id or "").strip()
    latest: dict[str, dict] = {}
    if not target_id:
        return latest
    entries = OrderAuditEntry.objects.filter(order_type="stock_move").order_by("-created_at")
    for entry in entries:
        payload = entry.payload or {}
        if not isinstance(payload, dict):
            continue
        processing_order_id = str(payload.get("processing_order_id") or "").strip()
        if processing_order_id != target_id:
            continue
        pallet_code = str(payload.get("pallet_code") or "").strip()
        if not pallet_code or pallet_code in latest:
            continue
        to_zone = _extract_zone_code(payload.get("to_location"), payload.get("to_zone"))
        from_zone = _extract_zone_code(payload.get("from_location"), payload.get("from_zone"))
        if not to_zone or to_zone == "OBR":
            continue
        if from_zone and from_zone != "OBR":
            continue
        to_location = payload.get("to_location") or {}
        latest[pallet_code] = {
            "status": str(payload.get("status") or payload.get("submit_action") or "").strip().lower(),
            "status_label": str(payload.get("status_label") or "").strip(),
            "to_zone": to_zone,
            "to_label": _move_location_label(payload.get("to_location") or {}),
            "to_row": _parse_int_value((to_location or {}).get("row")),
            "to_section": _parse_int_value((to_location or {}).get("section")),
            "to_tier": _parse_int_value((to_location or {}).get("tier")),
            "to_cell": _parse_int_value((to_location or {}).get("cell")),
            "from_zone": from_zone,
            "order_id": str(entry.order_id or "").strip(),
            "source": "audit",
        }
    for pallet_code, move_data in _warehouse_processing_moves_by_pallet(
        target_id,
        pallet_codes=pallet_codes,
        agency=agency,
    ).items():
        latest.setdefault(pallet_code, move_data)
    return latest


def _processing_warehouse_move_progress(order_id: str, placement_pallets=None, agency=None) -> dict:
    target_id = str(order_id or "").strip()
    pallets = placement_pallets if isinstance(placement_pallets, list) else []
    pallet_codes: list[str] = []
    zone_by_pallet: dict[str, str] = {}
    for pallet in pallets:
        if not isinstance(pallet, dict):
            continue
        code = str(pallet.get("code") or "").strip()
        if not code or code in zone_by_pallet:
            continue
        pallet_codes.append(code)
        zone_by_pallet[code] = _extract_zone_code(pallet.get("location"), pallet.get("zone"))
    latest_moves = _latest_warehouse_moves_by_pallet_for_processing(
        target_id,
        pallet_codes=pallet_codes,
        agency=agency,
    )
    created_count = 0
    in_progress_count = 0
    done_count = 0
    canceled_count = 0
    other_count = 0
    implicit_done_count = 0
    not_created_count = 0
    implicit_done_codes: set[str] = set()
    if pallet_codes:
        for pallet_code in pallet_codes:
            latest_move = latest_moves.get(pallet_code) or {}
            status = str(latest_move.get("status") or "").strip().lower()
            move_source = str(latest_move.get("source") or "").strip().lower()
            if status == "done":
                done_count += 1
                if move_source == "warehouse":
                    implicit_done_count += 1
                    implicit_done_codes.add(pallet_code)
                continue
            if status in {"canceled", "cancelled"}:
                canceled_count += 1
                not_created_count += 1
                continue
            if status == "in_progress":
                in_progress_count += 1
                continue
            if status == "created":
                created_count += 1
                continue
            if status:
                other_count += 1
                continue
            zone = zone_by_pallet.get(pallet_code) or ""
            if zone and zone != "OBR":
                done_count += 1
                implicit_done_count += 1
                implicit_done_codes.add(pallet_code)
                continue
            not_created_count += 1
    else:
        for latest_move in latest_moves.values():
            status = str((latest_move or {}).get("status") or "").strip().lower()
            if status == "done":
                done_count += 1
            elif status in {"canceled", "cancelled"}:
                canceled_count += 1
            elif status == "in_progress":
                in_progress_count += 1
            elif status == "created":
                created_count += 1
            elif status:
                other_count += 1
    total_pallets = len(pallet_codes)
    active_count = created_count + in_progress_count
    has_any_task = bool(active_count or done_count or other_count)
    all_done = bool(total_pallets) and done_count >= total_pallets
    return {
        "total_pallets": total_pallets,
        "created_count": created_count,
        "in_progress_count": in_progress_count,
        "active_count": active_count,
        "done_count": done_count,
        "canceled_count": canceled_count,
        "other_count": other_count,
        "implicit_done_count": implicit_done_count,
        "not_created_count": not_created_count,
        "pending_count": max(total_pallets - done_count, 0),
        "has_any_task": has_any_task,
        "all_done": all_done,
        "moves_by_pallet": latest_moves,
        "implicit_done_codes": implicit_done_codes,
    }


def _create_processing_warehouse_moves(
    order_id: str,
    entries: list[OrderAuditEntry],
    request,
    destinations_by_pallet: dict[str, dict] | None = None,
) -> tuple[int, int, int]:
    order_key = str(order_id or "").strip()
    if not order_key or not entries:
        return 0, 0, 0
    finish_checks = _processing_finish_checks(order_key, dict((entries[-1].payload or {})), entries)
    placement_payload = finish_checks.get("placement_payload") or {}
    placement_pallets = placement_payload.get("act_pallets") or []
    if not isinstance(placement_pallets, list):
        placement_pallets = []
    normalized_destinations: dict[str, dict] = {}
    if isinstance(destinations_by_pallet, dict):
        for raw_pallet_code, raw_destination in destinations_by_pallet.items():
            pallet_code = str(raw_pallet_code or "").strip()
            if not pallet_code:
                continue
            normalized_destinations[pallet_code] = _normalize_move_location(raw_destination)
    latest = entries[-1]
    placement_pallet_codes = [
        str((pallet or {}).get("code") or "").strip()
        for pallet in placement_pallets
        if isinstance(pallet, dict) and str((pallet or {}).get("code") or "").strip()
    ]
    latest_by_pallet = _latest_warehouse_moves_by_pallet_for_processing(
        order_key,
        pallet_codes=placement_pallet_codes,
        agency=latest.agency if latest else None,
    )
    creator = get_employee_for_user(request.user)
    creator_name = ""
    if creator and creator.full_name:
        creator_name = creator.full_name
    elif request.user and request.user.is_authenticated:
        creator_name = request.user.get_full_name().strip() or request.user.username or str(request.user)
    creator_role = get_request_role(request)
    created = 0
    skipped_existing = 0
    skipped_missing_destination = 0
    task_specs_by_destination: dict[tuple[str, int, int, int, int], list[dict]] = {}
    seen_os_destinations: set[tuple[int, int, int, int]] = set()
    for pallet in placement_pallets:
        if not isinstance(pallet, dict):
            continue
        pallet_code = str(pallet.get("code") or "").strip()
        if not pallet_code:
            continue
        if normalized_destinations and pallet_code not in normalized_destinations:
            continue
        latest_move = latest_by_pallet.get(pallet_code) or {}
        latest_status = str(latest_move.get("status") or "").strip().lower()
        latest_to_zone = _normalize_move_zone(latest_move.get("to_zone"))
        if latest_to_zone and latest_to_zone != "OBR" and latest_status in {"created", "in_progress", "done"}:
            skipped_existing += 1
            continue
        from_location = _normalize_move_location(pallet.get("location"), pallet)
        if _normalize_move_zone(from_location.get("zone")) != "OBR":
            from_location = {"zone": "OBR", "row": "", "section": "", "tier": "", "cell": ""}
        to_location = normalized_destinations.get(
            pallet_code,
            {"zone": "PR", "row": "", "section": "", "tier": "", "cell": ""},
        )
        to_zone = _normalize_move_zone(to_location.get("zone"))
        if to_zone == "PR":
            skipped_missing_destination += 1
            continue
        if to_zone == "MR" and not _parse_int_value(to_location.get("row")):
            skipped_missing_destination += 1
            continue
        if to_zone == "OS":
            row = _parse_int_value(to_location.get("row"))
            section = _parse_int_value(to_location.get("section"))
            tier = _parse_int_value(to_location.get("tier"))
            cell = _parse_int_value(to_location.get("cell"))
            if not (row and section and tier and cell):
                skipped_missing_destination += 1
                continue
            os_key = (row, section, tier, cell)
            if os_key in seen_os_destinations:
                skipped_missing_destination += 1
                continue
            seen_os_destinations.add(os_key)
        move_payload = {
            "status": "created",
            "status_label": "Ожидает перевозки",
            "pallet_code": pallet_code,
            "from_location": from_location,
            "to_location": to_location,
            "from_label": _move_location_label(from_location),
            "to_label": _move_location_label(to_location),
            "receiving_order_id": order_key,
            "requested_by_name": creator_name,
            "requested_by_role": creator_role,
            "pick_mode": "full",
            "requested_qty": "",
            "requested_sku": "",
            "requested_barcodes": [],
            "requested_goods_type": "",
            "available_qty": "",
            "processing_order_id": order_key,
        }
        destination_key = (
            str(to_location.get("zone") or "PR"),
            _parse_int_value(to_location.get("row")),
            _parse_int_value(to_location.get("section")),
            _parse_int_value(to_location.get("tier")),
            _parse_int_value(to_location.get("cell")),
        )
        task_specs_by_destination.setdefault(destination_key, []).append(
            {
                "description": f"Задание на перемещение паллеты {pallet_code} в склад",
                "payload": move_payload,
            }
        )
    for destination_key, task_specs in task_specs_by_destination.items():
        if not task_specs:
            continue
        destination = {
            "zone": destination_key[0],
            "row": destination_key[1] or "",
            "section": destination_key[2] or "",
            "tier": destination_key[3] or "",
            "cell": destination_key[4] or "",
        }
        _move_request, move_ids = create_batch_move_tasks(
            context_type="processing",
            context_id=order_key,
            agency=latest.agency,
            user=request.user if request.user.is_authenticated else None,
            requested_by_name=creator_name,
            requested_by_role=creator_role or "",
            destination=destination,
            comment=f"Автозапрос перемещения паллет по обработке #{order_key}",
            task_specs=task_specs,
        )
        created += len(move_ids)
    return created, skipped_existing, skipped_missing_destination


def _cancel_processing_warehouse_moves(
    order_id: str,
    entries: list[OrderAuditEntry],
    request,
) -> tuple[int, int]:
    order_key = str(order_id or "").strip()
    if not order_key or not entries:
        return 0, 0
    move_entries = (
        OrderAuditEntry.objects.filter(order_type="stock_move")
        .select_related("agency")
        .order_by("order_id", "created_at")
    )
    latest_by_move_order: dict[str, OrderAuditEntry] = {}
    for move_entry in move_entries:
        payload = move_entry.payload or {}
        processing_order_id = str(payload.get("processing_order_id") or "").strip()
        if processing_order_id != order_key:
            continue
        latest_by_move_order[str(move_entry.order_id)] = move_entry
    if not latest_by_move_order:
        return 0, 0
    actor = get_employee_for_user(request.user)
    if actor and actor.full_name:
        actor_name = actor.full_name
    elif request.user and request.user.is_authenticated:
        actor_name = request.user.get_full_name().strip() or request.user.username or str(request.user)
    else:
        actor_name = "Сотрудник"
    actor_role = get_request_role(request)
    latest = entries[-1]
    canceled = 0
    skipped = 0
    canceled_label = "Отменено руководителем обработки"
    for move_order_id, move_entry in latest_by_move_order.items():
        move_payload = dict(move_entry.payload or {})
        pallet_code = str(move_payload.get("pallet_code") or "").strip() or "-"
        move_status = str(
            move_payload.get("status") or move_payload.get("submit_action") or ""
        ).strip().lower()
        move_processing_order_id = str(move_payload.get("processing_order_id") or "").strip()
        if move_processing_order_id != order_key or move_status not in {"created", "in_progress"}:
            skipped += 1
            continue
        move_payload["status"] = "canceled"
        move_payload["status_label"] = canceled_label
        move_payload["canceled_at"] = timezone.localtime().isoformat()
        move_payload["canceled_by_name"] = actor_name
        move_payload["canceled_by_role"] = actor_role
        log_order_action(
            "status",
            order_id=move_order_id,
            order_type="stock_move",
            user=request.user if request.user.is_authenticated else None,
            agency=move_entry.agency if move_entry.agency else latest.agency,
            description=f"Задание {move_order_id} отменено: {pallet_code}",
            payload=move_payload,
        )
        log_stock_move(
            "update",
            user=request.user if request.user.is_authenticated else None,
            agency=move_entry.agency if move_entry.agency else latest.agency,
            description=f"Отменено задание ричтракеру {move_order_id} (паллета {pallet_code})",
            snapshot={
                "move_id": move_order_id,
                "pallet_code": pallet_code,
                "from_location": move_payload.get("from_location"),
                "to_location": move_payload.get("to_location"),
                "from_label": move_payload.get("from_label"),
                "to_label": move_payload.get("to_label"),
                "receiving_order_id": move_payload.get("receiving_order_id"),
                "status": "canceled",
                "canceled_by": actor_name,
            },
        )
        sync_task_status_by_legacy_order_id(
            move_order_id,
            status="canceled",
            assigned_to_name=actor_name,
        )
        canceled += 1
    return canceled, skipped


def _parse_processing_warehouse_destinations(raw) -> tuple[dict[str, dict] | None, str]:
    return parse_putaway_destinations(
        raw,
        allowed_zones={"PR", "MR", "OS"},
        allowed_zones_label="PR, MR или OS",
    )


def _marking_required_by_barcode(payload: dict) -> tuple[dict[str, int], int, int]:
    required: dict[str, int] = {}
    missing_barcodes = 0
    total_required = 0
    for item in extract_processing_items(payload):
        qty = item.get("qty") or 0
        if qty <= 0:
            continue
        barcode = str(item.get("barcode") or "").strip()
        if not barcode:
            missing_barcodes += qty
            continue
        total_required += qty
        required[barcode] = required.get(barcode, 0) + qty
    return required, total_required, missing_barcodes


def _marking_available_by_barcode(agency: Agency | None, order_id: str | None) -> dict[str, int]:
    if not agency:
        return {}
    qs = (
        MarkingCode.objects.filter(agency=agency, order_type="processing", used_at__isnull=True)
        .exclude(barcode="")
    )
    if order_id:
        qs = qs.filter(Q(order_id__isnull=True) | Q(order_id="") | Q(order_id=order_id))
    else:
        qs = qs.filter(Q(order_id__isnull=True) | Q(order_id=""))
    return {row["barcode"]: row["count"] for row in qs.values("barcode").annotate(count=Count("id"))}


def _marking_free_by_barcode(agency: Agency | None) -> dict[str, int]:
    if not agency:
        return {}
    qs = (
        MarkingCode.objects.filter(
            agency=agency,
            order_type="processing",
            used_at__isnull=True,
        )
        .exclude(barcode="")
        .filter(Q(order_id__isnull=True) | Q(order_id=""))
    )
    return {row["barcode"]: row["count"] for row in qs.values("barcode").annotate(count=Count("id"))}


def _reserve_marking_codes(
    agency: Agency | None,
    order_id: str | None,
    required_map: dict[str, int],
) -> tuple[bool, str | None]:
    if not agency or not order_id or not required_map:
        return True, None
    order_id = str(order_id).strip()
    if not order_id:
        return True, None
    with transaction.atomic():
        reserved_qs = MarkingCode.objects.filter(
            agency=agency,
            order_type="processing",
            order_id=order_id,
            used_at__isnull=True,
        )
        reserved_map = {
            row["barcode"]: row["count"]
            for row in reserved_qs.values("barcode").annotate(count=Count("id"))
        }
        for barcode, required_qty in required_map.items():
            reserved_qty = reserved_map.get(barcode, 0)
            if reserved_qty <= required_qty:
                continue
            release_qty = reserved_qty - required_qty
            release_ids = list(
                MarkingCode.objects.select_for_update()
                .filter(
                    agency=agency,
                    order_type="processing",
                    order_id=order_id,
                    used_at__isnull=True,
                    barcode=barcode,
                )
                .order_by("-created_at")
                .values_list("id", flat=True)[:release_qty]
            )
            if release_ids:
                MarkingCode.objects.filter(id__in=release_ids).update(order_id="")

        for barcode, required_qty in required_map.items():
            reserved_qty = reserved_map.get(barcode, 0)
            needed = max(required_qty - reserved_qty, 0)
            if not needed:
                continue
            free_ids = list(
                MarkingCode.objects.select_for_update()
                .filter(
                    agency=agency,
                    order_type="processing",
                    used_at__isnull=True,
                    barcode=barcode,
                )
                .filter(Q(order_id__isnull=True) | Q(order_id=""))
                .values_list("id", flat=True)[:needed]
            )
            if len(free_ids) < needed:
                return False, "Недостаточно свободных ЧЗ для брони."
            MarkingCode.objects.filter(id__in=free_ids).update(order_id=order_id)
    return True, None


def _normalize_marking_import_order_id(order_id: str | None) -> str:
    value = str(order_id or "").strip()
    if value.startswith("draft-"):
        return ""
    return value


def _normalize_excel_cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, int):
        return str(value)
    return str(value).strip()


def _normalize_marking_code(value: str) -> str:
    if not value:
        return ""
    # Excel encodes ASCII 29 (GS) as _x001D_ in XML exports.
    text = str(value)
    text = re.sub(r"_x001d_", "\x1d", text, flags=re.IGNORECASE)
    text = re.sub(r"<\s*gs\s*>", "\x1d", text, flags=re.IGNORECASE)
    text = text.replace("\\u001d", "\x1d").replace("\\x1d", "\x1d").replace("\\x001d", "\x1d")
    return text.strip()


def _import_marking_codes(file, payload: dict, order_id: str, agency: Agency, user):
    try:
        workbook = load_workbook(file, read_only=True, data_only=True)
    except Exception:
        return False, {"error": "Не удалось прочитать .xlsx файл."}
    sheet = workbook.active

    rows = []
    barcodes = set()
    invalid_rows = 0
    for idx, row in enumerate(sheet.iter_rows(values_only=True), start=1):
        if idx == 1:
            continue
        barcode = _normalize_excel_cell(row[0]) if row and len(row) > 0 else ""
        code = _normalize_excel_cell(row[1]) if row and len(row) > 1 else ""
        code = _normalize_marking_code(code)
        if not barcode or not code:
            if barcode or code:
                invalid_rows += 1
            continue
        rows.append((barcode, code))
        barcodes.add(barcode)

    if not rows:
        return False, {"error": "В файле нет данных для импорта."}

    barcode_qs = SKUBarcode.objects.select_related("sku").filter(value__in=barcodes)
    barcode_map = {item.value: item for item in barcode_qs}
    existing_codes = set(
        MarkingCode.objects.filter(code__in=[code for _, code in rows]).values_list("code", flat=True)
    )

    items = extract_processing_items(payload)
    allowed_pairs = {(item["sku_code"], item["size"]) for item in items}
    allowed_barcodes = {str(item.get("barcode") or "").strip() for item in items}
    allowed_barcodes.discard("")
    bind_order_id = _normalize_marking_import_order_id(order_id)
    added = 0
    duplicates = 0
    unknown_barcodes = 0
    mismatched_barcodes = 0
    seen_codes = set()
    to_create = []

    for barcode, code in rows:
        if code in seen_codes:
            duplicates += 1
            continue
        seen_codes.add(code)
        if code in existing_codes:
            duplicates += 1
            continue
        barcode_obj = barcode_map.get(barcode)
        if not barcode_obj or not barcode_obj.sku:
            unknown_barcodes += 1
            continue
        sku_obj = barcode_obj.sku
        if agency and sku_obj.agency and sku_obj.agency_id != agency.id:
            mismatched_barcodes += 1
            continue
        sku_code = sku_obj.sku_code
        size = (barcode_obj.size or sku_obj.size or "").strip()
        order_binding = ""
        if allowed_pairs:
            if (sku_code, size) in allowed_pairs:
                order_binding = bind_order_id
            elif (sku_code, "") in allowed_pairs:
                order_binding = bind_order_id
                size = ""
            else:
                matched = [pair for pair in allowed_pairs if pair[0] == sku_code]
                if len(matched) == 1:
                    order_binding = bind_order_id
                    size = matched[0][1]
        if not order_binding and barcode in allowed_barcodes:
            order_binding = bind_order_id
        to_create.append(
            MarkingCode(
                order_type="processing",
                order_id=order_binding,
                agency=agency,
                sku=sku_obj,
                sku_code=sku_code,
                size=size,
                barcode=barcode,
                code=code,
                source="import",
                created_by=user if getattr(user, "is_authenticated", False) else None,
            )
        )

    if to_create:
        with transaction.atomic():
            MarkingCode.objects.bulk_create(to_create, batch_size=500)
        added = len(to_create)

    return True, {
        "added": added,
        "duplicates": duplicates,
        "unknown_barcodes": unknown_barcodes,
        "mismatched_barcodes": mismatched_barcodes,
        "invalid_rows": invalid_rows,
    }




def _get_print_agent_token() -> str:
    return str(getattr(settings, "PRINT_AGENT_TOKEN", "")).strip()


def _check_print_agent_token(request):
    expected = _get_print_agent_token()
    if not expected:
        return False, JsonResponse({"ok": False, "error": "PRINT_AGENT_TOKEN not set"}, status=403)
    token = (
        request.headers.get("X-Print-Token")
        or request.GET.get("token")
        or request.POST.get("token")
        or ""
    )
    if token != expected:
        return False, JsonResponse({"ok": False, "error": "Invalid token"}, status=403)
    return True, None


def _require_print_admin(request):
    if not request.user.is_authenticated:
        return False, HttpResponseForbidden("Доступ запрещен")
    role = get_request_role(request)
    if role not in {"storekeeper", "processing_head", "head_manager", "director", "admin"}:
        return False, HttpResponseForbidden("Доступ запрещен")
    return True, None


def _print_queue_counts() -> dict:
    return {
        "pending": ProcessingPrintJob.objects.filter(status=ProcessingPrintJob.STATUS_PENDING).count(),
        "printing": ProcessingPrintJob.objects.filter(status=ProcessingPrintJob.STATUS_PRINTING).count(),
        "failed": ProcessingPrintJob.objects.filter(status=ProcessingPrintJob.STATUS_FAILED).count(),
    }


def _enqueue_agent_command(command: str, payload: dict | None = None, agent_id: str = "") -> None:
    if not command:
        return
    AgentCommand.objects.create(
        agent_id=agent_id or "",
        command=command,
        payload=payload or {},
    )


def _serialize_print_job(job: ProcessingPrintJob) -> dict:
    return {
        "id": job.id,
        "order_id": job.order_id,
        "orderId": job.order_id,
        "card_id": job.card_id,
        "cardId": job.card_id,
        "article": job.article,
        "barcode": job.barcode,
        "size": job.size,
        "printer_name": job.printer_name,
        "printerName": job.printer_name,
        "label_png_base64": job.label_png_base64,
        "labelPngBase64": job.label_png_base64,
        "label_width_mm": job.label_width_mm,
        "labelWidthMm": job.label_width_mm,
        "label_height_mm": job.label_height_mm,
        "labelHeightMm": job.label_height_mm,
        "status": job.status,
        "requested_by": job.requested_by,
        "requestedBy": job.requested_by,
        "created_at": job.created_at.isoformat() if job.created_at else "",
        "createdAt": job.created_at.isoformat() if job.created_at else "",
    }


def _short_city(value: str) -> str:
    text = str(value or "").strip()
    if not text:
        return ""
    city = text.split(",")[0].strip()
    lowered = city.lower()
    for prefix in ("г.", "г ", "город "):
        if lowered.startswith(prefix):
            city = city[len(prefix):].strip()
            break
    return city or text


def _processing_params_from_payload(payload: dict) -> list[dict]:
    def _add_param(rows: list[dict], label: str, value: str | None, extras: list[str] | None = None) -> None:
        base = _non_empty_text(value)
        extra_values = [item for item in (extras or []) if _non_empty_text(item)]
        if not base and not extra_values:
            return
        if base and extra_values:
            value_text = f"{base}; " + "; ".join(extra_values)
        elif base:
            value_text = base
        else:
            value_text = "; ".join(extra_values)
        rows.append({"label": label, "value": value_text})

    insert_types = payload.get("insert_types") or []
    if isinstance(insert_types, str):
        insert_types = [insert_types] if insert_types else []
    insert_other = (payload.get("insert_other") or "").strip()
    if insert_other:
        insert_types = list(insert_types) + [insert_other]

    marking_stickers = payload.get("marking_stickers") or []
    if isinstance(marking_stickers, str):
        marking_stickers = [marking_stickers] if marking_stickers else []
    marking_sizes = _processing_marking_sizes_from_payload(payload)

    processing_params: list[dict] = []
    _add_param(processing_params, "Маркетплейс", payload.get("marketplace"))
    defect_percent = _non_empty_text(payload.get("defect_percent"))
    if defect_percent and defect_percent.isdigit():
        defect_percent = f"{defect_percent}%"
    _add_param(processing_params, "Проверка на брак", defect_percent)
    for option in PROCESSING_MARKING_LABELS:
        _add_param(processing_params, option["label"], payload.get(option["field"]))
    _add_param(processing_params, "Маркировка 58/40 (шт/чз)", payload.get("marking_5840_each_qty"))
    _add_param(processing_params, "Замена бирок", payload.get("tag_owner"))

    def _add_pack_param(label: str, base_value: str | None, prefix: str) -> None:
        extras = []
        type_value = _non_empty_text(payload.get(f"{prefix}_type"))
        size_value = _non_empty_text(payload.get(f"{prefix}_size"))
        qty_value = _non_empty_text(payload.get(f"{prefix}_qty"))
        supply_value = _non_empty_text(payload.get(f"{prefix}_supply"))
        if type_value:
            extras.append(f"Тип: {type_value}")
        if size_value:
            extras.append(f"Размер: {size_value}")
        if qty_value:
            extras.append(f"Кол-во: {qty_value}")
        if supply_value:
            extras.append(f"Закупка: {supply_value}")
        _add_param(
            processing_params,
            label,
            base_value or payload.get(f"{prefix}_needed"),
            extras=extras,
        )

    _add_pack_param("Замена пакета", payload.get("bag_replace_type"), "bag_replace")
    _add_pack_param("Упаковка в Бабл пленку", payload.get("bubble_wrap_size"), "bubble_wrap")
    _add_pack_param("Упаковка в термо пленку", payload.get("shrink_wrap_size"), "shrink_wrap")
    _add_pack_param("Замена гофрокороба", payload.get("box_replace_type"), "box_replace")
    set_qty = _non_empty_text(payload.get("set_qty"))
    _add_param(
        processing_params,
        "Сборка набора",
        set_qty and f"Кол-во: {set_qty}" or payload.get("set_build"),
    )
    insert_qty = _non_empty_text(payload.get("insert_qty"))
    insert_types_label = _format_list_value(insert_types)
    if insert_other and insert_other not in insert_types:
        insert_types_label = _format_list_value(list(insert_types) + [insert_other])
    _add_param(
        processing_params,
        "Вложение",
        insert_qty and f"Кол-во: {insert_qty}" or payload.get("insert_needed"),
        extras=[f"Типы: {insert_types_label}" if insert_types_label else ""],
    )
    direction_mode = (payload.get("direction_needed") or "").strip()
    direction_file = _non_empty_text(payload.get("direction_file"))
    direction_addresses = _parse_json_value(payload.get("direction_addresses_json"), [])
    if isinstance(direction_addresses, dict):
        direction_addresses = direction_addresses.get("directions") or direction_addresses.get("addresses") or []
    if not isinstance(direction_addresses, list):
        direction_addresses = []
    direction_addresses = [str(item).strip() for item in direction_addresses if str(item).strip()]
    direction_plan = _parse_json_value(payload.get("direction_plan_json"), {})
    if not isinstance(direction_plan, dict):
        direction_plan = {}
    if not direction_addresses:
        plan_dirs = direction_plan.get("directions") or direction_plan.get("addresses") or []
        if isinstance(plan_dirs, list):
            direction_addresses = [str(item).strip() for item in plan_dirs if str(item).strip()]
    direction_count = _parse_qty_value(payload.get("direction_count"))
    if not direction_count and direction_addresses:
        direction_count = len(direction_addresses)
    direction_value = ""
    direction_extras: list[str] = []
    if direction_file or direction_mode in {"file", "Да"}:
        direction_value = direction_file and f"Файл: {direction_file}" or "Файл"
    elif direction_mode == "set" or direction_plan.get("rows"):
        if direction_addresses:
            direction_value = _format_list_value(direction_addresses)
        else:
            direction_value = "Задано"
    elif direction_mode in {"none", "Нет", "Отсутствует"}:
        direction_value = "Отсутствует"
    elif direction_mode:
        direction_value = direction_mode
    _add_param(
        processing_params,
        "Распределение по направлениям",
        direction_value,
        extras=direction_extras,
    )
    box_forming = _non_empty_text(payload.get("box_forming"))
    if box_forming == "other":
        box_forming = _non_empty_text(payload.get("box_forming_other"))
    _add_param(processing_params, "Формирование короба", box_forming)
    _add_param(processing_params, "Прочие", payload.get("comments"))
    _add_param(processing_params, "Маркировка", _format_list_value(marking_stickers))
    _add_param(processing_params, "Размеры стикеров", _format_list_value(marking_sizes))
    _add_param(processing_params, "Информационный", payload.get("marking_info"))
    _add_param(processing_params, "Вытянуть из мешка и наклеить ЧЗ", payload.get("pull_from_bag"))
    _add_param(processing_params, "Проверка на брак (кол-во)", payload.get("defect_qty"))
    _add_param(processing_params, "Обрезание ниток (кол-во)", payload.get("trim_threads_qty"))
    _add_param(processing_params, "Скрепление скотчем (кол-во)", payload.get("tape_qty"))
    _add_param(processing_params, "Удаление бирки", payload.get("remove_tag"))
    _add_param(processing_params, "Удаление бирки (кол-во)", payload.get("remove_tag_qty"))
    _add_param(processing_params, "Скрепление бирки", payload.get("attach_tag"))
    _add_param(processing_params, "Скрепление бирки (кол-во)", payload.get("attach_tag_qty"))
    return processing_params


def _parse_qty_value(raw: str | None) -> int | None:
    if raw is None:
        return None
    text = str(raw).strip()
    if not text:
        return None
    try:
        return int(text)
    except (TypeError, ValueError):
        return None


def _is_yes_value(raw) -> bool:
    text = str(raw or "").strip().lower()
    if not text:
        return False
    if text in {"1", "true", "yes", "y", "on", "да"}:
        return True
    if text in {"0", "false", "no", "n", "off", "нет", "-"}:
        return False
    return False


def _processing_has_direction_distribution(payload: dict) -> bool:
    direction_addresses = _parse_json_value(payload.get("direction_addresses_json"), [])
    if isinstance(direction_addresses, dict):
        direction_addresses = (
            direction_addresses.get("directions")
            or direction_addresses.get("addresses")
            or []
        )
    if not isinstance(direction_addresses, list):
        direction_addresses = []
    if any(str(item or "").strip() for item in direction_addresses):
        return True

    direction_plan = _parse_json_value(payload.get("direction_plan_json"), {})
    if not isinstance(direction_plan, dict):
        return False
    plan_dirs = direction_plan.get("directions") or direction_plan.get("addresses") or []
    if isinstance(plan_dirs, list) and any(str(item or "").strip() for item in plan_dirs):
        return True
    rows = direction_plan.get("rows") or []
    if not isinstance(rows, list):
        return False
    for row in rows:
        if not isinstance(row, dict):
            continue
        quantities = row.get("quantities") or []
        if not isinstance(quantities, (list, tuple)):
            continue
        for value in quantities:
            if (_parse_qty_value(value) or 0) > 0:
                return True
    return False


def _processing_direction_flow_data(payload: dict) -> dict:
    if not isinstance(payload, dict):
        return {"directions": [], "rows": [], "totals": {}}
    direction_addresses = _parse_json_value(payload.get("direction_addresses_json"), [])
    if isinstance(direction_addresses, dict):
        direction_addresses = (
            direction_addresses.get("directions")
            or direction_addresses.get("addresses")
            or []
        )
    if not isinstance(direction_addresses, list):
        direction_addresses = []
    direction_plan = _parse_json_value(payload.get("direction_plan_json"), {})
    if not isinstance(direction_plan, dict):
        direction_plan = {}
    plan_dirs = direction_plan.get("directions") or direction_plan.get("addresses") or []
    if not isinstance(plan_dirs, list):
        plan_dirs = []
    directions: list[str] = []
    seen: set[str] = set()
    for raw in [*direction_addresses, *plan_dirs]:
        name = str(raw or "").strip()
        if not name:
            continue
        key = name.lower()
        if key in seen:
            continue
        seen.add(key)
        directions.append(name)
    rows = direction_plan.get("rows") or []
    if not isinstance(rows, list):
        rows = []
    flow_rows: list[dict] = []
    totals: dict[str, int] = {name: 0 for name in directions}
    for row in rows:
        if not isinstance(row, dict):
            continue
        sku_code = str(row.get("article") or row.get("sku_code") or row.get("sku") or "").strip()
        name = str(row.get("product_name") or row.get("name") or "").strip()
        size = str(row.get("size") or "").strip()
        if not (sku_code or name or size):
            continue
        quantities = row.get("quantities") or []
        if not isinstance(quantities, (list, tuple)):
            quantities = []
        targets: dict[str, int] = {}
        for idx, direction_name in enumerate(directions):
            qty = _parse_qty_value(quantities[idx] if idx < len(quantities) else 0) or 0
            if qty <= 0:
                continue
            targets[direction_name] = targets.get(direction_name, 0) + qty
            totals[direction_name] = totals.get(direction_name, 0) + qty
        if not targets:
            continue
        flow_rows.append(
            {
                "sku_code": sku_code,
                "name": name,
                "size": size,
                "targets": targets,
            }
        )
    return {
        "directions": directions,
        "rows": flow_rows,
        "totals": totals,
    }


def _processing_result_requirements(payload: dict, has_direction_distribution: bool | None = None) -> dict[str, bool]:
    if has_direction_distribution is None:
        has_direction_distribution = _processing_has_direction_distribution(payload)

    defect_required = bool(
        str(payload.get("defect_percent") or "").strip()
        or str(payload.get("defect_qty") or "").strip()
    )
    label_qty_map = _processing_marking_qty_by_label_key(payload)
    labels_required = any(value > 0 for value in label_qty_map.values())
    tags_required = bool(
        _non_empty_text(payload.get("tag_owner"))
        or _is_yes_value(payload.get("tag_replace_needed"))
        or _is_yes_value(payload.get("remove_tag"))
        or _is_yes_value(payload.get("attach_tag"))
        or (_parse_qty_value(payload.get("remove_tag_qty")) or 0) > 0
        or (_parse_qty_value(payload.get("attach_tag_qty")) or 0) > 0
    )
    return {
        "quality": defect_required,
        "labels": labels_required,
        "shipping": bool(has_direction_distribution),
        "tags": tags_required,
    }


def _expected_processing_results(payload: dict) -> list[tuple[str, str, str, str]]:
    fallback_keys: set[tuple[str, str, str, str]] = set()
    stock_rows = payload.get("stock_rows") or []
    if isinstance(stock_rows, list):
        for row in stock_rows:
            if not isinstance(row, dict):
                continue
            article = str(row.get("article") or row.get("sku") or "").strip().lower()
            size = str(row.get("size") or "").strip().lower()
            qty = _parse_qty_value(row.get("qty")) or 0
            if qty <= 0:
                continue
            fallback_keys.add(("", article, size, "-"))
    if fallback_keys:
        return sorted(fallback_keys)

    cards = payload.get("cards") or []
    if isinstance(cards, list):
        for card in cards:
            if not isinstance(card, dict):
                continue
            card_id = _processing_card_id(card).lower()
            base_article = str(card.get("article") or "").strip().lower()
            for row in (card.get("rows") or []):
                if not isinstance(row, dict):
                    continue
                article = str(row.get("article") or base_article).strip().lower()
                size = str(row.get("size") or "").strip().lower()
                qty = _parse_qty_value(row.get("qty")) or 0
                if qty <= 0:
                    continue
                fallback_keys.add((card_id, article, size, "-"))
    if fallback_keys:
        return sorted(fallback_keys)

    # Direction plan is used on the unboxing stage; on processing card stage
    # results are filled once per article/size without direction split.
    direction_plan = _parse_json_value(payload.get("direction_plan_json"), {})
    plan_rows = direction_plan.get("rows") if isinstance(direction_plan, dict) else []
    if not isinstance(plan_rows, list):
        plan_rows = []
    for row in plan_rows:
        if not isinstance(row, dict):
            continue
        article = str(row.get("article") or row.get("product_name") or "").strip().lower()
        size = str(row.get("size") or "").strip().lower()
        if not article:
            continue
        qty_total = 0
        quantities = row.get("quantities") or []
        if isinstance(quantities, (list, tuple)):
            for value in quantities:
                qty_total += _parse_qty_value(value) or 0
        if qty_total <= 0:
            qty_total = (
                _parse_qty_value(row.get("qty"))
                or _parse_qty_value(row.get("processing_qty"))
                or 0
            )
        if qty_total <= 0:
            continue
        fallback_keys.add(("", article, size, "-"))
    if fallback_keys:
        return sorted(fallback_keys)
    return []


def _is_draft_payload(payload: dict | None) -> bool:
    payload = payload or {}
    status_value = (payload.get("status") or payload.get("submit_action") or "").lower()
    status_label = (payload.get("status_label") or "").lower()
    return status_value == "draft" or "черновик" in status_label


def _processing_work_allowed_by_payload(payload: dict, *, placement_completed: bool) -> bool:
    status_value = (payload.get("status") or payload.get("submit_action") or "").lower()
    status_label = (payload.get("status_label") or "").lower()
    return (
        status_value in {"processing_head", "processing_in_work"}
        or ("передан" in status_label and "обработ" in status_label)
        or "взята" in status_label
        or placement_completed
    )


def _processing_card_allowed_by_payload(payload: dict) -> bool:
    status_value = (payload.get("status") or payload.get("submit_action") or "").lower()
    status_label = (payload.get("status_label") or "").lower()
    return (
        status_value in {"processing_head", "processing_in_work", "done", "completed", "closed", "finished"}
        or ("передан" in status_label and "обработ" in status_label)
        or "взята" in status_label
        or "выполн" in status_label
    )


def _processing_dispatch_state(order_id: str, agency: Agency | None, payload: dict | None = None):
    if not agency:
        return None
    return WarehouseGoodsStateResolver.resolve_for_processing_order(
        order_id=str(order_id or ""),
        agency=agency,
        payload=payload if isinstance(payload, dict) else None,
    )


def _next_order_number(order_type: str = "receiving") -> str:
    order_ids = (
        OrderAuditEntry.objects.filter(order_type=order_type)
        .values_list("order_id", flat=True)
        .distinct()
    )
    max_number = 0
    for order_id in order_ids:
        if not order_id:
            continue
        candidate = str(order_id).strip()
        if not re.fullmatch(r"\d+", candidate):
            continue
        try:
            number = int(candidate)
        except (TypeError, ValueError):
            continue
        if number > max_number:
            max_number = number
    next_number = max_number + 1
    while OrderAuditEntry.objects.filter(order_type=order_type, order_id=str(next_number)).exists():
        next_number += 1
    return str(next_number)


def _manager_due_date(submitted_at):
    cutoff = submitted_at.replace(hour=14, minute=0, second=0, microsecond=0)
    if submitted_at <= cutoff:
        return submitted_at.replace(hour=18, minute=0, second=0, microsecond=0)
    next_day = submitted_at + timedelta(days=1)
    return next_day.replace(hour=13, minute=0, second=0, microsecond=0)


def _create_processing_manager_task(order_id, agency, request, submitted_at):
    if not agency:
        return
    manager = (
        Employee.objects.filter(role="manager", is_active=True)
        .order_by("full_name")
        .first()
    )
    if not manager:
        return
    route = f"/orders/processing/{order_id}/"
    existing = Task.objects.filter(route=route, assigned_to=manager).exclude(status="done")
    if existing.exists():
        return
    description = f"Клиент: {agency.agn_name or agency.inn or agency.id}"
    Task.objects.create(
        title=f"Подтвердите заявку на обработку №{order_id}",
        description=description,
        route=route,
        assigned_to=manager,
        created_by=request.user if request.user.is_authenticated else None,
        due_date=_manager_due_date(submitted_at),
    )


def _create_processing_head_task(order_id, agency, request, submitted_at):
    if not agency:
        return
    head = (
        Employee.objects.filter(role="processing_head", is_active=True)
        .order_by("full_name")
        .first()
    )
    if not head:
        return
    route = f"/orders/processing/{order_id}/"
    existing = Task.objects.filter(route=route, assigned_to=head).exclude(status="done")
    if existing.exists():
        return
    description = f"Клиент: {agency.agn_name or agency.inn or agency.id}"
    Task.objects.create(
        title=f"Заявка на обработку №{order_id}",
        description=description,
        route=route,
        assigned_to=head,
        created_by=request.user if request.user.is_authenticated else None,
        due_date=submitted_at or timezone.localtime(),
    )


def _create_processing_discrepancy_task(order_id, agency, request, mismatch_rows: list[dict] | None = None):
    if not agency:
        return
    head = (
        Employee.objects.filter(role="processing_head", is_active=True)
        .order_by("full_name")
        .first()
    )
    if not head:
        return
    if (
        request.user.is_authenticated
        and head.user_id
        and request.user.id == head.user_id
        and get_request_role(request) == "processing_head"
    ):
        return
    route = f"/orders/processing/{order_id}/"
    title = f"Разногласие по обработке №{order_id}"
    mismatch_rows = mismatch_rows or []
    summary_parts: list[str] = []
    for row in mismatch_rows[:3]:
        sku = (row.get("sku_code") or "-").strip()
        size = (row.get("size") or "-").strip()
        expected_qty = _parse_qty_value(row.get("expected_qty")) or 0
        factual_qty = _parse_qty_value(row.get("factual_qty")) or 0
        summary_parts.append(f"{sku} ({size}): {expected_qty}->{factual_qty}")
    summary_line = "; ".join(summary_parts) if summary_parts else "Проверьте акт разногласий."
    description = (
        f"Клиент: {agency.agn_name or agency.inn or agency.id}\n"
        f"Зафиксированы расхождения после обработки.\n"
        f"{summary_line}"
    )
    existing = (
        Task.objects.filter(route=route, assigned_to=head, title=title)
        .exclude(status="done")
        .order_by("-created_at")
        .first()
    )
    if existing:
        existing.description = description
        existing.priority = "high"
        existing.due_date = timezone.localtime()
        existing.save(update_fields=["description", "priority", "due_date", "updated_at"])
        return
    Task.objects.create(
        title=title,
        description=description,
        route=route,
        assigned_to=head,
        created_by=request.user if request.user.is_authenticated else None,
        due_date=timezone.localtime(),
        priority="high",
    )


def _processing_expected_qty_from_payload(payload: dict | None) -> dict[tuple[str, str], int]:
    payload = payload or {}
    expected: dict[tuple[str, str], int] = {}
    stock_rows = payload.get("stock_rows") or []
    if isinstance(stock_rows, list):
        for row in stock_rows:
            if not isinstance(row, dict):
                continue
            sku = (row.get("article") or row.get("sku") or "").strip()
            if not sku:
                continue
            size = (row.get("size") or "").strip()
            qty_value = _parse_qty_value(row.get("qty")) or 0
            if qty_value <= 0:
                continue
            key = (sku.lower(), size.lower())
            expected[key] = expected.get(key, 0) + qty_value
    if expected:
        return expected
    cards = payload.get("cards") or []
    if isinstance(cards, list):
        for card in cards:
            if not isinstance(card, dict):
                continue
            base_article = (card.get("article") or "").strip()
            for row in (card.get("rows") or []):
                if not isinstance(row, dict):
                    continue
                sku = (row.get("article") or base_article).strip()
                if not sku:
                    continue
                size = (row.get("size") or "").strip()
                qty_value = _parse_qty_value(row.get("qty")) or 0
                if qty_value <= 0:
                    continue
                key = (sku.lower(), size.lower())
                expected[key] = expected.get(key, 0) + qty_value
    return expected


def _processing_factual_qty_from_payload(
    payload: dict | None,
) -> tuple[dict[tuple[str, str], int], dict[tuple[str, str], str]]:
    payload = payload or {}
    factual: dict[tuple[str, str], int] = {}
    names: dict[tuple[str, str], str] = {}

    def add_item(item: dict):
        if not isinstance(item, dict):
            return
        sku = (item.get("sku") or item.get("sku_code") or "").strip()
        if not sku:
            return
        size = (item.get("size") or "").strip()
        qty_value = _parse_qty_value(item.get("qty"))
        if qty_value is None:
            qty_value = _parse_qty_value(item.get("actual_qty")) or 0
        key = (sku.lower(), size.lower())
        factual[key] = factual.get(key, 0) + max(qty_value, 0)
        if key not in names:
            names[key] = (item.get("name") or "").strip()

    boxes = payload.get("act_boxes") or []
    pallets = payload.get("act_pallets") or []
    for box in boxes if isinstance(boxes, list) else []:
        for item in (box or {}).get("items") or []:
            add_item(item)
    for pallet in pallets if isinstance(pallets, list) else []:
        for item in (pallet or {}).get("items") or []:
            add_item(item)
    if not boxes and not pallets:
        for item in payload.get("act_items") or []:
            add_item(item)
    return factual, names


def _processing_discrepancy_rows_from_payload(payload: dict | None) -> list[dict]:
    expected = _processing_expected_qty_from_payload(payload)
    factual, names = _processing_factual_qty_from_payload(payload)
    keys = set(expected.keys()) | set(factual.keys())
    rows: list[dict] = []
    for key in sorted(keys):
        expected_qty = expected.get(key, 0)
        factual_qty = factual.get(key, 0)
        if expected_qty == factual_qty:
            continue
        sku_lower, size_lower = key
        rows.append(
            {
                "sku_code": sku_lower,
                "size": size_lower,
                "name": names.get(key, ""),
                "expected_qty": expected_qty,
                "factual_qty": factual_qty,
                "delta_qty": factual_qty - expected_qty,
            }
        )
    return rows


def _processing_discrepancy_rows(expected_payload: dict | None, factual_payload: dict | None = None) -> list[dict]:
    comparison_payload = dict(expected_payload or {})
    factual_payload = factual_payload if isinstance(factual_payload, dict) else {}
    for key in ("act_items", "act_boxes", "act_pallets"):
        value = factual_payload.get(key)
        if isinstance(value, list):
            comparison_payload[key] = value
    return _processing_discrepancy_rows_from_payload(comparison_payload)


def _latest_non_empty_payload_value(entries: list[OrderAuditEntry], key: str):
    for entry in reversed(entries or []):
        payload = entry.payload or {}
        value = payload.get(key)
        if value not in (None, "", [], {}):
            return value
    return None


def _processing_work_payload_from_entries(entries: list[OrderAuditEntry]) -> dict:
    if not entries:
        return {}
    status_entry = _current_status_entry(entries) or entries[-1]
    payload = dict((status_entry.payload or {}) if status_entry else {})
    backfill_keys = (
        "cards",
        "stock_rows",
        "processing_results",
        "processed_cards",
        "placed_cards",
        "direction_plan_json",
        "direction_addresses_json",
        "defect_percent",
        "marking_5840_qty",
        "marking_5860_qty",
        "marking_75120_qty",
        "marking_5840_each_qty",
        "marking_sizes",
        "tag_owner",
        "goods_type",
        "goods_type_label",
        "product_name",
        "product_photo_url",
    )
    for key in backfill_keys:
        if payload.get(key) not in (None, "", [], {}):
            continue
        value = _latest_non_empty_payload_value(entries, key)
        if value not in (None, "", [], {}):
            payload[key] = value
    return payload


def _processing_reserve_rows_for_order(order_id: str, agency: Agency | None) -> list[dict]:
    if not order_id or not agency:
        return []
    reserve_groups = list(
        WarehouseReserve.objects.filter(
            agency=agency,
            reserve_type=WarehouseReserve.TYPE_PROCESSING,
            context_type="processing",
            context_id=str(order_id),
        )
        .exclude(
            status__in=[
                WarehouseReserve.STATUS_RELEASED,
                WarehouseReserve.STATUS_CANCELED,
            ]
        )
        .values("sku_code", "size", "barcode", "goods_type")
        .annotate(
            qty_reserved_sum=Sum("qty_reserved"),
            qty_satisfied_sum=Sum("qty_satisfied"),
        )
    )
    rows: list[dict] = []
    for item in reserve_groups:
        outstanding_qty = max(
            int(item.get("qty_reserved_sum") or 0) - int(item.get("qty_satisfied_sum") or 0),
            0,
        )
        if outstanding_qty <= 0:
            continue
        sku_value = str(item.get("sku_code") or "").strip()
        if not sku_value:
            continue
        rows.append(
            {
                "article": sku_value,
                "sku": sku_value,
                "size": str(item.get("size") or "").strip(),
                "barcode": str(item.get("barcode") or "").strip(),
                "goods_type": str(item.get("goods_type") or "").strip(),
                "qty": outstanding_qty,
            }
        )
    return rows


def _client_agency_from_request(request):
    if not request.user.is_authenticated:
        return None
    role = get_request_role(request)
    if is_staff_role(role):
        return None
    client_id = request.GET.get("client") or request.GET.get("agency")
    if client_id:
        return Agency.objects.filter(pk=client_id, portal_user=request.user).first()
    return Agency.objects.filter(portal_user=request.user).first()


def _inventory_items_for_agency(
    agency: Agency | None,
    exclude_order_id: str | None = None,
) -> list[dict]:
    return StockAvailabilityService.inventory_items_for_agency(
        agency=agency,
        exclude_processing_order_id=exclude_order_id,
    )


def _replace_processing_reserves(order_id: str, agency: Agency, stock_rows: list[dict]):
    if not order_id or not agency:
        return
    reserves: dict[tuple[str, str, str, str], int] = {}
    for row in stock_rows or []:
        if not isinstance(row, dict):
            continue
        sku = (row.get("article") or row.get("sku") or "").strip()
        if not sku:
            continue
        qty_value = _parse_qty_value(row.get("qty"))
        if qty_value is None or qty_value <= 0:
            continue
        size = str(row.get("size") or "").strip()
        barcode = str(row.get("barcode") or "").strip()
        goods_type = str(row.get("goods_type") or "").strip()
        key = (sku, size, barcode, goods_type)
        reserves[key] = reserves.get(key, 0) + qty_value
    if not reserves:
        WarehouseWritePathService.replace_processing_reserves(
            agency=agency,
            order_id=str(order_id),
            items=[],
        )
        return
    WarehouseWritePathService.replace_processing_reserves(
        agency=agency,
        order_id=str(order_id),
        items=[
            {
                "sku": sku,
                "sku_code": sku,
                "size": size,
                "barcode": barcode,
                "goods_type": goods_type,
                "qty": qty,
            }
            for (sku, size, barcode, goods_type), qty in reserves.items()
            if str(sku or "").strip()
        ],
    )


def _remaining_processing_stock_rows_for_reserve(
    stock_rows: list[dict] | None,
    factual_payload: dict | None,
) -> list[dict]:
    """Return only not-yet-processed quantities to keep in processing reserve."""
    rows = stock_rows or []
    if not isinstance(rows, list) or not rows:
        return []

    normalized_rows: list[dict] = []
    expected_by_sku_size: dict[tuple[str, str], int] = {}
    for row in rows:
        if not isinstance(row, dict):
            continue
        sku = (row.get("article") or row.get("sku") or "").strip()
        if not sku:
            continue
        size = (row.get("size") or "").strip()
        qty_value = _parse_qty_value(row.get("qty")) or 0
        if qty_value <= 0:
            continue
        sku_size_key = (sku.lower(), size.lower())
        expected_by_sku_size[sku_size_key] = expected_by_sku_size.get(sku_size_key, 0) + qty_value
        row_copy = dict(row)
        row_copy["qty"] = qty_value
        normalized_rows.append(row_copy)

    if not normalized_rows:
        return []

    factual_by_sku_size, _ = _processing_factual_qty_from_payload(factual_payload or {})
    remaining_by_sku_size: dict[tuple[str, str], int] = {}
    for key, expected_qty in expected_by_sku_size.items():
        factual_qty = factual_by_sku_size.get(key, 0)
        remaining_by_sku_size[key] = max(expected_qty - factual_qty, 0)

    remaining_rows: list[dict] = []
    for row in normalized_rows:
        sku = (row.get("article") or row.get("sku") or "").strip()
        size = (row.get("size") or "").strip()
        key = (sku.lower(), size.lower())
        remaining_qty = remaining_by_sku_size.get(key, 0)
        if remaining_qty <= 0:
            continue
        row_qty = _parse_qty_value(row.get("qty")) or 0
        keep_qty = min(row_qty, remaining_qty)
        if keep_qty <= 0:
            continue
        remaining_by_sku_size[key] = remaining_qty - keep_qty
        row_copy = dict(row)
        row_copy["qty"] = keep_qty
        remaining_rows.append(row_copy)
    return remaining_rows


def _submit_processing(request):
    return ProcessingWorkflowService.submit_processing(request=request)


class ProcessingHomeView(RoleRequiredMixin, TemplateView):
    template_name = "processing/processing.html"
    allowed_roles = ("manager", "storekeeper", "head_manager", "director", "admin")

    def _use_manager_template(self) -> bool:
        if getattr(self.request, "_client_agency", None):
            return False
        role = get_request_role(self.request)
        if role not in {"manager", "head_manager", "director", "admin"}:
            return False
        edit_flag = (self.request.GET.get("edit") or "").strip().lower()
        return edit_flag in {"1", "true", "yes"}

    def get_template_names(self):
        if self._use_manager_template():
            return ["processing/processing_manager.html"]
        return [self.template_name]

    def dispatch(self, request, *args, **kwargs):
        client_agency = _client_agency_from_request(request)
        if client_agency:
            request._client_agency = client_agency
            return TemplateView.dispatch(self, request, *args, **kwargs)
        return super().dispatch(request, *args, **kwargs)

    def get(self, request, *args, **kwargs):
        status = (request.GET.get("status") or "").lower()
        ok = request.GET.get("ok") == "1"
        context_kwargs = dict(kwargs)
        submitted = context_kwargs.get("submitted") or (ok and status != "draft")
        draft_saved = ok and status == "draft"
        error = context_kwargs.pop("error", None)
        order_id = request.GET.get("order")
        if order_id and not getattr(request, "_client_agency", None):
            latest_entry = (
                OrderAuditEntry.objects.filter(order_id=order_id, order_type="processing")
                .order_by("-created_at")
                .first()
            )
            if latest_entry and _is_draft_payload(latest_entry.payload or {}):
                return HttpResponseForbidden("Доступ запрещен")
        ctx = self.get_context_data(
            submitted=submitted,
            draft_saved=draft_saved,
            error=error,
            **context_kwargs,
        )
        return self.render_to_response(ctx)

    def post(self, request, *args, **kwargs):
        return _submit_processing(request)

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        ctx.update(
            ProcessingWorkflowService.build_processing_home_page_context(
                request=self.request,
                submitted=kwargs.get("submitted", False),
                draft_saved=kwargs.get("draft_saved", False),
                error=kwargs.get("error"),
            )
        )
        return ctx


class ProcessingDirectionsView(RoleRequiredMixin, TemplateView):
    template_name = "processing/processing_directions.html"
    allowed_roles = ("manager", "storekeeper", "head_manager", "director", "admin", "processing_head")

    def dispatch(self, request, *args, **kwargs):
        client_agency = _client_agency_from_request(request)
        if client_agency:
            request._client_agency = client_agency
            return TemplateView.dispatch(self, request, *args, **kwargs)
        return super().dispatch(request, *args, **kwargs)

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        ctx.update(
            ProcessingWorkflowService.build_processing_directions_page_context(
                request=self.request,
            )
        )
        return ctx


class ProcessingStockPickerView(RoleRequiredMixin, TemplateView):
    template_name = "processing/stock_picker.html"
    allowed_roles = ("manager", "storekeeper", "head_manager", "director", "admin")

    def dispatch(self, request, *args, **kwargs):
        client_agency = _client_agency_from_request(request)
        if client_agency:
            request._client_agency = client_agency
            return TemplateView.dispatch(self, request, *args, **kwargs)
        return super().dispatch(request, *args, **kwargs)

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        ctx.update(
            ProcessingWorkflowService.build_processing_stock_picker_page_context(
                request=self.request,
            )
        )
        return ctx


def delete_processing_draft(request, order_id: str):
    return ProcessingWorkflowService.delete_processing_draft(request=request, order_id=order_id)


def _processing_latest_packaging_assignment_state(entries) -> dict | None:
    for entry in reversed(entries or []):
        payload = entry.payload if isinstance(entry.payload, dict) else {}
        state = str(payload.get("packing_assignment_state") or "").strip().lower()
        if not state:
            continue
        assignee_id = str(payload.get("packing_assignee_id") or "").strip()
        return {
            "state": state,
            "assignee_id": assignee_id,
            "assignee_name": str(payload.get("packing_assignee") or "").strip(),
            "assignee_role": str(payload.get("packing_assignee_role") or "").strip(),
        }
    return None


def _processing_pending_packaging_assignment(entries) -> dict | None:
    state = _processing_latest_packaging_assignment_state(entries)
    if not state:
        return None
    if str(state.get("state") or "").strip().lower() != "pending":
        return None
    return state


def _processing_packaging_task_route(order_id: str) -> str:
    return f"/orders/processing/{order_id}/flow/"


def _processing_packaging_task_exists(order_id: str, assignee: Employee) -> bool:
    route = _processing_packaging_task_route(order_id)
    return Task.objects.filter(route=route, assigned_to=assignee).exclude(status="done").exists()


def _create_processing_packaging_task(order_id: str, assignee: Employee, request_user=None):
    description = f"Раскоробовка товара. Исполнитель: {assignee.full_name}."
    observer = get_employee_for_user(request_user) if request_user else None
    return Task.objects.create(
        title=f"Задача на раскоробовку товара по заявке №{order_id}",
        description=description,
        route=_processing_packaging_task_route(order_id),
        assigned_to=assignee,
        observer=observer,
        created_by=request_user if getattr(request_user, "is_authenticated", False) else None,
        due_date=timezone.localtime(),
    )


def _processing_auto_dispatch_pending_packaging(
    *,
    order_id: str,
    entries,
    can_open_processing_flow: bool,
    placement_completed: bool,
):
    pending = _processing_pending_packaging_assignment(entries)
    if not pending:
        return None, None
    if not can_open_processing_flow or placement_completed:
        return pending, None

    assignee_id = str(pending.get("assignee_id") or "").strip()
    if not assignee_id:
        return None, "invalid_worker"
    assignee = Employee.objects.filter(
        pk=assignee_id,
        role="processing_worker",
        is_active=True,
    ).first()
    latest = entries[-1] if entries else None
    if not assignee:
        log_order_action(
            "update",
            order_id=order_id,
            order_type="processing",
            user=None,
            agency=latest.agency if latest else None,
            description="Отложенное поручение на упаковку отменено: обработчик не найден.",
            payload={
                "packing_assignment_state": "invalid_worker",
                "packing_assignee_id": assignee_id,
                "packing_assignee": str(pending.get("assignee_name") or "").strip(),
                "packing_assignee_role": "processing_worker",
            },
        )
        return None, "invalid_worker"

    if _processing_packaging_task_exists(order_id, assignee):
        log_order_action(
            "update",
            order_id=order_id,
            order_type="processing",
            user=None,
            agency=latest.agency if latest else None,
            description=f"Отложенное поручение на упаковку уже активно: {assignee.full_name}",
            payload={
                "packing_assignment_state": "dispatched",
                "packing_assignee_id": assignee.id,
                "packing_assignee": assignee.full_name,
                "packing_assignee_role": assignee.role,
                "packing_assignment_dispatch_mode": "auto-existing",
            },
        )
        return None, None

    _create_processing_packaging_task(order_id, assignee)
    log_order_action(
        "update",
        order_id=order_id,
        order_type="processing",
        user=None,
        agency=latest.agency if latest else None,
        description=f"Отложенное поручение на упаковку автоматически отправлено: {assignee.full_name}",
        payload={
            "packing_assignment_state": "dispatched",
            "packing_assignee_id": assignee.id,
            "packing_assignee": assignee.full_name,
            "packing_assignee_role": assignee.role,
            "packing_assignment_dispatch_mode": "auto",
        },
    )
    return None, "auto_dispatched"


class ProcessingWorkView(RoleRequiredMixin, TemplateView):
    template_name = "processing/processing_work.html"
    allowed_roles = ("storekeeper", "processing_head", "head_manager", "director", "admin")

    def dispatch(self, request, *args, **kwargs):
        order_id = kwargs.get("order_id")
        if not order_id:
            return redirect("/orders/")
        if _client_agency_from_request(request):
            return HttpResponseForbidden("Доступ запрещен")
        entries = list(
            OrderAuditEntry.objects.filter(order_id=order_id, order_type="processing")
            .select_related("agency")
            .order_by("created_at")
        )
        if not entries:
            return redirect("/orders/")
        latest = entries[-1]
        payload = _processing_work_payload_from_entries(entries)
        placement_completed = _flow_closed_from_entries(entries)
        state_result = _processing_dispatch_state(str(order_id or ""), latest.agency, payload)
        is_ready = bool(state_result and state_result.code in _PROCESSING_WORK_WAREHOUSE_CODES)
        if not is_ready:
            is_ready = _processing_work_allowed_by_payload(payload, placement_completed=placement_completed)
        if not is_ready:
            return HttpResponseForbidden("Доступ запрещен")
        request._processing_work_entries = entries
        request._processing_work_payload = payload
        request._processing_work_agency = latest.agency
        return super().dispatch(request, *args, **kwargs)

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        order_id = kwargs.get("order_id")
        entries = getattr(self.request, "_processing_work_entries", None) or []
        payload = getattr(self.request, "_processing_work_payload", None) or {}
        agency = getattr(self.request, "_processing_work_agency", None)
        ctx.update(
            ProcessingWorkflowService.build_processing_work_page_context(
                order_id=str(order_id or ""),
                entries=entries,
                payload=payload,
                agency=agency,
                request=self.request,
                error=kwargs.get("error"),
            )
        )
        return ctx

    def post(self, request, *args, **kwargs):
        action = (request.POST.get("action") or "").strip().lower()
        requested_with = str(request.headers.get("X-Requested-With") or "").strip().lower()
        accepts = str(request.headers.get("Accept") or "").strip().lower()
        is_ajax = requested_with == "xmlhttprequest" or "application/json" in accepts
        if action not in {"finish_processing", "send_to_warehouse", "cancel_warehouse_moves"}:
            return HttpResponseForbidden("Доступ запрещен")
        role = get_request_role(request)
        if action == "send_to_warehouse":
            if role not in {"storekeeper", "processing_head"}:
                return HttpResponseForbidden("Доступ запрещен")
            order_id = kwargs.get("order_id")
            if not order_id:
                return redirect("/orders/")
            entries = list(
                OrderAuditEntry.objects.filter(order_id=order_id, order_type="processing")
                .select_related("agency")
                .order_by("created_at")
            )
            if not entries:
                return redirect("/orders/")
            result = ProcessingWorkflowService.send_processing_to_warehouse(
                order_id=str(order_id),
                entries=entries,
                request=request,
            )
            if result.status == "not_ready":
                return self.get(
                    request,
                    error=result.error_message,
                    order_id=order_id,
                )
            if result.status == "invalid_destination":
                if is_ajax:
                    return JsonResponse({"ok": False, "error": result.error_message}, status=400)
                return self.get(request, error=result.error_message, order_id=order_id)
            if result.status != "ok":
                if is_ajax:
                    return JsonResponse(
                        {
                            "ok": False,
                            "error": result.error_message,
                            "created": 0,
                            "skipped": result.skipped_existing_count,
                            "missing": result.skipped_missing_count,
                        },
                        status=400,
                    )
                return redirect(
                    f"/orders/processing/{order_id}/work/?warehouse_move={result.status}"
                    f"&warehouse_created=0"
                    f"&warehouse_skipped={result.skipped_existing_count}"
                    f"&warehouse_missing={result.skipped_missing_count}"
                )
            if is_ajax:
                return JsonResponse(
                    {
                        "ok": True,
                        "created": result.created_count,
                        "skipped": result.skipped_existing_count,
                        "missing": result.skipped_missing_count,
                    },
                    status=200,
                )
            return redirect(
                f"/orders/processing/{order_id}/work/?warehouse_move=ok"
                f"&warehouse_created={result.created_count}"
                f"&warehouse_skipped={result.skipped_existing_count}"
                f"&warehouse_missing={result.skipped_missing_count}"
            )
        if action == "cancel_warehouse_moves":
            if role not in {"storekeeper", "processing_head"}:
                return HttpResponseForbidden("Доступ запрещен")
            order_id = kwargs.get("order_id")
            if not order_id:
                return redirect("/orders/")
            entries = list(
                OrderAuditEntry.objects.filter(order_id=order_id, order_type="processing")
                .select_related("agency")
                .order_by("created_at")
            )
            if not entries:
                return redirect("/orders/")
            result = ProcessingWorkflowService.cancel_processing_warehouse_moves(
                order_id=str(order_id),
                entries=entries,
                request=request,
            )
            if result.status == "cancel_none":
                return redirect(
                    f"/orders/processing/{order_id}/work/?warehouse_move=cancel_none"
                    f"&warehouse_canceled=0&warehouse_skipped={result.skipped_count}"
                )
            return redirect(
                f"/orders/processing/{order_id}/work/?warehouse_move=canceled"
                f"&warehouse_canceled={result.canceled_count}&warehouse_skipped={result.skipped_count}"
            )
        if role not in {"storekeeper", "processing_head"}:
            return HttpResponseForbidden("Доступ запрещен")
        order_id = kwargs.get("order_id")
        if not order_id:
            return redirect("/orders/")
        entries = list(
            OrderAuditEntry.objects.filter(order_id=order_id, order_type="processing")
            .select_related("agency")
            .order_by("created_at")
        )
        if not entries:
            return redirect("/orders/")
        result = ProcessingWorkflowService.finish_processing(
            order_id=str(order_id),
            entries=entries,
            request=request,
            role=role,
        )
        if result.status == "already_done":
            return redirect(resolve_cabinet_url(role))
        if result.status in {"blocked", "discrepancy_reported_head", "discrepancy_reported_wait"}:
            return self.get(
                request,
                error=result.error_message,
                order_id=order_id,
            )
        return redirect(resolve_cabinet_url(role))


class ProcessingPlacementActView(OrdersProcessingPlacementActView):
    template_name = "processing/processing_placement_act.html"


class ProcessingFlowView(OrdersReceivingFlowView):
    template_name = "processing/processing_flow.html"
    directional_template_name = "processing/processing_flow_directions.html"
    order_type = "processing"
    allowed_roles = ("storekeeper", "processing_head", "processing_worker", "head_manager", "director", "admin", "manager")
    finish_roles = ("processing_head", "processing_worker")
    mismatch_finish_roles = ("processing_head",)
    box_reassign_roles = ("processing_head",)

    def _can_finish_flow_role(self, role: str) -> bool:
        return role in self.finish_roles

    def _can_finish_flow_with_mismatch_role(self, role: str) -> bool:
        return role in self.mismatch_finish_roles

    def _can_reassign_boxes_between_pallets_role(self, role: str) -> bool:
        return role in self.box_reassign_roles

    def _is_directional_unboxing_payload(self, payload: dict | None) -> bool:
        if not isinstance(payload, dict):
            return False
        return _processing_has_direction_distribution(payload)

    def _select_flow_template(self, payload: dict | None) -> None:
        if self._is_directional_unboxing_payload(payload):
            self.template_name = self.directional_template_name
        else:
            self.template_name = "processing/processing_flow.html"

    def _normalize_flow_state(self, boxes_data, pallets_data, active_box, active_pallet):
        boxes_data, pallets_data = _dedupe_pallet_box_links(boxes_data, pallets_data)
        state = super()._normalize_flow_state(boxes_data, pallets_data, active_box, active_pallet)
        owner_boxes = {}
        fixed_labels = {}
        box_directions = {}
        for raw in boxes_data or []:
            if not isinstance(raw, dict):
                continue
            code = str(raw.get("code") or "").strip()
            if not code:
                continue
            owner_boxes[code] = {
                "owner_agent_id": raw.get("owner_agent_id") or "",
                "owner_user_id": raw.get("owner_user_id"),
                "owner_user_label": raw.get("owner_user_label") or "",
            }
            fixed_label = str(raw.get("fixed_label") or "").strip()
            if fixed_label:
                fixed_labels[code] = fixed_label
            direction_name = str(raw.get("direction") or raw.get("direction_name") or "").strip()
            if direction_name:
                box_directions[code] = direction_name
        owner_pallets = {}
        closed_pallets = {}
        pallet_directions = {}
        for raw in pallets_data or []:
            if not isinstance(raw, dict):
                continue
            code = str(raw.get("code") or "").strip()
            if not code:
                continue
            owner_pallets[code] = {
                "owner_agent_id": raw.get("owner_agent_id") or "",
                "owner_user_id": raw.get("owner_user_id"),
                "owner_user_label": raw.get("owner_user_label") or "",
            }
            closed_pallets[code] = {
                "closed_by_agent_id": raw.get("closed_by_agent_id") or "",
                "closed_by_user_id": raw.get("closed_by_user_id"),
                "closed_by_user_label": raw.get("closed_by_user_label") or "",
            }
            direction_name = str(raw.get("direction") or raw.get("direction_name") or "").strip()
            if direction_name:
                pallet_directions[code] = direction_name
        for box in state.get("boxes") or []:
            box_code = box.get("code") or ""
            owner = owner_boxes.get(box_code)
            if owner:
                if owner.get("owner_agent_id") and not box.get("owner_agent_id"):
                    box["owner_agent_id"] = owner.get("owner_agent_id")
                if owner.get("owner_user_id") and not box.get("owner_user_id"):
                    box["owner_user_id"] = owner.get("owner_user_id")
                if owner.get("owner_user_label") and not box.get("owner_user_label"):
                    box["owner_user_label"] = owner.get("owner_user_label")
            if fixed_labels.get(box_code) and not box.get("fixed_label"):
                box["fixed_label"] = fixed_labels.get(box_code)
            if box_directions.get(box_code) and not str(box.get("direction") or "").strip():
                box["direction"] = box_directions.get(box_code)
        boxes_by_code = {}
        for box in state.get("boxes") or []:
            if not isinstance(box, dict):
                continue
            code = str(box.get("code") or "").strip()
            if code:
                boxes_by_code[code] = box
        for pallet in state.get("pallets") or []:
            owner = owner_pallets.get(pallet.get("code") or "")
            if not owner:
                continue
            if owner.get("owner_agent_id") and not pallet.get("owner_agent_id"):
                pallet["owner_agent_id"] = owner.get("owner_agent_id")
            if owner.get("owner_user_id") and not pallet.get("owner_user_id"):
                pallet["owner_user_id"] = owner.get("owner_user_id")
            if owner.get("owner_user_label") and not pallet.get("owner_user_label"):
                pallet["owner_user_label"] = owner.get("owner_user_label")
            closed = closed_pallets.get(pallet.get("code") or "")
            if not closed:
                continue
            if closed.get("closed_by_agent_id") and not pallet.get("closed_by_agent_id"):
                pallet["closed_by_agent_id"] = closed.get("closed_by_agent_id")
            if closed.get("closed_by_user_id") and not pallet.get("closed_by_user_id"):
                pallet["closed_by_user_id"] = closed.get("closed_by_user_id")
            if closed.get("closed_by_user_label") and not pallet.get("closed_by_user_label"):
                pallet["closed_by_user_label"] = closed.get("closed_by_user_label")
        for pallet in state.get("pallets") or []:
            if not isinstance(pallet, dict):
                continue
            owner_label = _resolve_actor_label(
                pallet.get("owner_user_label"),
                pallet.get("owner_agent_id"),
                pallet.get("owner_user_id"),
                unknown="",
            )
            if owner_label and not str(pallet.get("owner_user_label") or "").strip():
                pallet["owner_user_label"] = owner_label
            if not pallet.get("sealed"):
                continue
            closer_label = _resolve_actor_label(
                pallet.get("closed_by_user_label"),
                pallet.get("closed_by_agent_id"),
                pallet.get("closed_by_user_id"),
                unknown="",
            )
            if not closer_label:
                closer_label = owner_label
            if closer_label and not str(pallet.get("closed_by_user_label") or "").strip():
                pallet["closed_by_user_label"] = closer_label
            if not str(pallet.get("closed_by_user_id") or "").strip() and str(pallet.get("owner_user_id") or "").strip():
                pallet["closed_by_user_id"] = pallet.get("owner_user_id") or ""
            if not str(pallet.get("closed_by_agent_id") or "").strip() and str(pallet.get("owner_agent_id") or "").strip():
                pallet["closed_by_agent_id"] = pallet.get("owner_agent_id") or ""
        for pallet in state.get("pallets") or []:
            if not isinstance(pallet, dict):
                continue
            code = str(pallet.get("code") or "").strip()
            if pallet_directions.get(code) and not str(pallet.get("direction") or "").strip():
                pallet["direction"] = pallet_directions.get(code)
            direction_name = str(pallet.get("direction") or "").strip()
            if not direction_name:
                inferred = ""
                for box_code in pallet.get("boxes") or []:
                    box = boxes_by_code.get(str(box_code or "").strip())
                    if not box:
                        continue
                    box_direction = str(box.get("direction") or "").strip()
                    if box_direction:
                        inferred = box_direction
                        break
                if inferred:
                    pallet["direction"] = inferred
                    direction_name = inferred
            if not direction_name:
                continue
            for box_code in pallet.get("boxes") or []:
                box = boxes_by_code.get(str(box_code or "").strip())
                if not box:
                    continue
                if not str(box.get("direction") or "").strip():
                    box["direction"] = direction_name
        return state

    def _can_start(self, entries):
        if not entries:
            return False
        if _flow_closed_from_entries(entries):
            return False
        latest = entries[-1] if entries else None
        payload = _latest_payload_from_entries(entries)
        if not _processing_results_are_ready(payload, include_shipping=False):
            return False
        cards_total_count = _processing_cards_total(payload)
        processed_cards, placed_cards = _processing_card_sets(payload)
        if cards_total_count > 0 and len(processed_cards) < cards_total_count:
            return False
        ready_cards = processed_cards - placed_cards if processed_cards else set()
        items = _processing_receiving_items(
            payload,
            latest.agency_id if latest else None,
            ready_cards if ready_cards else None,
        )
        if not items:
            items = self._items_from_placement_act(entries)
        return bool(items)

    def _items_from_placement_act(self, entries):
        placement_entry = self._placement_act_entry(entries)
        if not placement_entry:
            return []
        placement_payload = placement_entry.payload or {}
        raw_items = placement_payload.get("act_items") or []
        if not isinstance(raw_items, list):
            return []
        items = []
        for raw in raw_items:
            if not isinstance(raw, dict):
                continue
            qty = _parse_qty_value(raw.get("actual_qty"))
            if qty is None:
                qty = _parse_qty_value(raw.get("qty")) or 0
            if qty <= 0:
                continue
            sku_code = str(raw.get("sku_code") or raw.get("sku") or "").strip()
            name = str(raw.get("name") or "").strip()
            size = str(raw.get("size") or "").strip()
            if not (sku_code or name or size):
                continue
            items.append(
                {
                    "sku_code": sku_code,
                    "name": name or "-",
                    "size": size,
                    "actual_qty": qty,
                }
            )
        return items

    def _mark_in_progress(self, request, order_id: str | None):
        return

    def _save_flow_draft(self, request, order_id, entries):
        result = ProcessingWorkflowService.save_processing_flow_draft(
            order_id=str(order_id),
            entries=entries,
            request=request,
            normalize_flow_state=self._normalize_flow_state,
            can_reassign_boxes=self._can_reassign_boxes_between_pallets_role(get_request_role(request)),
        )
        if result.status == "forbidden":
            return JsonResponse({"ok": False, "error": "forbidden"}, status=403)
        if result.status == "box_move_head_only":
            return JsonResponse({"ok": False, "error": "box_move_head_only"}, status=403)
        if result.status == "closed":
            return JsonResponse({"ok": False, "error": "closed"}, status=400)
        if result.status == "not_allowed":
            return JsonResponse({"ok": False, "error": "not_allowed"}, status=400)
        if result.status == "invalid_json":
            return JsonResponse({"ok": False, "error": "invalid_json"}, status=400)
        if result.status == "missing_agent_id":
            return JsonResponse({"ok": False, "error": "missing_agent_id"}, status=400)
        if result.status == "session_not_found":
            return JsonResponse({"ok": False, "error": "session_not_found"}, status=404)
        return JsonResponse({"ok": True, "session_id": result.session_id})

    def _reopen_flow(self, request, order_id, entries):
        result = ProcessingWorkflowService.reopen_processing_flow(
            order_id=str(order_id),
            entries=entries,
            request=request,
            normalize_flow_state=self._normalize_flow_state,
        )
        if result.status == "forbidden":
            return HttpResponseForbidden("Доступ запрещен")
        if result.status == "not_closed":
            return redirect(f"/orders/processing/{order_id}/flow/")
        return redirect(f"/orders/processing/{order_id}/flow/?reopen=1")

    def get(self, request, *args, **kwargs):
        order_id = kwargs.get("order_id")
        if get_request_role(request) == "storekeeper":
            self._mark_in_progress(request, order_id)
        entries = self._load_entries(order_id)
        if not entries:
            return redirect("/orders/")
        payload = _latest_payload_from_entries(entries)
        self._select_flow_template(payload)
        if not _flow_closed_from_entries(entries) and not self._can_start(entries):
            role = get_request_role(request)
            if role in {"storekeeper", "processing_head", "head_manager", "director", "admin"}:
                query = urlencode(
                    {
                        "error": (
                            "Нельзя открыть раскоробовку: сначала завершите все карты обработки "
                            "и заполните результаты (кроме этапа раскоробовки)."
                        )
                    }
                )
                return redirect(f"/orders/processing/{order_id}/work/?{query}")
            return HttpResponseForbidden("Раскоробовка недоступна: не заполнены карты обработки.")
        ok = request.GET.get("ok") == "1"
        error = request.GET.get("error")
        ctx = self.get_context_data(ok=ok, error=error, **kwargs)
        return self.render_to_response(ctx)

    def post(self, request, *args, **kwargs):
        order_id = kwargs.get("order_id")
        entries = self._load_entries(order_id)
        if not entries:
            return redirect("/orders/")
        if request.POST.get("flow_action") == "draft":
            return self._save_flow_draft(request, order_id, entries)
        if request.POST.get("flow_action") == "reopen":
            return self._reopen_flow(request, order_id, entries)
        role = get_request_role(request)
        can_finish = self._can_finish_flow_role(role)
        can_finish_with_mismatch = self._can_finish_flow_with_mismatch_role(role)
        can_reassign_boxes = self._can_reassign_boxes_between_pallets_role(role)
        if not can_finish:
            return HttpResponseForbidden("Доступ запрещен")
        if _flow_closed_from_entries(entries):
            return redirect(f"/orders/processing/{order_id}/flow/")
        def flow_error(code: str):
            return redirect(f"/orders/processing/{order_id}/flow/?error={code}")

        result = ProcessingWorkflowService.complete_processing_flow(
            order_id=str(order_id),
            entries=entries,
            request=request,
            normalize_flow_state=self._normalize_flow_state,
            items_from_placement_act=self._items_from_placement_act,
            can_start=self._can_start(entries),
            can_finish_with_mismatch=can_finish_with_mismatch,
            can_reassign_boxes=can_reassign_boxes,
            order_type=self.order_type,
        )
        if result.status != "ok":
            return flow_error(result.error_code)
        return redirect(f"/orders/processing/{order_id}/flow/?ok=1")

    def get_context_data(self, **kwargs):
        ctx = TemplateView.get_context_data(self, **kwargs)
        order_id = kwargs.get("order_id")
        role = get_request_role(self.request)
        entries = self._load_entries(order_id)
        ctx.update(
            ProcessingWorkflowService.build_processing_flow_page_context(
                order_id=str(order_id or ""),
                entries=entries,
                request=self.request,
                can_finish_flow=self._can_finish_flow_role(role),
                can_finish_flow_mismatch=self._can_finish_flow_with_mismatch_role(role),
                can_reassign_boxes=self._can_reassign_boxes_between_pallets_role(role),
                is_directional_unboxing_payload=self._is_directional_unboxing_payload,
                items_from_placement_act=self._items_from_placement_act,
                normalize_flow_state=self._normalize_flow_state,
                find_flow_state=self._find_flow_state,
                placement_act_entry=self._placement_act_entry,
                ok=kwargs.get("ok", False),
                error=kwargs.get("error"),
            )
        )
        return ctx


@login_required
@require_GET
def processing_flow_session(request, order_id: str):
    role = get_request_role(request)
    if role not in {"storekeeper", "processing_head", "processing_worker", "head_manager", "director", "admin", "manager"}:
        return HttpResponseForbidden("Доступ запрещен")
    result = ProcessingWorkflowService.get_processing_flow_session(
        order_id=str(order_id),
        request=request,
    )
    if result.status == "missing_agent_id":
        return JsonResponse({"ok": False, "error": "missing_agent_id"}, status=400)
    if result.status == "session_not_found":
        return JsonResponse({"ok": False, "error": "session_not_found"}, status=404)
    return JsonResponse(
        {
            "ok": True,
            "session_id": result.session_id,
            "flow_state": result.flow_state,
            "updated_at": result.updated_at,
        }
    )


@login_required
@require_GET
def processing_flow_shared(request, order_id: str):
    role = get_request_role(request)
    if role not in {"storekeeper", "processing_head", "processing_worker", "head_manager", "director", "admin", "manager"}:
        return HttpResponseForbidden("Доступ запрещен")
    result = ProcessingWorkflowService.get_processing_flow_shared_state(order_id=str(order_id))
    return JsonResponse({"ok": True, "flow_state": result.flow_state})


def processing_flow_box_action(request, order_id: str):
    if request.method != "POST":
        return JsonResponse({"ok": False, "error": "method_not_allowed"}, status=405)
    if not request.user.is_authenticated:
        return HttpResponseForbidden("Доступ запрещен")
    role = get_request_role(request)
    if role != "storekeeper":
        return HttpResponseForbidden("Доступ запрещен")
    result = ProcessingWorkflowService.log_processing_flow_box_action(
        order_id=str(order_id),
        request=request,
    )
    if result.status == "invalid_action":
        return JsonResponse({"ok": False, "error": "invalid_action"}, status=400)
    if result.status == "missing_box":
        return JsonResponse({"ok": False, "error": "missing_box"}, status=400)
    return JsonResponse({"ok": True})


@login_required
@require_POST
def processing_flow_marking_scan(request, order_id: str):
    result = ProcessingWorkflowService.scan_processing_flow_marking(
        order_id=str(order_id),
        request=request,
    )
    return JsonResponse(result.payload, status=result.http_status)


@login_required
@require_POST
def processing_assign_packaging(request, order_id: str):
    role = get_request_role(request)
    if role not in {"processing_head", "head_manager", "director", "admin"}:
        return HttpResponseForbidden("Доступ запрещен")
    result = ProcessingWorkflowService.assign_processing_packaging(
        order_id=str(order_id),
        request=request,
    )
    return redirect(result.redirect_to)


class ProcessingCardView(RoleRequiredMixin, TemplateView):
    template_name = "processing/processing_card.html"
    allowed_roles = ("storekeeper", "processing_head", "head_manager", "director", "admin")

    def dispatch(self, request, *args, **kwargs):
        order_id = kwargs.get("order_id")
        if not order_id:
            return redirect("/orders/")
        if _client_agency_from_request(request):
            return HttpResponseForbidden("Доступ запрещен")
        latest = (
            OrderAuditEntry.objects.filter(order_id=order_id, order_type="processing")
            .select_related("agency")
            .order_by("-created_at")
            .first()
        )
        if not latest:
            return redirect("/orders/")
        payload = latest.payload or {}
        state_result = _processing_dispatch_state(str(order_id or ""), latest.agency, payload)
        allowed = bool(state_result and state_result.code in _PROCESSING_CARD_WAREHOUSE_CODES)
        if not allowed:
            allowed = _processing_card_allowed_by_payload(payload)
        if not allowed:
            return HttpResponseForbidden("Доступ запрещен")
        request._processing_card_payload = payload
        request._processing_card_agency = latest.agency
        request._processing_card_status_label = (
            (state_result.label_for("processing") if state_result else "") or payload.get("status_label")
        )
        return super().dispatch(request, *args, **kwargs)

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        payload = getattr(self.request, "_processing_card_payload", None) or {}
        agency = getattr(self.request, "_processing_card_agency", None)
        ctx.update(
            ProcessingWorkflowService.build_processing_card_page_context(
                order_id=str(kwargs.get("order_id") or ""),
                card_id=str(kwargs.get("card_id") or ""),
                request=self.request,
                payload=payload,
                agency=agency,
            )
        )
        return ctx

    def render_to_response(self, context, **response_kwargs):
        response = super().render_to_response(context, **response_kwargs)
        response["Cache-Control"] = "no-store, no-cache, must-revalidate, max-age=0"
        response["Pragma"] = "no-cache"
        response["Vary"] = "Cookie"
        return response

    def post(self, request, *args, **kwargs):
        action = (request.POST.get("action") or "").strip().lower()
        if action not in {"finish_card", "save_results", "return_to_processing"}:
            return HttpResponseForbidden("Доступ запрещен")
        order_id = kwargs.get("order_id")
        if not order_id:
            return redirect("/orders/")
        result = ProcessingWorkflowService.handle_processing_card_action(
            order_id=str(order_id),
            request=request,
            action=action,
        )
        if result.status == "forbidden":
            return HttpResponseForbidden("Доступ запрещен")
        return redirect(result.redirect_to)


class ProcessingTechnicalCardView(ProcessingCardView):
    template_name = "processing/processing_technical_card.html"

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        payload = getattr(self.request, "_processing_card_payload", None) or {}
        ctx.update(
            ProcessingWorkflowService.build_processing_technical_card_page_context(
                ctx=ctx,
                order_id=str(kwargs.get("order_id") or ctx.get("order_id") or ""),
                card_id=str(kwargs.get("card_id") or ctx.get("card_id") or "card"),
                request=self.request,
                payload=payload,
            )
        )
        return ctx


class ProcessingLabelPrintView(ProcessingCardView):
    template_name = "processing/processing_label_print.html"

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        payload = getattr(self.request, "_processing_card_payload", None) or {}
        agency = getattr(self.request, "_processing_card_agency", None)
        ctx.update(
            ProcessingWorkflowService.build_processing_label_print_page_context(
                ctx=ctx,
                order_id=str(kwargs.get("order_id") or ctx.get("order_id") or ""),
                card_id=str(kwargs.get("card_id") or ctx.get("card_id") or ""),
                request=self.request,
                payload=payload,
                agency=agency,
            )
        )
        return ctx


class ProcessingDetailView(OrdersDetailView):
    order_type = "processing"
    allowed_roles = ("manager", "storekeeper", "head_manager", "director", "admin", "processing_head")

    def dispatch(self, request, *args, **kwargs):
        order_id = kwargs.get("order_id")
        client_agency = _client_agency_from_request(request)
        if client_agency:
            request._client_agency = client_agency
        if order_id and not client_agency:
            latest_entry = (
                OrderAuditEntry.objects.filter(order_id=order_id, order_type=self.order_type)
                .order_by("-created_at")
                .first()
            )
            if latest_entry and _is_draft_payload(latest_entry.payload or {}):
                return HttpResponseForbidden("Доступ запрещен")
        return super().dispatch(request, *args, **kwargs)

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        order_id = kwargs.get("order_id")
        entries_list = list(
            OrderAuditEntry.objects.filter(order_id=order_id, order_type=self.order_type)
            .select_related("user", "agency")
            .order_by("created_at")
        )
        ctx.update(
            ProcessingWorkflowService.build_processing_detail_page_context(
                ctx=ctx,
                order_id=str(order_id or ""),
                entries_list=entries_list,
                request=self.request,
                payload_from_entries=self._payload_from_entries,
            )
        )
        return ctx

    def post(self, request, *args, **kwargs):
        order_id = kwargs.get("order_id")
        if not order_id:
            return redirect("/orders/")
        result = ProcessingWorkflowService.handle_processing_detail_action(
            order_id=str(order_id),
            request=request,
            order_type=self.order_type,
            payload_from_entries=self._payload_from_entries,
        )
        if result.status == "forbidden":
            return HttpResponseForbidden("Доступ запрещен")
        return redirect(result.redirect_to)


@login_required
@require_POST
def processing_marking_availability(request):
    data = _parse_json_body(request)
    if data is None:
        return HttpResponseBadRequest("Invalid JSON")
    items = data.get("items") or []
    if not isinstance(items, list):
        return HttpResponseBadRequest("Invalid items")
    result = ProcessingWorkflowService.processing_marking_availability(request=request, data=data)
    if result.http_status == 403:
        return HttpResponseForbidden("Доступ запрещен")
    if result.http_status == 400 and result.payload.get("error") == "Клиент не выбран":
        return HttpResponseBadRequest("Клиент не выбран")
    return JsonResponse(result.payload, status=result.http_status)


@login_required
@require_POST
def processing_marking_import(request):
    cz_file = request.FILES.get("marking_cz_file") or request.FILES.get("file")
    if not cz_file:
        return JsonResponse({"ok": False, "error": "Файл не выбран."}, status=400)
    cards_payload = _parse_json_value(request.POST.get("cards_json"), [])
    result = ProcessingWorkflowService.processing_marking_import(
        request=request,
        cz_file=cz_file,
        cards_payload=cards_payload,
        order_id=str(request.POST.get("order_id") or "").strip(),
    )
    if result.http_status == 403:
        return HttpResponseForbidden("Доступ запрещен")
    if result.http_status == 400 and result.payload.get("error") == "Клиент не выбран":
        return HttpResponseBadRequest("Клиент не выбран")
    return JsonResponse(result.payload, status=result.http_status)



@login_required
@require_POST
def enqueue_processing_print_job(request):
    data = _parse_json_body(request)
    if data is None:
        return HttpResponseBadRequest("Invalid JSON")
    result = ProcessingWorkflowService.enqueue_processing_print_job(request=request, data=data)
    return JsonResponse(result.payload, status=result.http_status)


@require_GET
def processing_print_jobs_next(request):
    ok, response = _check_print_agent_token(request)
    if not ok:
        return response
    agent_name = (request.GET.get("agent") or request.headers.get("X-Print-Agent") or "").strip()
    result = ProcessingWorkflowService.processing_print_jobs_next(agent_name=agent_name)
    return JsonResponse(result.payload, status=result.http_status)


@csrf_exempt
@require_POST
def processing_print_jobs_complete(request):
    ok, response = _check_print_agent_token(request)
    if not ok:
        return response
    data = _parse_json_body(request)
    if data is None:
        data = request.POST
    result = ProcessingWorkflowService.processing_print_jobs_complete(data=data)
    return JsonResponse(result.payload, status=result.http_status)


@login_required
@require_GET
def processing_print_jobs_status(request):
    ok, response = _require_print_admin(request)
    if not ok:
        return response
    agent_name = str(request.GET.get("agent_id") or "").strip()
    result = ProcessingWorkflowService.processing_print_jobs_status(agent_name=agent_name)
    return JsonResponse(result.payload, status=result.http_status)


@login_required
@require_POST
def processing_print_jobs_pause(request):
    ok, response = _require_print_admin(request)
    if not ok:
        return response
    data = _parse_json_body(request) or {}
    result = ProcessingWorkflowService.processing_print_jobs_pause(request=request, data=data)
    return JsonResponse(result.payload, status=result.http_status)


@login_required
@require_POST
def processing_print_jobs_resume(request):
    ok, response = _require_print_admin(request)
    if not ok:
        return response
    data = _parse_json_body(request) or {}
    result = ProcessingWorkflowService.processing_print_jobs_resume(request=request, data=data)
    return JsonResponse(result.payload, status=result.http_status)


@login_required
@require_POST
def processing_print_jobs_clear(request):
    ok, response = _require_print_admin(request)
    if not ok:
        return response
    data = _parse_json_body(request) or {}
    result = ProcessingWorkflowService.processing_print_jobs_clear(request=request, data=data)
    return JsonResponse(result.payload, status=result.http_status)


@login_required
@require_POST
def processing_print_jobs_reset(request):
    ok, response = _require_print_admin(request)
    if not ok:
        return response
    data = _parse_json_body(request) or {}
    result = ProcessingWorkflowService.processing_print_jobs_reset(request=request, data=data)
    return JsonResponse(result.payload, status=result.http_status)


@login_required
@require_POST
def processing_print_jobs_recover(request):
    ok, response = _require_print_admin(request)
    if not ok:
        return response
    data = _parse_json_body(request) or {}
    result = ProcessingWorkflowService.processing_print_jobs_recover(request=request, data=data)
    return JsonResponse(result.payload, status=result.http_status)


@login_required
@require_GET
def download_processing_print_agent_cmd(request):
    token = _get_print_agent_token()
    if not token:
        return HttpResponseBadRequest("PRINT_AGENT_TOKEN not set")
    server_url = request.build_absolute_uri("/").rstrip("/")
    script_url = f"{server_url}/orders/processing/print-agent/script/?token={token}"
    content = (
        "@echo off\r\n"
        f"set SERVER_URL={server_url}\r\n"
        f"set TOKEN={token}\r\n"
        "set AGENT_DIR=%~dp0\r\n"
        "powershell -ExecutionPolicy Bypass -Command "
        f"\"Invoke-WebRequest -Uri '{script_url}' -OutFile '%AGENT_DIR%\\print_agent.ps1'\""
        "\r\n"
        "powershell -ExecutionPolicy Bypass -File \"%AGENT_DIR%print_agent.ps1\" "
        "-ServerUrl \"%SERVER_URL%\" -Token \"%TOKEN%\"\r\n"
        "pause\r\n"
    )
    response = HttpResponse(content, content_type="application/octet-stream")
    response["Content-Disposition"] = "attachment; filename=print_agent_start.cmd"
    return response


@login_required
@require_GET
def download_processing_print_agent_install_cmd(request):
    token = _get_print_agent_token()
    if not token:
        return HttpResponseBadRequest("PRINT_AGENT_TOKEN not set")
    server_url = request.build_absolute_uri("/").rstrip("/")
    script_url = f"{server_url}/orders/processing/print-agent/script/?token={token}"
    content = (
        "@echo off\r\n"
        "setlocal\r\n"
        f"set SERVER_URL={server_url}\r\n"
        f"set TOKEN={token}\r\n"
        "set AGENT_DIR=%APPDATA%\\FullboxPrintAgent\r\n"
        "if not exist \"%AGENT_DIR%\" mkdir \"%AGENT_DIR%\"\r\n"
        "powershell -ExecutionPolicy Bypass -Command "
        f"\"Invoke-WebRequest -Uri '{script_url}' -OutFile '%AGENT_DIR%\\print_agent.ps1'\""
        "\r\n"
        "set STARTUP_DIR=%APPDATA%\\Microsoft\\Windows\\Start Menu\\Programs\\Startup\r\n"
        "set VBS_FILE=%STARTUP_DIR%\\FullboxPrintAgent.vbs\r\n"
        "> \"%VBS_FILE%\" echo Set WshShell = CreateObject(\"WScript.Shell\")\r\n"
        ">> \"%VBS_FILE%\" echo WshShell.Run \"powershell -WindowStyle Hidden -ExecutionPolicy Bypass -File \"\"%AGENT_DIR%\\print_agent.ps1\"\" -ServerUrl \"\"%SERVER_URL%\"\" -Token \"\"%TOKEN%\"\"\", 0\r\n"
        ">> \"%VBS_FILE%\" echo Set WshShell = Nothing\r\n"
        "wscript \"%VBS_FILE%\"\r\n"
        "echo Autostart installed.\r\n"
        "pause\r\n"
    )
    response = HttpResponse(content, content_type="application/octet-stream")
    response["Content-Disposition"] = "attachment; filename=print_agent_install.cmd"
    return response


@login_required
@require_GET
def download_processing_printer_sync_package(request):
    path = settings.BASE_DIR.parent / "sync_printers.ps1"
    if not path.exists():
        return HttpResponseBadRequest("sync_printers.ps1 not found")
    ps1_content = path.read_text(encoding="utf-8")
    encoded = base64.b64encode(ps1_content.encode("utf-16le")).decode("ascii")
    cmd_content = (
        "@echo off\r\n"
        f"powershell -NoProfile -ExecutionPolicy Bypass -EncodedCommand {encoded}\r\n"
        "pause\r\n"
    )
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("printer_sync.ps1", ps1_content)
        archive.writestr("run_printer_sync.cmd", cmd_content)
        archive.writestr(
            "README.txt",
            "Запускайте run_printer_sync.cmd. Файл printer_sync.ps1 вручную не запускать.\r\n",
        )
    buffer.seek(0)
    response = HttpResponse(buffer.getvalue(), content_type="application/zip")
    response["Content-Disposition"] = "attachment; filename=printer_sync_package.zip"
    return response

@require_GET
def download_processing_print_agent_script(request):
    ok, response = _check_print_agent_token(request)
    if not ok:
        return response
    path = settings.BASE_DIR.parent / "print_agent.ps1"
    if not path.exists():
        return HttpResponseBadRequest("print_agent.ps1 not found")
    content = path.read_text(encoding="utf-8")
    response = HttpResponse(content, content_type="text/plain")
    response["Content-Disposition"] = "attachment; filename=print_agent.ps1"
    return response

