import json
import base64
import io
import re
import uuid
import zipfile
from datetime import datetime, timedelta
from pathlib import Path
from urllib.parse import urlparse, urlencode

from django.conf import settings
from django.contrib.auth.decorators import login_required
from django.db import transaction, DatabaseError
from django.db.models import Count, Q
from django.shortcuts import redirect
from django.http import HttpResponse, HttpResponseBadRequest, HttpResponseForbidden, JsonResponse
from django.utils import timezone
from django.views.decorators.csrf import csrf_exempt
from django.views.decorators.http import require_GET, require_POST
from django.views.generic import TemplateView

from audit.models import OrderAuditEntry, log_order_action, log_staff_overaction
from employees.access import (
    RoleRequiredMixin,
    get_employee_for_user,
    get_request_role,
    resolve_cabinet_url,
    is_staff_role,
)
from employees.models import Employee
from labels.utils import (
    LABEL_SIZES,
    SCANNER_EOLS,
    load_available_printers_data,
    load_label_settings,
    load_print_agent_status,
    load_scanner_settings,
    save_print_agent_status,
    set_print_agent_pause,
)
from marking.models import MarkingCode
from marking.utils import extract_processing_items
from orders.views import (
    OrdersDetailView,
    ProcessingPlacementActView as OrdersProcessingPlacementActView,
    ReceivingFlowView as OrdersReceivingFlowView,
    _current_status_entry,
    _flow_closed_from_entries,
    _item_key,
    _latest_payload_from_entries,
    _parse_int_value,
    _processing_receiving_items,
    _shorten_ip_name,
    _status_label_from_entry,
)
from sku.models import Agency, SKU, SKUBarcode
from sklad.models import InventoryState
from todo.models import Task
from .models import ProcessingFlowSession, ProcessingPrintJob
from agent.models import AgentCommand, DeviceAgent
from openpyxl import load_workbook

GOODS_TYPE_LABELS = {
    "op": "Оптовый",
    "gv": "Готовый",
    "br": "Брак",
    "vz": "Возврат",
    "rh": "Расходный",
    "no": "Не обработанный",
}


def _format_payload_value(value):
    if value is None or value == "":
        return "-"
    return str(value).strip()


def _format_payload_list(value):
    if value is None or value == "":
        return "-"
    if isinstance(value, (list, tuple, set)):
        items = [str(item).strip() for item in value if str(item).strip()]
        return ", ".join(items) if items else "-"
    text = str(value).strip()
    return text if text else "-"


_ORG_FORM_RE = re.compile(r"\bобщество\s+с\s+ограниченной\s+ответственностью\b", re.IGNORECASE)


def _normalize_org_name(value) -> str:
    if value is None:
        return ""
    text = str(value)
    if not text.strip():
        return text.strip()
    text = _ORG_FORM_RE.sub("ООО", text)
    text = re.sub(r"\s{2,}", " ", text).strip()
    return text


def _non_empty_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Да" if value else ""
    if isinstance(value, (int, float)):
        return "" if value == 0 else str(value)
    text = str(value).strip()
    if not text or text in {"-", "0", "0.0"}:
        return ""
    if text.isdigit() and int(text) == 0:
        return ""
    return text


def _format_list_value(value) -> str:
    if value is None:
        return ""
    if isinstance(value, (list, tuple, set)):
        items = [str(item).strip() for item in value if str(item).strip()]
        return ", ".join(items)
    return _non_empty_text(value)


def _processing_card_id(card: dict) -> str:
    if not isinstance(card, dict):
        return ""
    value = card.get("id") or card.get("article") or card.get("sku") or ""
    return str(value).strip()


def _processing_card_sets(payload: dict) -> tuple[set[str], set[str]]:
    payload = payload or {}
    processed = set()
    placed = set()
    for card in payload.get("cards") or []:
        card_id = _processing_card_id(card)
        if not card_id:
            continue
        if card.get("processed_at") or card.get("processed_done") or card.get("processed"):
            processed.add(card_id)
        if card.get("placed_at") or card.get("placed_done"):
            placed.add(card_id)
    for value in payload.get("processed_cards") or []:
        text = str(value).strip()
        if text:
            processed.add(text)
    for value in payload.get("placed_cards") or []:
        text = str(value).strip()
        if text:
            placed.add(text)
    return processed, placed


def _parse_json_value(raw, fallback):
    if raw is None:
        return fallback
    if isinstance(raw, (dict, list)):
        return raw
    text = str(raw).strip()
    if not text:
        return fallback
    try:
        return json.loads(text)
    except (TypeError, ValueError):
        return fallback


def _parse_json_body(request):
    if not request.body:
        return {}
    try:
        return json.loads(request.body.decode("utf-8"))
    except (UnicodeDecodeError, json.JSONDecodeError):
        return None


def _default_flow_state() -> dict:
    return {
        "boxes": [],
        "pallets": [],
        "activeBox": "",
        "activePallet": "",
    }


def _flow_session_for_request(order_id: str, agent_id: str, request, create: bool = False):
    agent_id = str(agent_id or "").strip()
    if not agent_id:
        return None
    user = request.user if request.user.is_authenticated else None
    employee = get_employee_for_user(request.user) if user else None
    qs = ProcessingFlowSession.objects.filter(
        order_id=order_id,
        order_type="processing",
        agent_id=agent_id,
        status=ProcessingFlowSession.STATUS_OPEN,
    )
    if user:
        qs = qs.filter(user=user)
    session = qs.order_by("-updated_at").first()
    if not session and create:
        session = ProcessingFlowSession.objects.create(
            order_id=order_id,
            order_type="processing",
            agent_id=agent_id,
            user=user,
            employee=employee,
            flow_state=_default_flow_state(),
            status=ProcessingFlowSession.STATUS_OPEN,
            last_seen=timezone.localtime(),
        )
        return session
    if session:
        update_fields = []
        now = timezone.localtime()
        if session.last_seen != now:
            session.last_seen = now
            update_fields.append("last_seen")
        if user and session.user_id != user.id:
            session.user = user
            update_fields.append("user")
        if employee and session.employee_id != employee.id:
            session.employee = employee
            update_fields.append("employee")
        if update_fields:
            session.save(update_fields=update_fields + ["updated_at"])
    return session


def _flow_owner_label(session: ProcessingFlowSession | None) -> str:
    if not session:
        return ""
    employee = getattr(session, "employee", None)
    if employee and getattr(employee, "full_name", None):
        return employee.full_name
    user = getattr(session, "user", None)
    if user:
        full_name = user.get_full_name() or ""
        if full_name:
            return full_name
        return getattr(user, "username", "") or ""
    return ""


def _apply_flow_owner(values: list, session: ProcessingFlowSession | None) -> list:
    if not session:
        return values
    owner_agent = session.agent_id or ""
    owner_user_id = session.user_id
    owner_label = _flow_owner_label(session)
    patched = []
    for raw in values or []:
        if not isinstance(raw, dict):
            continue
        entry = dict(raw)
        if owner_agent and not entry.get("owner_agent_id"):
            entry["owner_agent_id"] = owner_agent
        if owner_user_id and not entry.get("owner_user_id"):
            entry["owner_user_id"] = owner_user_id
        if owner_label and not entry.get("owner_user_label"):
            entry["owner_user_label"] = owner_label
        patched.append(entry)
    return patched


def _merge_flow_sessions(sessions: list[ProcessingFlowSession]) -> tuple[list, list]:
    def _session_sort_key(session: ProcessingFlowSession) -> float:
        ts = getattr(session, "updated_at", None) or getattr(session, "last_seen", None)
        if not ts:
            return 0.0
        try:
            return float(ts.timestamp())
        except Exception:
            return 0.0

    sessions_sorted = sorted(sessions, key=_session_sort_key)
    boxes_raw: list = []
    pallets_raw: list = []
    for session in sessions_sorted:
        state = session.flow_state if isinstance(session.flow_state, dict) else {}
        if isinstance(state.get("boxes"), list):
            boxes_raw.extend(_apply_flow_owner(state.get("boxes") or [], session))
        if isinstance(state.get("pallets"), list):
            pallets_raw.extend(_apply_flow_owner(state.get("pallets") or [], session))
    boxes_map: dict[str, dict] = {}
    for box in boxes_raw:
        if not isinstance(box, dict):
            continue
        code = str(box.get("code") or "").strip()
        if not code:
            continue
        boxes_map[code] = box
    pallets_map: dict[str, dict] = {}
    for pallet in pallets_raw:
        if not isinstance(pallet, dict):
            continue
        code = str(pallet.get("code") or "").strip()
        if not code:
            continue
        pallets_map[code] = pallet
    boxes = list(boxes_map.values())
    pallets = list(pallets_map.values())
    box_codes = {box.get("code") for box in boxes}
    for pallet in pallets:
        if not isinstance(pallet, dict):
            continue
        pallet_boxes = [
            str(code).strip()
            for code in (pallet.get("boxes") or [])
            if str(code or "").strip()
        ]
        pallet["boxes"] = [code for code in pallet_boxes if code in box_codes]
    return boxes, pallets


def _state_has_open_box_with_items(flow_state: dict) -> bool:
    if not isinstance(flow_state, dict):
        return False
    boxes = flow_state.get("boxes") or []
    if not isinstance(boxes, list):
        return False
    for box in boxes:
        if not isinstance(box, dict):
            continue
        if box.get("sealed"):
            continue
        items = box.get("items") or []
        if not isinstance(items, list):
            continue
        total = sum(_parse_qty_value(item.get("qty")) or 0 for item in items if isinstance(item, dict))
        if total > 0:
            return True
    return False


def _order_has_open_boxes(order_id: str) -> bool:
    sessions = ProcessingFlowSession.objects.filter(
        order_id=order_id,
        order_type="processing",
        status=ProcessingFlowSession.STATUS_OPEN,
    )
    for session in sessions:
        if _state_has_open_box_with_items(session.flow_state or {}):
            return True
    return False


def _marking_required_by_barcode(payload: dict) -> tuple[dict[str, int], int, int]:
    required: dict[str, int] = {}
    missing_barcodes = 0
    total_required = 0
    for item in extract_processing_items(payload):
        qty = item.get("qty") or 0
        if qty <= 0:
            continue
        barcode = str(item.get("barcode") or "").strip()
        if not barcode:
            missing_barcodes += qty
            continue
        total_required += qty
        required[barcode] = required.get(barcode, 0) + qty
    return required, total_required, missing_barcodes


def _marking_available_by_barcode(agency: Agency | None, order_id: str | None) -> dict[str, int]:
    if not agency:
        return {}
    qs = (
        MarkingCode.objects.filter(agency=agency, order_type="processing", used_at__isnull=True)
        .exclude(barcode="")
    )
    if order_id:
        qs = qs.filter(Q(order_id__isnull=True) | Q(order_id="") | Q(order_id=order_id))
    else:
        qs = qs.filter(Q(order_id__isnull=True) | Q(order_id=""))
    return {row["barcode"]: row["count"] for row in qs.values("barcode").annotate(count=Count("id"))}


def _marking_free_by_barcode(agency: Agency | None) -> dict[str, int]:
    if not agency:
        return {}
    qs = (
        MarkingCode.objects.filter(
            agency=agency,
            order_type="processing",
            used_at__isnull=True,
        )
        .exclude(barcode="")
        .filter(Q(order_id__isnull=True) | Q(order_id=""))
    )
    return {row["barcode"]: row["count"] for row in qs.values("barcode").annotate(count=Count("id"))}


def _reserve_marking_codes(
    agency: Agency | None,
    order_id: str | None,
    required_map: dict[str, int],
) -> tuple[bool, str | None]:
    if not agency or not order_id or not required_map:
        return True, None
    order_id = str(order_id).strip()
    if not order_id:
        return True, None
    with transaction.atomic():
        reserved_qs = MarkingCode.objects.filter(
            agency=agency,
            order_type="processing",
            order_id=order_id,
            used_at__isnull=True,
        )
        reserved_map = {
            row["barcode"]: row["count"]
            for row in reserved_qs.values("barcode").annotate(count=Count("id"))
        }
        for barcode, required_qty in required_map.items():
            reserved_qty = reserved_map.get(barcode, 0)
            if reserved_qty <= required_qty:
                continue
            release_qty = reserved_qty - required_qty
            release_ids = list(
                MarkingCode.objects.select_for_update()
                .filter(
                    agency=agency,
                    order_type="processing",
                    order_id=order_id,
                    used_at__isnull=True,
                    barcode=barcode,
                )
                .order_by("-created_at")
                .values_list("id", flat=True)[:release_qty]
            )
            if release_ids:
                MarkingCode.objects.filter(id__in=release_ids).update(order_id="")

        for barcode, required_qty in required_map.items():
            reserved_qty = reserved_map.get(barcode, 0)
            needed = max(required_qty - reserved_qty, 0)
            if not needed:
                continue
            free_ids = list(
                MarkingCode.objects.select_for_update()
                .filter(
                    agency=agency,
                    order_type="processing",
                    used_at__isnull=True,
                    barcode=barcode,
                )
                .filter(Q(order_id__isnull=True) | Q(order_id=""))
                .values_list("id", flat=True)[:needed]
            )
            if len(free_ids) < needed:
                return False, "Недостаточно свободных ЧЗ для брони."
            MarkingCode.objects.filter(id__in=free_ids).update(order_id=order_id)
    return True, None


def _normalize_marking_import_order_id(order_id: str | None) -> str:
    value = str(order_id or "").strip()
    if value.startswith("draft-"):
        return ""
    return value


def _normalize_excel_cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, int):
        return str(value)
    return str(value).strip()


def _normalize_marking_code(value: str) -> str:
    if not value:
        return ""
    # Excel encodes ASCII 29 (GS) as _x001D_ in XML exports.
    text = str(value)
    text = re.sub(r"_x001d_", "\x1d", text, flags=re.IGNORECASE)
    text = re.sub(r"<\s*gs\s*>", "\x1d", text, flags=re.IGNORECASE)
    text = text.replace("\\u001d", "\x1d").replace("\\x1d", "\x1d").replace("\\x001d", "\x1d")
    return text.strip()


def _import_marking_codes(file, payload: dict, order_id: str, agency: Agency, user):
    try:
        workbook = load_workbook(file, read_only=True, data_only=True)
    except Exception:
        return False, {"error": "Не удалось прочитать .xlsx файл."}
    sheet = workbook.active

    rows = []
    barcodes = set()
    invalid_rows = 0
    for idx, row in enumerate(sheet.iter_rows(values_only=True), start=1):
        if idx == 1:
            continue
        barcode = _normalize_excel_cell(row[0]) if row and len(row) > 0 else ""
        code = _normalize_excel_cell(row[1]) if row and len(row) > 1 else ""
        code = _normalize_marking_code(code)
        if not barcode or not code:
            if barcode or code:
                invalid_rows += 1
            continue
        rows.append((barcode, code))
        barcodes.add(barcode)

    if not rows:
        return False, {"error": "В файле нет данных для импорта."}

    barcode_qs = SKUBarcode.objects.select_related("sku").filter(value__in=barcodes)
    barcode_map = {item.value: item for item in barcode_qs}
    existing_codes = set(
        MarkingCode.objects.filter(code__in=[code for _, code in rows]).values_list("code", flat=True)
    )

    items = extract_processing_items(payload)
    allowed_pairs = {(item["sku_code"], item["size"]) for item in items}
    allowed_barcodes = {str(item.get("barcode") or "").strip() for item in items}
    allowed_barcodes.discard("")
    bind_order_id = _normalize_marking_import_order_id(order_id)
    added = 0
    duplicates = 0
    unknown_barcodes = 0
    mismatched_barcodes = 0
    seen_codes = set()
    to_create = []

    for barcode, code in rows:
        if code in seen_codes:
            duplicates += 1
            continue
        seen_codes.add(code)
        if code in existing_codes:
            duplicates += 1
            continue
        barcode_obj = barcode_map.get(barcode)
        if not barcode_obj or not barcode_obj.sku:
            unknown_barcodes += 1
            continue
        sku_obj = barcode_obj.sku
        if agency and sku_obj.agency and sku_obj.agency_id != agency.id:
            mismatched_barcodes += 1
            continue
        sku_code = sku_obj.sku_code
        size = (barcode_obj.size or sku_obj.size or "").strip()
        order_binding = ""
        if allowed_pairs:
            if (sku_code, size) in allowed_pairs:
                order_binding = bind_order_id
            elif (sku_code, "") in allowed_pairs:
                order_binding = bind_order_id
                size = ""
            else:
                matched = [pair for pair in allowed_pairs if pair[0] == sku_code]
                if len(matched) == 1:
                    order_binding = bind_order_id
                    size = matched[0][1]
        if not order_binding and barcode in allowed_barcodes:
            order_binding = bind_order_id
        to_create.append(
            MarkingCode(
                order_type="processing",
                order_id=order_binding,
                agency=agency,
                sku=sku_obj,
                sku_code=sku_code,
                size=size,
                barcode=barcode,
                code=code,
                source="import",
                created_by=user if getattr(user, "is_authenticated", False) else None,
            )
        )

    if to_create:
        with transaction.atomic():
            MarkingCode.objects.bulk_create(to_create, batch_size=500)
        added = len(to_create)

    return True, {
        "added": added,
        "duplicates": duplicates,
        "unknown_barcodes": unknown_barcodes,
        "mismatched_barcodes": mismatched_barcodes,
        "invalid_rows": invalid_rows,
    }




def _get_print_agent_token() -> str:
    return str(getattr(settings, "PRINT_AGENT_TOKEN", "")).strip()


def _check_print_agent_token(request):
    expected = _get_print_agent_token()
    if not expected:
        return False, JsonResponse({"ok": False, "error": "PRINT_AGENT_TOKEN not set"}, status=403)
    token = (
        request.headers.get("X-Print-Token")
        or request.GET.get("token")
        or request.POST.get("token")
        or ""
    )
    if token != expected:
        return False, JsonResponse({"ok": False, "error": "Invalid token"}, status=403)
    return True, None


def _require_print_admin(request):
    if not request.user.is_authenticated:
        return False, HttpResponseForbidden("Доступ запрещен")
    role = get_request_role(request)
    if role not in {"storekeeper", "processing_head", "head_manager", "director", "admin"}:
        return False, HttpResponseForbidden("Доступ запрещен")
    return True, None


def _print_queue_counts() -> dict:
    return {
        "pending": ProcessingPrintJob.objects.filter(status=ProcessingPrintJob.STATUS_PENDING).count(),
        "printing": ProcessingPrintJob.objects.filter(status=ProcessingPrintJob.STATUS_PRINTING).count(),
        "failed": ProcessingPrintJob.objects.filter(status=ProcessingPrintJob.STATUS_FAILED).count(),
    }


def _enqueue_agent_command(command: str, payload: dict | None = None, agent_id: str = "") -> None:
    if not command:
        return
    AgentCommand.objects.create(
        agent_id=agent_id or "",
        command=command,
        payload=payload or {},
    )


def _serialize_print_job(job: ProcessingPrintJob) -> dict:
    return {
        "id": job.id,
        "order_id": job.order_id,
        "card_id": job.card_id,
        "article": job.article,
        "barcode": job.barcode,
        "size": job.size,
        "printer_name": job.printer_name,
        "label_png_base64": job.label_png_base64,
        "label_width_mm": job.label_width_mm,
        "label_height_mm": job.label_height_mm,
        "status": job.status,
        "requested_by": job.requested_by,
        "created_at": job.created_at.isoformat() if job.created_at else "",
    }


def _short_city(value: str) -> str:
    text = str(value or "").strip()
    if not text:
        return ""
    city = text.split(",")[0].strip()
    lowered = city.lower()
    for prefix in ("г.", "г ", "город "):
        if lowered.startswith(prefix):
            city = city[len(prefix):].strip()
            break
    return city or text


def _processing_params_from_payload(payload: dict) -> list[dict]:
    def _add_param(rows: list[dict], label: str, value: str | None, extras: list[str] | None = None) -> None:
        base = _non_empty_text(value)
        extra_values = [item for item in (extras or []) if _non_empty_text(item)]
        if not base and not extra_values:
            return
        if base and extra_values:
            value_text = f"{base}; " + "; ".join(extra_values)
        elif base:
            value_text = base
        else:
            value_text = "; ".join(extra_values)
        rows.append({"label": label, "value": value_text})

    insert_types = payload.get("insert_types") or []
    if isinstance(insert_types, str):
        insert_types = [insert_types] if insert_types else []
    insert_other = (payload.get("insert_other") or "").strip()
    if insert_other:
        insert_types = list(insert_types) + [insert_other]

    marking_stickers = payload.get("marking_stickers") or []
    if isinstance(marking_stickers, str):
        marking_stickers = [marking_stickers] if marking_stickers else []
    marking_sizes = payload.get("marking_sizes") or []
    if isinstance(marking_sizes, str):
        marking_sizes = [marking_sizes] if marking_sizes else []

    processing_params: list[dict] = []
    _add_param(processing_params, "Маркетплейс", payload.get("marketplace"))
    defect_percent = _non_empty_text(payload.get("defect_percent"))
    if defect_percent and defect_percent.isdigit():
        defect_percent = f"{defect_percent}%"
    _add_param(processing_params, "Проверка на брак", defect_percent)
    _add_param(processing_params, "Маркировка 58/40", payload.get("marking_5840_qty"))
    _add_param(processing_params, "Маркировка 58/40 (шт/чз)", payload.get("marking_5840_each_qty"))
    _add_param(processing_params, "Замена бирок", payload.get("tag_owner"))

    def _add_pack_param(label: str, base_value: str | None, prefix: str) -> None:
        extras = []
        type_value = _non_empty_text(payload.get(f"{prefix}_type"))
        size_value = _non_empty_text(payload.get(f"{prefix}_size"))
        qty_value = _non_empty_text(payload.get(f"{prefix}_qty"))
        supply_value = _non_empty_text(payload.get(f"{prefix}_supply"))
        if type_value:
            extras.append(f"Тип: {type_value}")
        if size_value:
            extras.append(f"Размер: {size_value}")
        if qty_value:
            extras.append(f"Кол-во: {qty_value}")
        if supply_value:
            extras.append(f"Закупка: {supply_value}")
        _add_param(
            processing_params,
            label,
            base_value or payload.get(f"{prefix}_needed"),
            extras=extras,
        )

    _add_pack_param("Замена пакета", payload.get("bag_replace_type"), "bag_replace")
    _add_pack_param("Упаковка в Бабл пленку", payload.get("bubble_wrap_size"), "bubble_wrap")
    _add_pack_param("Упаковка в термо пленку", payload.get("shrink_wrap_size"), "shrink_wrap")
    _add_pack_param("Замена гофрокороба", payload.get("box_replace_type"), "box_replace")
    set_qty = _non_empty_text(payload.get("set_qty"))
    _add_param(
        processing_params,
        "Сборка набора",
        set_qty and f"Кол-во: {set_qty}" or payload.get("set_build"),
    )
    insert_qty = _non_empty_text(payload.get("insert_qty"))
    insert_types_label = _format_list_value(insert_types)
    if insert_other and insert_other not in insert_types:
        insert_types_label = _format_list_value(list(insert_types) + [insert_other])
    _add_param(
        processing_params,
        "Вложение",
        insert_qty and f"Кол-во: {insert_qty}" or payload.get("insert_needed"),
        extras=[f"Типы: {insert_types_label}" if insert_types_label else ""],
    )
    direction_mode = (payload.get("direction_needed") or "").strip()
    direction_file = _non_empty_text(payload.get("direction_file"))
    direction_addresses = _parse_json_value(payload.get("direction_addresses_json"), [])
    if isinstance(direction_addresses, dict):
        direction_addresses = direction_addresses.get("directions") or direction_addresses.get("addresses") or []
    if not isinstance(direction_addresses, list):
        direction_addresses = []
    direction_addresses = [str(item).strip() for item in direction_addresses if str(item).strip()]
    direction_plan = _parse_json_value(payload.get("direction_plan_json"), {})
    if not isinstance(direction_plan, dict):
        direction_plan = {}
    if not direction_addresses:
        plan_dirs = direction_plan.get("directions") or direction_plan.get("addresses") or []
        if isinstance(plan_dirs, list):
            direction_addresses = [str(item).strip() for item in plan_dirs if str(item).strip()]
    direction_count = _parse_qty_value(payload.get("direction_count"))
    if not direction_count and direction_addresses:
        direction_count = len(direction_addresses)
    direction_value = ""
    direction_extras: list[str] = []
    if direction_file or direction_mode in {"file", "Да"}:
        direction_value = direction_file and f"Файл: {direction_file}" or "Файл"
    elif direction_mode == "set" or direction_plan.get("rows"):
        if direction_addresses:
            direction_value = _format_list_value(direction_addresses)
        else:
            direction_value = "Задано"
    elif direction_mode in {"none", "Нет", "Отсутствует"}:
        direction_value = "Отсутствует"
    elif direction_mode:
        direction_value = direction_mode
    _add_param(
        processing_params,
        "Распределение по направлениям",
        direction_value,
        extras=direction_extras,
    )
    box_forming = _non_empty_text(payload.get("box_forming"))
    if box_forming == "other":
        box_forming = _non_empty_text(payload.get("box_forming_other"))
    _add_param(processing_params, "Формирование короба", box_forming)
    _add_param(processing_params, "Прочие", payload.get("comments"))
    _add_param(processing_params, "Маркировка", _format_list_value(marking_stickers))
    _add_param(processing_params, "Размеры стикеров", _format_list_value(marking_sizes))
    _add_param(processing_params, "Информационный", payload.get("marking_info"))
    _add_param(processing_params, "Вытянуть из мешка и наклеить ЧЗ", payload.get("pull_from_bag"))
    _add_param(processing_params, "Проверка на брак (кол-во)", payload.get("defect_qty"))
    _add_param(processing_params, "Обрезание ниток (кол-во)", payload.get("trim_threads_qty"))
    _add_param(processing_params, "Скрепление скотчем (кол-во)", payload.get("tape_qty"))
    _add_param(processing_params, "Удаление бирки", payload.get("remove_tag"))
    _add_param(processing_params, "Удаление бирки (кол-во)", payload.get("remove_tag_qty"))
    _add_param(processing_params, "Скрепление бирки", payload.get("attach_tag"))
    _add_param(processing_params, "Скрепление бирки (кол-во)", payload.get("attach_tag_qty"))
    return processing_params


def _parse_qty_value(raw: str | None) -> int | None:
    if raw is None:
        return None
    text = str(raw).strip()
    if not text:
        return None
    try:
        return int(text)
    except (TypeError, ValueError):
        return None


def _expected_processing_results(payload: dict) -> list[tuple[str, str, str]]:
    direction_plan = _parse_json_value(payload.get("direction_plan_json"), {})
    direction_addresses = _parse_json_value(payload.get("direction_addresses_json"), [])
    if isinstance(direction_addresses, dict):
        direction_addresses = (
            direction_addresses.get("directions")
            or direction_addresses.get("addresses")
            or []
        )
    if not isinstance(direction_addresses, list):
        direction_addresses = []
    if not direction_addresses and isinstance(direction_plan, dict):
        plan_dirs = direction_plan.get("directions") or direction_plan.get("addresses") or []
        if isinstance(plan_dirs, list):
            direction_addresses = plan_dirs
    direction_labels = [_short_city(item) or str(item).strip() for item in direction_addresses]
    plan_rows = []
    if isinstance(direction_plan, dict):
        plan_rows = direction_plan.get("rows") or []
    if not isinstance(plan_rows, list):
        plan_rows = []
    expected = []
    for row in plan_rows:
        if not isinstance(row, dict):
            continue
        quantities = row.get("quantities") or []
        if not isinstance(quantities, (list, tuple)):
            quantities = []
        article = str(row.get("article") or row.get("product_name") or "").strip().lower()
        size = str(row.get("size") or "").strip().lower()
        for idx, label in enumerate(direction_labels):
            qty = _parse_qty_value(quantities[idx] if idx < len(quantities) else None) or 0
            if qty <= 0:
                continue
            expected.append((article, size, label.lower()))
    return expected


def _is_draft_payload(payload: dict | None) -> bool:
    payload = payload or {}
    status_value = (payload.get("status") or payload.get("submit_action") or "").lower()
    status_label = (payload.get("status_label") or "").lower()
    return status_value == "draft" or "черновик" in status_label


def _normalize_goods_type(value):
    text = str(value or "").strip()
    if not text or text == "-":
        return ""
    lowered = text.lower()
    if lowered in GOODS_TYPE_LABELS:
        return GOODS_TYPE_LABELS[lowered].lower()
    return lowered


def _processing_reserve_maps(
    agency: Agency | None,
    exclude_order_id: str | None = None,
) -> tuple[dict[tuple[str, str, str], int], dict[tuple[str, str], int]]:
    if not agency:
        return {}, {}
    reserves = InventoryState.objects.filter(agency=agency, state="processing")
    if exclude_order_id:
        reserves = reserves.exclude(order_type="processing", order_id=str(exclude_order_id))
    reserve_map: dict[tuple[str, str, str], int] = {}
    reserve_any: dict[tuple[str, str], int] = {}
    for entry in reserves:
        sku = (entry.sku or "").strip()
        if not sku:
            continue
        size = (entry.size or "").strip()
        goods_key = _normalize_goods_type(entry.goods_type)
        key = (sku.lower(), size.lower(), goods_key)
        reserve_map[key] = reserve_map.get(key, 0) + (entry.qty or 0)
        any_key = (sku.lower(), size.lower())
        reserve_any[any_key] = reserve_any.get(any_key, 0) + (entry.qty or 0)
    return reserve_map, reserve_any


def _reserved_qty(
    reserve_map: dict[tuple[str, str, str], int],
    reserve_any: dict[tuple[str, str], int],
    sku: str,
    size: str,
    goods_type: str | None = None,
) -> int:
    sku_key = (sku or "").strip().lower()
    if not sku_key:
        return 0
    size_key = (size or "").strip().lower()
    goods_key = _normalize_goods_type(goods_type)
    if not goods_key:
        return reserve_any.get((sku_key, size_key), 0)
    return reserve_map.get((sku_key, size_key, goods_key), 0) + reserve_map.get(
        (sku_key, size_key, ""),
        0,
    )


def _next_order_number(order_type: str = "receiving") -> str:
    order_ids = (
        OrderAuditEntry.objects.filter(order_type=order_type)
        .values_list("order_id", flat=True)
        .distinct()
    )
    max_number = 0
    for order_id in order_ids:
        if not order_id:
            continue
        candidate = str(order_id).strip()
        if not re.fullmatch(r"\d+", candidate):
            continue
        try:
            number = int(candidate)
        except (TypeError, ValueError):
            continue
        if number > max_number:
            max_number = number
    next_number = max_number + 1
    while OrderAuditEntry.objects.filter(order_type=order_type, order_id=str(next_number)).exists():
        next_number += 1
    return str(next_number)


def _manager_due_date(submitted_at):
    cutoff = submitted_at.replace(hour=14, minute=0, second=0, microsecond=0)
    if submitted_at <= cutoff:
        return submitted_at.replace(hour=18, minute=0, second=0, microsecond=0)
    next_day = submitted_at + timedelta(days=1)
    return next_day.replace(hour=13, minute=0, second=0, microsecond=0)


def _create_processing_manager_task(order_id, agency, request, submitted_at):
    if not agency:
        return
    manager = (
        Employee.objects.filter(role="manager", is_active=True)
        .order_by("full_name")
        .first()
    )
    if not manager:
        return
    route = f"/orders/processing/{order_id}/"
    existing = Task.objects.filter(route=route, assigned_to=manager).exclude(status="done")
    if existing.exists():
        return
    description = f"Клиент: {agency.agn_name or agency.inn or agency.id}"
    Task.objects.create(
        title=f"Подтвердите заявку на обработку №{order_id}",
        description=description,
        route=route,
        assigned_to=manager,
        created_by=request.user if request.user.is_authenticated else None,
        due_date=_manager_due_date(submitted_at),
    )


def _create_processing_head_task(order_id, agency, request, submitted_at):
    if not agency:
        return
    head = (
        Employee.objects.filter(role="processing_head", is_active=True)
        .order_by("full_name")
        .first()
    )
    if not head:
        return
    route = f"/orders/processing/{order_id}/"
    existing = Task.objects.filter(route=route, assigned_to=head).exclude(status="done")
    if existing.exists():
        return
    description = f"Клиент: {agency.agn_name or agency.inn or agency.id}"
    Task.objects.create(
        title=f"Заявка на обработку №{order_id}",
        description=description,
        route=route,
        assigned_to=head,
        created_by=request.user if request.user.is_authenticated else None,
        due_date=submitted_at or timezone.localtime(),
    )


def _client_agency_from_request(request):
    if not request.user.is_authenticated:
        return None
    role = get_request_role(request)
    if is_staff_role(role):
        return None
    client_id = request.GET.get("client") or request.GET.get("agency")
    if client_id:
        return Agency.objects.filter(pk=client_id, portal_user=request.user).first()
    return Agency.objects.filter(portal_user=request.user).first()


def _barcode_value_for_sku(sku, size: str | None) -> str:
    if not sku:
        return "-"
    barcodes = list(getattr(sku, "barcodes", []).all())
    if not barcodes:
        return "-"
    size_value = (size or "").strip()
    if size_value:
        for barcode in barcodes:
            if (barcode.size or "").strip() == size_value:
                return barcode.value
    primary = next((barcode for barcode in barcodes if barcode.is_primary), None)
    return primary.value if primary else barcodes[0].value


def _inventory_items_for_agency(
    agency: Agency | None,
    exclude_order_id: str | None = None,
) -> list[dict]:
    if not agency:
        return []
    order_ids = list(
        OrderAuditEntry.objects.filter(order_type__in=("receiving", "processing"), agency=agency)
        .values_list("order_id", flat=True)
        .distinct()
    )
    if not order_ids:
        return []
    entries = list(
        OrderAuditEntry.objects.filter(order_type__in=("receiving", "processing"), order_id__in=order_ids)
        .select_related("agency")
        .order_by("-created_at")
    )
    goods_type_labels = GOODS_TYPE_LABELS
    goods_type_by_order = {}
    for entry in entries:
        if entry.order_id in goods_type_by_order:
            continue
        payload = entry.payload or {}
        goods_type = (payload.get("goods_type") or "").strip().lower()
        goods_label = (payload.get("goods_type_label") or "").strip()
        if not goods_label and goods_type in goods_type_labels:
            goods_label = goods_type_labels[goods_type]
        if goods_label or goods_type:
            goods_type_by_order[entry.order_id] = goods_label or goods_type
    latest_by_order = {}
    blocked_orders = set()
    for entry in entries:
        if entry.order_id in latest_by_order or entry.order_id in blocked_orders:
            continue
        payload = entry.payload or {}
        if payload.get("act") != "placement":
            continue
        state = (payload.get("act_state") or "closed").lower()
        if state != "closed":
            blocked_orders.add(entry.order_id)
            continue
        latest_by_order[entry.order_id] = entry

    totals = {}

    def add_item(item, goods_label: str):
        sku = (item.get("sku") or item.get("sku_code") or "").strip()
        name = (item.get("name") or "").strip()
        size = (item.get("size") or "").strip()
        qty = _parse_qty_value(item.get("qty"))
        if qty is None:
            qty = _parse_qty_value(item.get("actual_qty")) or 0
        if not any((sku, name, size)):
            return
        goods_label = goods_label or "-"
        key = (sku, name, size, goods_label)
        entry = totals.setdefault(
            key,
            {"sku": sku, "name": name, "size": size, "qty": 0, "goods_type": goods_label},
        )
        entry["qty"] += qty

    for entry in latest_by_order.values():
        payload = entry.payload or {}
        goods_label = goods_type_by_order.get(entry.order_id, "-")
        boxes = payload.get("act_boxes") or []
        pallets = payload.get("act_pallets") or []
        for box in boxes:
            for item in (box or {}).get("items") or []:
                add_item(item, goods_label)
        for pallet in pallets:
            for item in (pallet or {}).get("items") or []:
                add_item(item, goods_label)
        if not boxes and not pallets:
            if payload.get("act_items_removed"):
                continue
            for item in payload.get("act_items") or []:
                add_item(item, goods_label)

    sku_codes = {item["sku"] for item in totals.values() if item.get("sku")}
    sku_map = {}
    if sku_codes:
        for sku in (
            SKU.objects.filter(agency=agency, deleted=False, sku_code__in=sku_codes)
            .prefetch_related("barcodes", "photos")
        ):
            sku_map[sku.sku_code] = sku

    def normalize_photo_url(url: str) -> str:
        if not url:
            return ""
        if url.startswith(("http://", "https://", "/")):
            return url
        return f"/{url}"

    def sku_photo_url(sku_obj: SKU | None) -> str:
        if not sku_obj:
            return ""
        url = (sku_obj.img or "").strip()
        if url:
            return normalize_photo_url(url)
        photos = list(getattr(sku_obj, "photos", []).all())
        if photos:
            url = (photos[0].url or "").strip()
            return normalize_photo_url(url)
        return ""

    reserve_map, reserve_any = _processing_reserve_maps(agency, exclude_order_id=exclude_order_id)
    items = []
    for item in totals.values():
        sku_obj = sku_map.get(item.get("sku"))
        barcode = _barcode_value_for_sku(sku_obj, item.get("size")) if sku_obj else "-"
        photo_url = sku_photo_url(sku_obj)
        reserved_qty = _reserved_qty(
            reserve_map,
            reserve_any,
            item.get("sku") or "",
            item.get("size") or "",
            item.get("goods_type") or "",
        )
        available_qty = max((item.get("qty") or 0) - reserved_qty, 0)
        if available_qty <= 0:
            continue
        items.append(
            {
                "sku": item.get("sku") or "",
                "name": item.get("name") or "",
                "size": item.get("size") or "",
                "barcode": barcode or "-",
                "qty": available_qty,
                "goods_type": item.get("goods_type") or "-",
                "photo": photo_url,
            }
        )
    items.sort(
        key=lambda row: (
            row.get("name") or "",
            row.get("goods_type") or "",
            row.get("size") or "",
            row.get("sku") or "",
        )
    )
    return items


def _replace_processing_reserves(order_id: str, agency: Agency, stock_rows: list[dict]):
    if not order_id or not agency:
        return
    InventoryState.objects.filter(
        agency=agency,
        order_type="processing",
        order_id=str(order_id),
        state="processing",
    ).delete()
    reserves: dict[tuple[str, str, str, str], int] = {}
    for row in stock_rows or []:
        if not isinstance(row, dict):
            continue
        sku = (row.get("article") or row.get("sku") or "").strip()
        if not sku:
            continue
        qty_value = _parse_qty_value(row.get("qty"))
        if qty_value is None or qty_value <= 0:
            continue
        size = (row.get("size") or "").strip()
        barcode = (row.get("barcode") or "").strip()
        goods_type = (row.get("goods_type") or "").strip()
        key = (sku, size, barcode, goods_type)
        reserves[key] = reserves.get(key, 0) + qty_value
    if not reserves:
        return
    InventoryState.objects.bulk_create(
        [
            InventoryState(
                agency=agency,
                order_type="processing",
                order_id=str(order_id),
                sku=sku,
                size=size,
                barcode=barcode,
                goods_type=goods_type,
                qty=qty,
                state="processing",
            )
            for (sku, size, barcode, goods_type), qty in reserves.items()
        ]
    )


def _submit_processing(request):
    autosave = (request.POST.get("draft_autosave") or "").strip() == "1"

    def autosave_error(message: str):
        if not autosave:
            return None
        return JsonResponse({"ok": False, "error": message}, status=400)

    submit_action = (request.POST.get("submit_action") or "send").strip().lower()
    is_draft = submit_action == "draft"
    draft_order_id = (request.POST.get("draft_order_id") or "").strip()
    edit_order_id = (request.POST.get("edit_order_id") or "").strip()
    client_agency = getattr(request, "_client_agency", None) or _client_agency_from_request(request)
    if client_agency:
        agency = client_agency
    else:
        agency_id = request.POST.get("agency_id")
        agency = Agency.objects.filter(pk=agency_id).first()
    if not agency:
        error_response = autosave_error("Выберите клиента.")
        if error_response:
            return error_response
        return ProcessingHomeView().get(request, error="Выберите клиента.")

    role = get_request_role(request)
    existing_entries = []
    preserved_status = ""
    preserved_label = ""
    preserved_submit_action = ""
    if edit_order_id:
        if role not in {"manager", "head_manager", "director", "admin"}:
            return HttpResponseForbidden("Доступ запрещен")
        existing_entries = list(
            OrderAuditEntry.objects.filter(order_id=edit_order_id, order_type="processing")
            .order_by("created_at")
        )
        if not existing_entries:
            error_response = autosave_error("Заявка не найдена.")
            if error_response:
                return error_response
            return ProcessingHomeView().get(request, error="Заявка не найдена.")
        latest_payload = existing_entries[-1].payload or {}
        preserved_status = (latest_payload.get("status") or latest_payload.get("submit_action") or "").strip()
        preserved_label = (latest_payload.get("status_label") or "").strip()
        preserved_submit_action = (latest_payload.get("submit_action") or "").strip()
        status_lower = preserved_status.lower()
        label_lower = preserved_label.lower()
        if status_lower in {"done", "completed", "closed", "finished"} or "выполн" in label_lower:
            error_response = autosave_error(
                "Заявка уже утверждена и недоступна для редактирования.",
            )
            if error_response:
                return error_response
            return ProcessingHomeView().get(
                request,
                error="Заявка уже утверждена и недоступна для редактирования.",
            )
        is_draft = False

    existing_draft_entries = []
    if draft_order_id:
        draft_qs = OrderAuditEntry.objects.filter(
            order_id=draft_order_id,
            order_type="processing",
            agency=agency,
        ).order_by("created_at")
        existing_draft_entries = list(draft_qs)
        latest_draft = existing_draft_entries[-1] if existing_draft_entries else None
        if not latest_draft or not _is_draft_payload(latest_draft.payload or {}):
            draft_order_id = ""
            existing_draft_entries = []

    cards_payload = []
    cards_json = (request.POST.get("cards_json") or "").strip()
    if cards_json:
        try:
            parsed_cards = json.loads(cards_json)
        except json.JSONDecodeError:
            parsed_cards = []
        if isinstance(parsed_cards, list):
            for entry in parsed_cards:
                if not isinstance(entry, dict):
                    continue
                card_id = str(entry.get("id") or "").strip()
                article = str(entry.get("article") or "").strip()
                product_name = str(entry.get("product_name") or "").strip()
                photo_url = str(entry.get("photo_url") or "").strip()
                goods_type = str(entry.get("goods_type") or "").strip()
                rows = []
                for row in entry.get("rows") or []:
                    if not isinstance(row, dict):
                        continue
                    rows.append(
                        {
                            "article": str(row.get("article") or row.get("sku") or "").strip(),
                            "size": str(row.get("size") or "").strip(),
                            "barcode": str(row.get("barcode") or "").strip(),
                            "qty": row.get("qty"),
                        }
                    )
                photo_field = str(entry.get("photo_field") or "").strip()
                card_payload = {
                    "id": card_id,
                    "article": article,
                    "product_name": product_name,
                    "photo_url": photo_url,
                    "goods_type": goods_type,
                    "rows": rows,
                }
                if photo_field:
                    photo_file = request.FILES.get(photo_field)
                    if photo_file and getattr(photo_file, "name", ""):
                        card_payload["product_photo"] = photo_file.name
                cards_payload.append(card_payload)

    product_name = (request.POST.get("product_name") or "").strip()
    if cards_payload:
        product_name = cards_payload[0].get("product_name") or product_name
    if not is_draft:
        if cards_payload:
            if not any(card.get("product_name") for card in cards_payload):
                error_response = autosave_error("Укажите наименование товара.")
                if error_response:
                    return error_response
                return ProcessingHomeView().get(request, error="Укажите наименование товара.")
        elif not product_name:
            error_response = autosave_error("Укажите наименование товара.")
            if error_response:
                return error_response
            return ProcessingHomeView().get(request, error="Укажите наименование товара.")

    def collect_rows(field_map: dict) -> list[dict]:
        row_count = 0
        for values in field_map.values():
            row_count = max(row_count, len(values))
        rows = []
        for idx in range(row_count):
            row = {}
            for key, values in field_map.items():
                row[key] = values[idx] if idx < len(values) else ""
            if any(str(value).strip() for value in row.values()):
                rows.append(row)
        return rows

    size_rows = collect_rows(
        {
            "size_no": request.POST.getlist("size_no[]"),
            "size_value": request.POST.getlist("size_value[]"),
            "barcode": request.POST.getlist("size_barcode[]"),
            "recount_qty": request.POST.getlist("size_recount[]"),
            "processing_qty": request.POST.getlist("size_processing[]"),
            "unboxing_qty": request.POST.getlist("size_unpacking[]"),
            "defect_qty": request.POST.getlist("size_defect[]"),
        }
    )
    stock_rows = collect_rows(
        {
            "article": request.POST.getlist("stock_article[]"),
            "size": request.POST.getlist("stock_size[]"),
            "barcode": request.POST.getlist("stock_barcode[]"),
            "qty": request.POST.getlist("stock_qty[]"),
        }
    )
    global_article = (request.POST.get("article") or "").strip()
    if cards_payload:
        stock_rows = []
        for card in cards_payload:
            base_article = (card.get("article") or "").strip()
            goods_type = (card.get("goods_type") or "").strip()
            for row in card.get("rows") or []:
                article_value = (row.get("article") or base_article).strip()
                stock_rows.append(
                    {
                        "article": article_value,
                        "size": (row.get("size") or "").strip(),
                        "barcode": (row.get("barcode") or "").strip(),
                        "qty": row.get("qty"),
                        "goods_type": goods_type,
                    }
                )
    elif global_article:
        for row in stock_rows:
            if not (row.get("article") or "").strip():
                row["article"] = global_article
    unboxing_rows = collect_rows(
        {
            "date": request.POST.getlist("unboxing_date[]"),
            "box_size": request.POST.getlist("unboxing_box_size[]"),
            "multiple": request.POST.getlist("unboxing_multiple[]"),
            "box_qty": request.POST.getlist("unboxing_box_qty[]"),
            "pallet_qty": request.POST.getlist("unboxing_pallet_qty[]"),
            "storage_zone": request.POST.getlist("unboxing_storage_zone[]"),
        }
    )
    if not is_draft and stock_rows:
        available_map = {}
        barcode_map = {}
        for item in _inventory_items_for_agency(agency, exclude_order_id=edit_order_id or None):
            sku_key = (item.get("sku") or "").strip().lower()
            size_key = (item.get("size") or "").strip().lower()
            qty_value = _parse_qty_value(item.get("qty")) or 0
            if sku_key:
                key = (sku_key, size_key)
                available_map[key] = available_map.get(key, 0) + qty_value
            barcode_value = (item.get("barcode") or "").strip()
            if barcode_value:
                barcode_map[barcode_value] = barcode_map.get(barcode_value, 0) + qty_value
        for row in stock_rows:
            qty_value = _parse_qty_value(row.get("qty"))
            if qty_value is None:
                continue
            sku_key = (row.get("article") or "").strip().lower()
            size_key = (row.get("size") or "").strip().lower()
            barcode_value = (row.get("barcode") or "").strip()
            max_qty = None
            if sku_key:
                max_qty = available_map.get((sku_key, size_key))
            if max_qty is None and barcode_value:
                max_qty = barcode_map.get(barcode_value)
            if max_qty is None:
                max_qty = 0
            if qty_value > max_qty:
                sku_label = row.get("article") or "-"
                size_label = row.get("size") or "-"
                return ProcessingHomeView().get(
                    request,
                    error=(
                        f"Количество для {sku_label} ({size_label}) превышает остаток: {max_qty}."
                    ),
                )

    primary_article = (request.POST.get("article") or "").strip()
    primary_photo_url = (request.POST.get("product_photo_url") or "").strip()
    if cards_payload:
        primary_article = cards_payload[0].get("article") or primary_article
        primary_photo_url = cards_payload[0].get("photo_url") or primary_photo_url

    payload = {
        "email": (request.POST.get("email") or "").strip(),
        "fio": (request.POST.get("fio") or "").strip(),
        "org": (request.POST.get("org") or "").strip(),
        "product_name": product_name,
        "product_photo_url": primary_photo_url,
        "marketplace": request.POST.get("marketplace"),
        "supplier": request.POST.get("supplier"),
        "brand": request.POST.get("brand"),
        "subject": request.POST.get("subject"),
        "article": primary_article,
        "wb_article": request.POST.get("wb_article"),
        "color": request.POST.get("color"),
        "composition": request.POST.get("composition"),
        "gender": request.POST.get("gender"),
        "season": request.POST.get("season"),
        "order_no": request.POST.get("order_no"),
        "purchase_1c_no": request.POST.get("purchase_1c_no"),
        "purchase_1c_date": request.POST.get("purchase_1c_date"),
        "project_manager": request.POST.get("project_manager"),
        "warehouse_receiving": request.POST.get("warehouse_receiving"),
        "warehouse_packing": request.POST.get("warehouse_packing"),
        "warehouse_unpacking": request.POST.get("warehouse_unpacking"),
        "size_rows": size_rows,
        "stock_rows": stock_rows,
        "measure_needed": request.POST.get("measure_needed"),
        "measure_weight": request.POST.get("measure_weight"),
        "measure_width": request.POST.get("measure_width"),
        "measure_height": request.POST.get("measure_height"),
        "measure_depth": request.POST.get("measure_depth"),
        "defect_check": request.POST.get("defect_check"),
        "defect_percent": request.POST.get("defect_percent"),
        "defect_qty": request.POST.get("defect_qty"),
        "trim_threads_qty": request.POST.get("trim_threads_qty"),
        "tape_qty": request.POST.get("tape_qty"),
        "remove_tag": request.POST.get("remove_tag"),
        "remove_tag_qty": request.POST.get("remove_tag_qty"),
        "attach_tag": request.POST.get("attach_tag"),
        "attach_tag_qty": request.POST.get("attach_tag_qty"),
        "tag_replace_needed": request.POST.get("tag_replace_needed"),
        "tag_owner": request.POST.get("tag_owner"),
        "marking_stickers": request.POST.getlist("marking_stickers[]"),
        "marking_sizes": request.POST.getlist("marking_sizes[]"),
        "marking_info": request.POST.get("marking_info"),
        "marking_5840_needed": request.POST.get("marking_5840_needed"),
        "marking_5840_qty": request.POST.get("marking_5840_qty"),
        "marking_5840_each_needed": request.POST.get("marking_5840_each_needed"),
        "marking_5840_each_qty": request.POST.get("marking_5840_each_qty"),
        "set_build": request.POST.get("set_build"),
        "set_qty": request.POST.get("set_qty"),
        "insert_needed": request.POST.get("insert_needed"),
        "insert_types": request.POST.getlist("insert_types[]"),
        "insert_other": request.POST.get("insert_other"),
        "insert_qty": request.POST.get("insert_qty"),
        "pull_from_bag": request.POST.get("pull_from_bag"),
        "bubble_wrap_needed": request.POST.get("bubble_wrap_needed"),
        "bubble_wrap_type": request.POST.get("bubble_wrap_type"),
        "bubble_wrap_size": request.POST.get("bubble_wrap_size"),
        "bubble_wrap_qty": request.POST.get("bubble_wrap_qty"),
        "bubble_wrap_supply": request.POST.get("bubble_wrap_supply"),
        "bag_replace_needed": request.POST.get("bag_replace_needed"),
        "bag_replace_type": request.POST.get("bag_replace_type"),
        "bag_replace_size": request.POST.get("bag_replace_size"),
        "bag_replace_qty": request.POST.get("bag_replace_qty"),
        "bag_replace_supply": request.POST.get("bag_replace_supply"),
        "box_replace_needed": request.POST.get("box_replace_needed"),
        "box_replace_type": request.POST.get("box_replace_type"),
        "box_replace_size": request.POST.get("box_replace_size"),
        "box_replace_qty": request.POST.get("box_replace_qty"),
        "box_replace_supply": request.POST.get("box_replace_supply"),
        "shrink_wrap_needed": request.POST.get("shrink_wrap_needed"),
        "shrink_wrap_type": request.POST.get("shrink_wrap_type"),
        "shrink_wrap_size": request.POST.get("shrink_wrap_size"),
        "shrink_wrap_qty": request.POST.get("shrink_wrap_qty"),
        "shrink_wrap_supply": request.POST.get("shrink_wrap_supply"),
        "unboxing_rows": unboxing_rows,
        "wholesale_places_qty": request.POST.get("wholesale_places_qty"),
        "invoice_no": request.POST.get("invoice_no"),
        "invoice_date": request.POST.get("invoice_date"),
        "payment_date": request.POST.get("payment_date"),
        "accountant": request.POST.get("accountant"),
        "archive_date": request.POST.get("archive_date"),
        "executor_name": request.POST.get("executor_name"),
        "start_date": request.POST.get("start_date"),
        "end_date": request.POST.get("end_date"),
        "receive_date": request.POST.get("receive_date"),
        "responsible_name": request.POST.get("responsible_name"),
        "direction_needed": request.POST.get("direction_needed"),
        "direction_count": request.POST.get("direction_count"),
        "direction_addresses_json": request.POST.get("direction_addresses_json"),
        "direction_plan_json": request.POST.get("direction_plan_json"),
        "box_forming": request.POST.get("box_forming"),
        "box_forming_other": request.POST.get("box_forming_other"),
        "comments": request.POST.get("comments"),
    }
    if cards_payload:
        payload["cards"] = cards_payload
    photo_file = request.FILES.get("product_photo")
    if photo_file and getattr(photo_file, "name", ""):
        payload["product_photo"] = photo_file.name
    direction_file = request.FILES.get("direction_file")
    if direction_file and getattr(direction_file, "name", ""):
        payload["direction_file"] = direction_file.name
    if is_draft:
        status_value = "draft"
        status_label = "Черновик"
        payload["submit_action"] = "draft"
    else:
        status_value = "sent_unconfirmed"
        status_label = "Ждет подтверждения"
        payload["submit_action"] = "submitted"
    if edit_order_id and preserved_status:
        status_value = preserved_status
        if preserved_label:
            status_label = preserved_label
        if preserved_submit_action:
            payload["submit_action"] = preserved_submit_action
    payload["status"] = status_value
    payload["status_label"] = status_label
    if edit_order_id:
        order_id = edit_order_id
        action = "update"
        description = f"Исправление заявки на обработку №{order_id}"
    elif is_draft:
        if not draft_order_id:
            draft_order_id = f"draft-{uuid.uuid4().hex[:12]}"
        order_id = draft_order_id
        action = "update" if existing_draft_entries else "create"
        description = "Черновик заявки на обработку"
    else:
        order_id = _next_order_number(order_type="processing")
        action = "create"
        description = f"Заявка на обработку №{order_id}"
    marking_each = _non_empty_text(payload.get("marking_5840_each_qty"))
    cz_file = request.FILES.get("marking_cz_file")
    if cz_file and getattr(cz_file, "name", ""):
        payload["marking_cz_file"] = cz_file.name
    if marking_each:
        required_map, required_total, missing_barcodes = _marking_required_by_barcode(payload)
        if not is_draft and required_total <= 0:
            error_response = autosave_error("Добавьте товары для проверки ЧЗ.")
            if error_response:
                return error_response
            return ProcessingHomeView().get(request, error="Добавьте товары для проверки ЧЗ.")
        if not is_draft and missing_barcodes > 0:
            error_response = autosave_error("Для маркировки ЧЗ заполните штрихкоды товара.")
            if error_response:
                return error_response
            return ProcessingHomeView().get(request, error="Для маркировки ЧЗ заполните штрихкоды товара.")
        if cz_file and getattr(cz_file, "name", ""):
            ok, import_result = _import_marking_codes(cz_file, payload, order_id, agency, request.user)
            if not ok:
                message = import_result.get("error") or "Ошибка импорта ЧЗ."
                error_response = autosave_error(message)
                if error_response:
                    return error_response
                return ProcessingHomeView().get(request, error=message)
            payload["marking_cz_import"] = import_result
        if not is_draft:
            available_map = _marking_available_by_barcode(agency, order_id)
            missing_total = 0
            for barcode, required_qty in required_map.items():
                available_qty = available_map.get(barcode, 0)
                missing_total += max(required_qty - available_qty, 0)
            if missing_total > 0:
                message = f"Не хватает ЧЗ: {missing_total}. Загрузите файл с ЧЗ."
                error_response = autosave_error(message)
                if error_response:
                    return error_response
                return ProcessingHomeView().get(request, error=message)
            ok, reserve_error = _reserve_marking_codes(agency, order_id, required_map)
            if not ok:
                message = reserve_error or "Не удалось забронировать ЧЗ."
                error_response = autosave_error(message)
                if error_response:
                    return error_response
                return ProcessingHomeView().get(request, error=message)
    log_order_action(
        action,
        order_id=order_id,
        order_type="processing",
        user=request.user if request.user.is_authenticated else None,
        agency=agency,
        description=description,
        payload=payload,
    )
    if not is_draft:
        _replace_processing_reserves(order_id, agency, stock_rows)
    if not is_draft and draft_order_id and not edit_order_id:
        OrderAuditEntry.objects.filter(
            order_id=draft_order_id,
            order_type="processing",
            agency=agency,
        ).delete()
    if edit_order_id:
        if autosave:
            return JsonResponse(
                {
                    "ok": True,
                    "order_id": order_id,
                    "draft_order_id": "",
                    "status": status_value,
                    "status_label": status_label,
                }
            )
        return redirect(f"/orders/processing/{order_id}/")
    if not is_draft and client_agency:
        _create_processing_manager_task(order_id, agency, request, timezone.localtime())
    if autosave:
        return JsonResponse(
            {
                "ok": True,
                "order_id": order_id,
                "draft_order_id": order_id if is_draft else "",
                "status": status_value,
                "status_label": status_label,
            }
        )
    return redirect(
        f"/orders/processing/?client={agency.id}&ok=1&status={status_value}&order={order_id}"
    )


class ProcessingHomeView(RoleRequiredMixin, TemplateView):
    template_name = "processing/processing.html"
    allowed_roles = ("manager", "storekeeper", "head_manager", "director", "admin")

    def _use_manager_template(self) -> bool:
        if getattr(self.request, "_client_agency", None):
            return False
        role = get_request_role(self.request)
        if role not in {"manager", "head_manager", "director", "admin"}:
            return False
        edit_flag = (self.request.GET.get("edit") or "").strip().lower()
        return edit_flag in {"1", "true", "yes"}

    def get_template_names(self):
        if self._use_manager_template():
            return ["processing/processing_manager.html"]
        return [self.template_name]

    def dispatch(self, request, *args, **kwargs):
        client_agency = _client_agency_from_request(request)
        if client_agency:
            request._client_agency = client_agency
            return TemplateView.dispatch(self, request, *args, **kwargs)
        return super().dispatch(request, *args, **kwargs)

    def get(self, request, *args, **kwargs):
        status = (request.GET.get("status") or "").lower()
        ok = request.GET.get("ok") == "1"
        submitted = kwargs.get("submitted") or (ok and status != "draft")
        draft_saved = ok and status == "draft"
        error = kwargs.get("error")
        order_id = request.GET.get("order")
        if order_id and not getattr(request, "_client_agency", None):
            latest_entry = (
                OrderAuditEntry.objects.filter(order_id=order_id, order_type="processing")
                .order_by("-created_at")
                .first()
            )
            if latest_entry and _is_draft_payload(latest_entry.payload or {}):
                return HttpResponseForbidden("Доступ запрещен")
        ctx = self.get_context_data(
            submitted=submitted,
            draft_saved=draft_saved,
            error=error,
            **kwargs,
        )
        return self.render_to_response(ctx)

    def post(self, request, *args, **kwargs):
        return _submit_processing(request)

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        ctx["submitted"] = kwargs.get("submitted", False)
        ctx["draft_saved"] = kwargs.get("draft_saved", False)
        ctx["error"] = kwargs.get("error")
        ctx["cabinet_url"] = resolve_cabinet_url(get_request_role(self.request))
        ctx["can_assign_packaging"] = get_request_role(self.request) in {
            "processing_head",
            "head_manager",
            "director",
            "admin",
        }

        def resolve_status_label(status_value: str | None, fallback_label: str | None = None) -> str:
            if fallback_label:
                return fallback_label
            value = (status_value or "").strip().lower()
            if value == "draft":
                return "Черновик"
            if value in {"sent_unconfirmed", "send", "submitted"}:
                return "Ждет подтверждения"
            return "Подготовка заявки" if not value else value

        status = self.request.GET.get("status")
        status_label = resolve_status_label(status)
        ctx["order_number"] = self.request.GET.get("order", "")

        client_id = self.request.GET.get("client")
        agency_id = self.request.GET.get("agency")
        agency_key = client_id or agency_id
        client_agency = getattr(self.request, "_client_agency", None) or _client_agency_from_request(self.request)
        agency = client_agency or (Agency.objects.filter(pk=agency_key).first() if agency_key else None)
        ctx["agency"] = agency
        ctx["client_view"] = bool(client_agency)
        ctx["draft_order_id"] = ""
        ctx["edit_order_id"] = ""
        order_id = self.request.GET.get("order")
        draft_payload = None
        role = get_request_role(self.request)
        edit_flag = (self.request.GET.get("edit") or "").strip().lower()
        edit_mode = edit_flag in {"1", "true", "yes"} and role in {
            "manager",
            "head_manager",
            "director",
            "admin",
        }
        if order_id:
            draft_entries = OrderAuditEntry.objects.filter(order_id=order_id, order_type="processing")
            if client_agency:
                draft_entries = draft_entries.filter(agency=client_agency)
            elif agency:
                draft_entries = draft_entries.filter(agency=agency)
            draft_entry = draft_entries.order_by("-created_at").first()
            if draft_entry:
                payload = draft_entry.payload or {}
                status_value = (payload.get("submit_action") or payload.get("status") or "").lower()
                status_label = (payload.get("status_label") or "").lower()
                is_draft = status_value == "draft" or "черновик" in status_label
                if is_draft:
                    ctx["draft_order_id"] = order_id
                    draft_payload = payload
                elif edit_mode:
                    ctx["edit_order_id"] = order_id
                    draft_payload = payload
        ctx["draft_payload"] = draft_payload or {}
        ctx["draft_payload_json"] = json.dumps(draft_payload or {}, ensure_ascii=True)
        if draft_payload:
            payload_status = draft_payload.get("status") or draft_payload.get("submit_action")
            payload_label = (draft_payload.get("status_label") or "").strip()
            status_label = resolve_status_label(payload_status, payload_label)
        ctx["status_label"] = status_label
        return ctx


class ProcessingDirectionsView(RoleRequiredMixin, TemplateView):
    template_name = "processing/processing_directions.html"
    allowed_roles = ("manager", "storekeeper", "head_manager", "director", "admin", "processing_head")

    def dispatch(self, request, *args, **kwargs):
        client_agency = _client_agency_from_request(request)
        if client_agency:
            request._client_agency = client_agency
            return TemplateView.dispatch(self, request, *args, **kwargs)
        return super().dispatch(request, *args, **kwargs)

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        ctx["cabinet_url"] = resolve_cabinet_url(get_request_role(self.request))
        return_url = (self.request.GET.get("return") or "").strip()
        ctx["return_url"] = return_url
        ctx["return_url_json"] = json.dumps(return_url, ensure_ascii=True)
        return ctx


class ProcessingStockPickerView(RoleRequiredMixin, TemplateView):
    template_name = "processing/stock_picker.html"
    allowed_roles = ("manager", "storekeeper", "head_manager", "director", "admin")

    def dispatch(self, request, *args, **kwargs):
        client_agency = _client_agency_from_request(request)
        if client_agency:
            request._client_agency = client_agency
            return TemplateView.dispatch(self, request, *args, **kwargs)
        return super().dispatch(request, *args, **kwargs)

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        ctx["cabinet_url"] = resolve_cabinet_url(get_request_role(self.request))

        client_id = self.request.GET.get("client")
        agency_id = self.request.GET.get("agency")
        agency_key = client_id or agency_id
        client_agency = getattr(self.request, "_client_agency", None) or _client_agency_from_request(self.request)
        agency = client_agency or (Agency.objects.filter(pk=agency_key).first() if agency_key else None)
        ctx["agency"] = agency
        ctx["client_view"] = bool(client_agency)
        exclude_order_id = (self.request.GET.get("order") or "").strip() or None
        ctx["inventory_items_json"] = json.dumps(
            _inventory_items_for_agency(agency, exclude_order_id=exclude_order_id),
            ensure_ascii=True,
        )
        ctx["return_url"] = f"/orders/processing/?client={agency.id}" if agency else "/orders/processing/"
        return ctx


def delete_processing_draft(request, order_id: str):
    if request.method != "POST":
        return HttpResponseForbidden("Доступ запрещен")
    if not request.user.is_authenticated:
        return HttpResponseForbidden("Доступ запрещен")
    client_agency = _client_agency_from_request(request)
    if not client_agency:
        return HttpResponseForbidden("Доступ запрещен")
    entries = OrderAuditEntry.objects.filter(
        order_id=order_id,
        order_type="processing",
        agency=client_agency,
    ).order_by("-created_at")
    latest = entries.first()
    if not latest or not _is_draft_payload(latest.payload or {}):
        return HttpResponseForbidden("Доступ запрещен")
    MarkingCode.objects.filter(
        order_type="processing",
        order_id=order_id,
        agency=client_agency,
        used_at__isnull=True,
    ).update(order_id="")
    entries.delete()
    return redirect(f"/client/dashboard/?client={client_agency.id}")


class ProcessingWorkView(RoleRequiredMixin, TemplateView):
    template_name = "processing/processing_work.html"
    allowed_roles = ("storekeeper", "processing_head", "head_manager", "director", "admin")

    def dispatch(self, request, *args, **kwargs):
        order_id = kwargs.get("order_id")
        if not order_id:
            return redirect("/orders/")
        if _client_agency_from_request(request):
            return HttpResponseForbidden("Доступ запрещен")
        entries = list(
            OrderAuditEntry.objects.filter(order_id=order_id, order_type="processing")
            .select_related("agency")
            .order_by("created_at")
        )
        if not entries:
            return redirect("/orders/")
        latest = entries[-1]
        status_entry = _current_status_entry(entries) or latest
        payload = status_entry.payload or {}
        status_value = (payload.get("status") or payload.get("submit_action") or "").lower()
        status_label = (payload.get("status_label") or "").lower()
        is_ready = (
            status_value in {"processing_head", "processing_in_work"}
            or ("передан" in status_label and "обработ" in status_label)
            or "взята" in status_label
        )
        if not is_ready:
            return HttpResponseForbidden("Доступ запрещен")
        request._processing_work_payload = payload
        request._processing_work_agency = latest.agency
        return super().dispatch(request, *args, **kwargs)

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        order_id = kwargs.get("order_id")
        payload = getattr(self.request, "_processing_work_payload", None) or {}
        agency = getattr(self.request, "_processing_work_agency", None)
        status_value = (payload.get("status") or payload.get("submit_action") or "").strip().lower()
        status_label = (payload.get("status_label") or "").strip()
        if not status_label:
            if status_value == "processing_in_work":
                status_label = "Взята в работу"
            elif status_value == "processing_head":
                status_label = "Передано в обработку"
        ctx["submitted"] = False
        ctx["draft_saved"] = False
        ctx["error"] = kwargs.get("error") or self.request.GET.get("error")
        ctx["cabinet_url"] = resolve_cabinet_url(get_request_role(self.request))
        ctx["order_number"] = order_id or ""
        ctx["agency"] = agency
        ctx["client_view"] = False
        ctx["draft_order_id"] = ""
        ctx["edit_order_id"] = ""
        ctx["draft_payload"] = payload
        ctx["draft_payload_json"] = json.dumps(payload or {}, ensure_ascii=True)
        ctx["status_label"] = status_label or "Обработка товара"
        processed_cards, placed_cards = _processing_card_sets(payload)
        ready_cards = processed_cards - placed_cards
        results = payload.get("processing_results") or []
        has_ready_results = False
        if isinstance(results, list):
            for row in results:
                if not isinstance(row, dict):
                    continue
                processed_qty = _parse_qty_value(row.get("processed")) or 0
                shipped_qty = _parse_qty_value(row.get("shipped_qty")) or 0
                if processed_qty - shipped_qty > 0:
                    has_ready_results = True
                    break
        ctx["can_place_processed"] = bool(ready_cards) or has_ready_results
        ctx["processed_cards_count"] = len(processed_cards)
        ctx["placed_cards_count"] = len(placed_cards)
        ctx["ready_cards_count"] = len(ready_cards)
        marking_items = extract_processing_items(payload)
        counts = (
            MarkingCode.objects.filter(order_type="processing", order_id=order_id, used_at__isnull=False)
            .values("sku_code", "size")
            .annotate(count=Count("id"))
        )
        counts_map = {
            (item["sku_code"], item["size"] or ""): item["count"] for item in counts
        }
        total_count = 0
        for item in marking_items:
            key = (item["sku_code"], item["size"] or "")
            scanned = counts_map.get(key, 0)
            expected = item.get("qty") or 0
            item["scanned"] = scanned
            item["remaining"] = max(expected - scanned, 0)
            total_count += scanned
        ctx["marking_items"] = marking_items
        ctx["marking_total_count"] = total_count
        ctx["marking_api_base"] = f"/marking/processing/{order_id}/"
        ctx["can_assign_packaging"] = get_request_role(self.request) in {
            "processing_head",
            "head_manager",
            "director",
            "admin",
        }
        ctx["processing_workers"] = list(
            Employee.objects.filter(role="processing_worker", is_active=True).order_by("full_name")
        )
        ctx["packaging_tasks"] = list(
            Task.objects.select_related("assigned_to")
            .filter(route=f"/orders/processing/{order_id}/flow/", assigned_to__role="processing_worker")
            .exclude(status="done")
            .order_by("-created_at")
        )
        ctx["assign_status"] = self.request.GET.get("assign")
        ctx["assign_error"] = self.request.GET.get("assign_error")
        return ctx

    def post(self, request, *args, **kwargs):
        action = (request.POST.get("action") or "").strip().lower()
        if action != "finish_processing":
            return HttpResponseForbidden("Доступ запрещен")
        role = get_request_role(request)
        if role not in {"storekeeper", "processing_head"}:
            return HttpResponseForbidden("Доступ запрещен")
        order_id = kwargs.get("order_id")
        if not order_id:
            return redirect("/orders/")
        entries = list(
            OrderAuditEntry.objects.filter(order_id=order_id, order_type="processing")
            .select_related("agency")
            .order_by("created_at")
        )
        if not entries:
            return redirect("/orders/")
        latest = entries[-1]
        payload = dict(latest.payload or {})
        status_value = (payload.get("status") or payload.get("submit_action") or "").lower()
        status_label = (payload.get("status_label") or "").lower()
        if status_value in {"done", "completed", "closed", "finished"} or "выполн" in status_label:
            return redirect(resolve_cabinet_url(role))
        if _order_has_open_boxes(order_id):
            return self.get(
                request,
                error="Есть открытые короба с товаром. Закройте все короба перед завершением заявки.",
                order_id=order_id,
            )

        expected_results = _expected_processing_results(payload)
        if expected_results:
            results = payload.get("processing_results") or []
            results_map = {}
            if isinstance(results, list):
                for item in results:
                    if not isinstance(item, dict):
                        continue
                    article_key = str(item.get("article") or "").strip().lower()
                    size_key = str(item.get("size") or "").strip().lower()
                    dest_key = str(item.get("destination") or "").strip().lower()
                    results_map[(article_key, size_key, dest_key)] = item
            required_fields = ("processed", "defect", "shortage", "shipped_qty")
            for key in expected_results:
                saved = results_map.get(key)
                if not saved:
                    return self.get(
                        request,
                        error="Заполните результаты обработки перед закрытием заявки.",
                        order_id=order_id,
                        )
                for field in required_fields:
                    if _parse_qty_value(saved.get(field)) is None:
                        return self.get(
                            request,
                            error="Заполните результаты обработки перед закрытием заявки.",
                            order_id=order_id,
                        )
        placement_entry = next(
            (
                entry
                for entry in reversed(entries)
                if (entry.payload or {}).get("act") == "placement"
            ),
            None,
        )
        if not placement_entry or (
            (placement_entry.payload or {}).get("act_state") or "closed"
        ).lower() != "closed":
            return redirect(f"/orders/processing/{order_id}/placement/?error=1")
        MarkingCode.objects.filter(
            order_type="processing",
            order_id=order_id,
            used_at__isnull=True,
        ).update(order_id="", printed_at=None, printed_by=None)
        payload["status"] = "done"
        payload["status_label"] = "Выполнена"
        payload["completed_at"] = timezone.localtime().isoformat()
        log_order_action(
            "status",
            order_id=order_id,
            order_type="processing",
            user=request.user if request.user.is_authenticated else None,
            agency=latest.agency if latest else None,
            description="Обработка завершена",
            payload=payload,
        )
        Task.objects.filter(
            route=f"/orders/processing/{order_id}/",
        ).exclude(status="done").update(status="done")
        return redirect(resolve_cabinet_url(role))


class ProcessingPlacementActView(OrdersProcessingPlacementActView):
    template_name = "processing/processing_placement_act.html"


class ProcessingFlowView(OrdersReceivingFlowView):
    template_name = "processing/processing_flow.html"
    order_type = "processing"
    allowed_roles = ("storekeeper", "processing_head", "processing_worker", "head_manager", "director", "admin", "manager")

    def _normalize_flow_state(self, boxes_data, pallets_data, active_box, active_pallet):
        state = super()._normalize_flow_state(boxes_data, pallets_data, active_box, active_pallet)
        owner_boxes = {}
        for raw in boxes_data or []:
            if not isinstance(raw, dict):
                continue
            code = str(raw.get("code") or "").strip()
            if not code:
                continue
            owner_boxes[code] = {
                "owner_agent_id": raw.get("owner_agent_id") or "",
                "owner_user_id": raw.get("owner_user_id"),
                "owner_user_label": raw.get("owner_user_label") or "",
            }
        owner_pallets = {}
        closed_pallets = {}
        for raw in pallets_data or []:
            if not isinstance(raw, dict):
                continue
            code = str(raw.get("code") or "").strip()
            if not code:
                continue
            owner_pallets[code] = {
                "owner_agent_id": raw.get("owner_agent_id") or "",
                "owner_user_id": raw.get("owner_user_id"),
                "owner_user_label": raw.get("owner_user_label") or "",
            }
            closed_pallets[code] = {
                "closed_by_agent_id": raw.get("closed_by_agent_id") or "",
                "closed_by_user_id": raw.get("closed_by_user_id"),
                "closed_by_user_label": raw.get("closed_by_user_label") or "",
            }
        for box in state.get("boxes") or []:
            owner = owner_boxes.get(box.get("code") or "")
            if not owner:
                continue
            if owner.get("owner_agent_id") and not box.get("owner_agent_id"):
                box["owner_agent_id"] = owner.get("owner_agent_id")
            if owner.get("owner_user_id") and not box.get("owner_user_id"):
                box["owner_user_id"] = owner.get("owner_user_id")
            if owner.get("owner_user_label") and not box.get("owner_user_label"):
                box["owner_user_label"] = owner.get("owner_user_label")
        for pallet in state.get("pallets") or []:
            owner = owner_pallets.get(pallet.get("code") or "")
            if not owner:
                continue
            if owner.get("owner_agent_id") and not pallet.get("owner_agent_id"):
                pallet["owner_agent_id"] = owner.get("owner_agent_id")
            if owner.get("owner_user_id") and not pallet.get("owner_user_id"):
                pallet["owner_user_id"] = owner.get("owner_user_id")
            if owner.get("owner_user_label") and not pallet.get("owner_user_label"):
                pallet["owner_user_label"] = owner.get("owner_user_label")
            closed = closed_pallets.get(pallet.get("code") or "")
            if not closed:
                continue
            if closed.get("closed_by_agent_id") and not pallet.get("closed_by_agent_id"):
                pallet["closed_by_agent_id"] = closed.get("closed_by_agent_id")
            if closed.get("closed_by_user_id") and not pallet.get("closed_by_user_id"):
                pallet["closed_by_user_id"] = closed.get("closed_by_user_id")
            if closed.get("closed_by_user_label") and not pallet.get("closed_by_user_label"):
                pallet["closed_by_user_label"] = closed.get("closed_by_user_label")
        return state

    def _can_start(self, entries):
        if not entries:
            return False
        if _flow_closed_from_entries(entries):
            return False
        latest = entries[-1] if entries else None
        payload = _latest_payload_from_entries(entries)
        processed_cards, placed_cards = _processing_card_sets(payload)
        ready_cards = processed_cards - placed_cards if processed_cards else set()
        items = _processing_receiving_items(
            payload,
            latest.agency_id if latest else None,
            ready_cards if ready_cards else None,
        )
        return bool(items)

    def _mark_in_progress(self, request, order_id: str | None):
        return

    def _save_flow_draft(self, request, order_id, entries):
        role = get_request_role(request)
        if role not in {"storekeeper", "processing_head", "processing_worker", "head_manager", "director", "admin", "manager"}:
            return JsonResponse({"ok": False, "error": "forbidden"}, status=403)
        if _flow_closed_from_entries(entries):
            return JsonResponse({"ok": False, "error": "closed"}, status=400)
        if not self._can_start(entries):
            return JsonResponse({"ok": False, "error": "not_allowed"}, status=400)

        boxes_raw = request.POST.get("boxes_json") or "[]"
        pallets_raw = request.POST.get("pallets_json") or "[]"
        active_box = request.POST.get("active_box") or ""
        active_pallet = request.POST.get("active_pallet") or ""
        try:
            boxes_data = json.loads(boxes_raw)
            pallets_data = json.loads(pallets_raw)
        except json.JSONDecodeError:
            return JsonResponse({"ok": False, "error": "invalid_json"}, status=400)
        if not isinstance(boxes_data, list):
            boxes_data = []
        if not isinstance(pallets_data, list):
            pallets_data = []

        agent_id = (request.POST.get("agent_id") or "").strip()
        if not agent_id:
            return JsonResponse({"ok": False, "error": "missing_agent_id"}, status=400)
        flow_state = self._normalize_flow_state(boxes_data, pallets_data, active_box, active_pallet)
        session = _flow_session_for_request(order_id, agent_id, request, create=True)
        if not session:
            return JsonResponse({"ok": False, "error": "session_not_found"}, status=404)
        session.flow_state = flow_state
        session.last_seen = timezone.localtime()
        session.status = ProcessingFlowSession.STATUS_OPEN
        session.save(update_fields=["flow_state", "last_seen", "status", "updated_at"])
        return JsonResponse({"ok": True, "session_id": session.id})

    def _reopen_flow(self, request, order_id, entries):
        can_finish = (
            request.user.is_authenticated
            and Employee.objects.filter(user=request.user, is_active=True, role="processing_head").exists()
        )
        if not can_finish:
            return HttpResponseForbidden("Доступ запрещен")
        if not _flow_closed_from_entries(entries):
            return redirect(f"/orders/processing/{order_id}/flow/")
        latest = entries[-1] if entries else None
        closed_entry = next(
            (entry for entry in reversed(entries) if (entry.payload or {}).get("flow_closed")),
            None,
        )
        closed_payload = closed_entry.payload if closed_entry else {}
        snapshot = {
            "order_id": order_id,
            "flow_closed_at": closed_payload.get("flow_closed_at"),
        }
        log_staff_overaction(
            "update",
            user=request.user if request.user.is_authenticated else None,
            agency=latest.agency if latest else None,
            description=(
                "Избыточное действие: повторное открытие размещения после обработки "
                f"(заявка {order_id})"
            ),
            snapshot=snapshot,
        )
        log_order_action(
            "update",
            order_id=order_id,
            order_type=self.order_type,
            user=request.user if request.user.is_authenticated else None,
            agency=latest.agency if latest else None,
            description="Повторное открытие размещения после обработки",
            payload={
                "flow_reopened": True,
                "flow_reopened_at": timezone.localtime().isoformat(),
            },
        )
        return redirect(f"/orders/processing/{order_id}/flow/?ok=1")

    def get(self, request, *args, **kwargs):
        order_id = kwargs.get("order_id")
        if get_request_role(request) == "storekeeper":
            self._mark_in_progress(request, order_id)
        entries = self._load_entries(order_id)
        if not entries:
            return redirect("/orders/")
        ok = request.GET.get("ok") == "1"
        error = request.GET.get("error")
        ctx = self.get_context_data(ok=ok, error=error, **kwargs)
        return self.render_to_response(ctx)

    def post(self, request, *args, **kwargs):
        order_id = kwargs.get("order_id")
        entries = self._load_entries(order_id)
        if not entries:
            return redirect("/orders/")
        if request.POST.get("flow_action") == "draft":
            return self._save_flow_draft(request, order_id, entries)
        if request.POST.get("flow_action") == "reopen":
            return self._reopen_flow(request, order_id, entries)
        can_finish = (
            request.user.is_authenticated
            and Employee.objects.filter(user=request.user, is_active=True, role="processing_head").exists()
        )
        if not can_finish:
            return HttpResponseForbidden("Доступ запрещен")
        if _flow_closed_from_entries(entries):
            return redirect(f"/orders/processing/{order_id}/flow/")
        if not self._can_start(entries):
            return redirect(f"/orders/processing/{order_id}/flow/?error=1")

        status_entry = _current_status_entry(entries)
        latest = entries[-1]
        base_payload = dict((status_entry.payload or latest.payload or {}))
        processed_cards, placed_cards = _processing_card_sets(base_payload)
        ready_cards = processed_cards - placed_cards if processed_cards else set()
        act_items = _processing_receiving_items(
            base_payload,
            latest.agency_id,
            ready_cards if ready_cards else None,
        )
        if not act_items:
            return redirect(f"/orders/processing/{order_id}/flow/?error=1")

        boxes_raw = request.POST.get("boxes_json") or "[]"
        pallets_raw = request.POST.get("pallets_json") or "[]"
        try:
            boxes_data = json.loads(boxes_raw)
            pallets_data = json.loads(pallets_raw)
        except json.JSONDecodeError:
            return redirect(f"/orders/processing/{order_id}/flow/?error=1")
        if not isinstance(boxes_data, list):
            boxes_data = []
        if not isinstance(pallets_data, list):
            pallets_data = []

        agent_id = (request.POST.get("agent_id") or "").strip()
        if agent_id:
            session = _flow_session_for_request(order_id, agent_id, request, create=True)
            if session:
                flow_state = self._normalize_flow_state(boxes_data, pallets_data, "", "")
                session.flow_state = flow_state
                session.last_seen = timezone.localtime()
                session.status = ProcessingFlowSession.STATUS_OPEN
                session.save(update_fields=["flow_state", "last_seen", "status", "updated_at"])

        sessions = list(
            ProcessingFlowSession.objects.filter(
                order_id=order_id,
                order_type="processing",
                status=ProcessingFlowSession.STATUS_OPEN,
            )
        )
        if sessions:
            boxes_data, pallets_data = _merge_flow_sessions(sessions)

        if _state_has_open_box_with_items({"boxes": boxes_data}):
            return redirect(f"/orders/processing/{order_id}/flow/?error=open_boxes")

        def normalize_items(raw_items):
            items = []
            for raw in raw_items or []:
                if not isinstance(raw, dict):
                    continue
                qty = _parse_qty_value(raw.get("qty")) or 0
                if qty <= 0:
                    continue
                sku_code = (raw.get("sku_code") or raw.get("sku") or "").strip()
                name = (raw.get("name") or "").strip()
                size = (raw.get("size") or "").strip()
                if not (sku_code or name or size):
                    continue
                items.append(
                    {
                        "sku_code": sku_code,
                        "sku": sku_code,
                        "name": name,
                        "size": size,
                        "qty": qty,
                    }
                )
            return items

        cleaned_boxes = []
        seen_box_codes = set()
        for idx, box in enumerate(boxes_data):
            if not isinstance(box, dict):
                continue
            items = normalize_items(box.get("items") or [])
            if not items:
                continue
            code = str(box.get("code") or "").strip() or f"BOX-{idx + 1}"
            if code in seen_box_codes:
                code = f"{code}-{idx + 1}"
            seen_box_codes.add(code)
            cleaned_boxes.append(
                {
                    "code": code,
                    "items": items,
                    "sealed": True,
                }
            )

        cleaned_pallets = []
        seen_pallet_codes = set()
        for idx, pallet in enumerate(pallets_data):
            if not isinstance(pallet, dict):
                continue
            code = str(pallet.get("code") or "").strip() or f"PALLET-{idx + 1}"
            if code in seen_pallet_codes:
                code = f"{code}-{idx + 1}"
            seen_pallet_codes.add(code)
            boxes = [
                str(box_code).strip()
                for box_code in (pallet.get("boxes") or [])
                if str(box_code or "").strip()
            ]
            boxes = [box_code for box_code in boxes if box_code in seen_box_codes]
            items = normalize_items(pallet.get("items") or [])
            if not boxes and not items:
                continue
            location = pallet.get("location")
            if isinstance(location, dict):
                location = dict(location)
            elif isinstance(location, str):
                location = {"zone": location}
            else:
                location = {}
            location.setdefault("zone", "PR")
            cleaned_pallets.append(
                {
                    "code": code,
                    "boxes": boxes,
                    "items": items,
                    "sealed": True,
                    "location": location,
                }
            )

        if not cleaned_boxes or not cleaned_pallets:
            return redirect(f"/orders/processing/{order_id}/flow/?error=1")

        pallet_box_codes = set()
        for pallet in cleaned_pallets:
            for box_code in pallet.get("boxes") or []:
                if box_code:
                    pallet_box_codes.add(box_code)
        unassigned_boxes = [box for box in cleaned_boxes if box["code"] not in pallet_box_codes]
        if unassigned_boxes:
            return redirect(f"/orders/processing/{order_id}/flow/?error=1")

        totals = {}

        def add_total(item, qty, field):
            key = _item_key(item.get("sku_code") or item.get("sku"), item.get("name"), item.get("size"))
            if not key:
                return
            entry = totals.setdefault(key, {"box": 0, "pallet": 0, "total": 0})
            entry[field] += qty
            entry["total"] += qty

        for box in cleaned_boxes:
            for item in box.get("items") or []:
                qty = _parse_qty_value(item.get("qty")) or 0
                add_total(item, qty, "box")
        for pallet in cleaned_pallets:
            for item in pallet.get("items") or []:
                qty = _parse_qty_value(item.get("qty")) or 0
                add_total(item, qty, "pallet")

        placement_items = []
        for item in act_items:
            key = _item_key(item.get("sku_code"), item.get("name"), item.get("size"))
            entry = totals.get(key, {"box": 0, "pallet": 0, "total": 0})
            actual_qty = _parse_qty_value(item.get("actual_qty")) or 0
            if entry["total"] > actual_qty or entry["total"] != actual_qty:
                return redirect(f"/orders/processing/{order_id}/flow/?error=1")
            placement_items.append(
                {
                    "sku_code": item.get("sku_code"),
                    "name": item.get("name"),
                    "size": item.get("size"),
                    "actual_qty": actual_qty,
                    "box_qty": entry["box"],
                    "pallet_qty": entry["pallet"],
                }
            )

        placed_cards_list = base_payload.get("placed_cards") or []
        if not isinstance(placed_cards_list, list):
            placed_cards_list = []
        for card_id in ready_cards:
            if card_id and card_id not in placed_cards_list:
                placed_cards_list.append(card_id)
        base_payload["placed_cards"] = placed_cards_list
        now = timezone.localtime().isoformat()
        for card in base_payload.get("cards") or []:
            card_id = _processing_card_id(card)
            if card_id and card_id in ready_cards:
                card["placed_at"] = card.get("placed_at") or now
                card["placed_done"] = True

        act_payload = dict(base_payload)
        act_payload["act"] = "placement"
        act_payload["act_label"] = "Акт размещения"
        act_payload["act_state"] = "closed"
        act_payload["act_items"] = placement_items
        act_payload["act_boxes"] = cleaned_boxes
        act_payload["act_pallets"] = cleaned_pallets
        act_payload["flow_state"] = self._normalize_flow_state(boxes_data, pallets_data, "", "")
        act_payload["flow_closed"] = True
        act_payload["flow_closed_at"] = timezone.localtime().isoformat()
        log_order_action(
            "update",
            order_id=order_id,
            order_type=self.order_type,
            user=request.user if request.user.is_authenticated else None,
            agency=latest.agency,
            description="Создан акт размещения после обработки (поток)",
            payload=act_payload,
        )
        ProcessingFlowSession.objects.filter(
            order_id=order_id,
            order_type="processing",
            status=ProcessingFlowSession.STATUS_OPEN,
        ).update(
            status=ProcessingFlowSession.STATUS_CLOSED,
            last_seen=timezone.localtime(),
        )
        return redirect(f"/orders/processing/{order_id}/flow/?ok=1")

    def get_context_data(self, **kwargs):
        ctx = TemplateView.get_context_data(self, **kwargs)
        order_id = kwargs.get("order_id")
        ctx["can_finish_flow"] = (
            self.request.user.is_authenticated
            and Employee.objects.filter(user=self.request.user, is_active=True, role="processing_head").exists()
        )
        entries = self._load_entries(order_id)
        latest = entries[-1] if entries else None
        status_entry = _current_status_entry(entries)
        payload = _latest_payload_from_entries(entries)
        processed_cards, placed_cards = _processing_card_sets(payload)
        ready_cards = processed_cards - placed_cards if processed_cards else set()
        items = _processing_receiving_items(
            payload,
            latest.agency_id if latest else None,
            ready_cards if ready_cards else None,
        )
        display_items = []
        for item in items:
            display_items.append(
                {
                    "sku_code": item.get("sku_code") or "",
                    "name": item.get("name") or "",
                    "size": item.get("size") or "",
                    "qty": _parse_qty_value(item.get("actual_qty")) or 0,
                    "comment": "",
                }
            )
        client_label = "-"
        client_prefix = ""
        if latest and latest.agency:
            name = latest.agency.agn_name or latest.agency.fio_agn or str(latest.agency)
            client_label = _shorten_ip_name(name)
            client_prefix = (latest.agency.pref or "").strip()
        status_payload = status_entry.payload or {} if status_entry else {}
        goods_type = (status_payload.get("goods_type") or payload.get("goods_type") or "").strip().lower()
        cz_required = bool(_parse_qty_value(payload.get("marking_5840_each_qty")))
        catalog_items = []
        barcode_map = {}
        if latest and latest.agency_id:
            for sku in SKU.objects.filter(agency_id=latest.agency_id, deleted=False).only(
                "sku_code",
                "name",
                "size",
            ):
                catalog_items.append(
                    {
                        "sku_code": sku.sku_code,
                        "name": sku.name,
                        "size": sku.size,
                    }
                )
            for barcode in (
                SKUBarcode.objects.select_related("sku")
                .filter(sku__agency_id=latest.agency_id, sku__deleted=False)
            ):
                value = (barcode.value or "").strip()
                if not value:
                    continue
                sku = barcode.sku
                barcode_map[value] = {
                    "sku_code": sku.sku_code,
                    "name": sku.name,
                    "size": (barcode.size or sku.size or "").strip(),
                }
        flow_locked = _flow_closed_from_entries(entries)
        flow_state: dict = {}
        flow_state_shared: dict = {"boxes": [], "pallets": []}
        if flow_locked:
            flow_state = self._find_flow_state(entries) or {}
            if not flow_state or not (flow_state.get("boxes") or flow_state.get("pallets")):
                placement_entry = self._placement_act_entry(entries)
                if placement_entry:
                    placement_payload = placement_entry.payload or {}
                    placement_boxes = placement_payload.get("act_boxes") or []
                    placement_pallets = placement_payload.get("act_pallets") or []
                    if placement_boxes or placement_pallets:
                        active_box = next(
                            (box.get("code") for box in placement_boxes if not box.get("sealed")),
                            "",
                        )
                        active_pallet = next(
                            (pallet.get("code") for pallet in placement_pallets if not pallet.get("sealed")),
                            "",
                        )
                        flow_state = {
                            "boxes": placement_boxes,
                            "pallets": placement_pallets,
                            "activeBox": active_box,
                            "activePallet": active_pallet,
                        }
        else:
            agent_id = (self.request.GET.get("agent_id") or "").strip()
            current_session = None
            if agent_id:
                current_session = _flow_session_for_request(order_id, agent_id, self.request, create=False)
                if current_session and isinstance(current_session.flow_state, dict):
                    flow_state = current_session.flow_state
            shared_sessions = list(
                ProcessingFlowSession.objects.filter(
                    order_id=order_id,
                    order_type="processing",
                    status=ProcessingFlowSession.STATUS_OPEN,
                )
            )
            if agent_id and self.request.user.is_authenticated:
                shared_sessions = [
                    session
                    for session in shared_sessions
                    if not (
                        session.agent_id == agent_id
                        and session.user_id == self.request.user.id
                    )
                ]
            if shared_sessions:
                boxes_data, pallets_data = _merge_flow_sessions(shared_sessions)
                flow_state_shared = {
                    "boxes": boxes_data,
                    "pallets": pallets_data,
                }
        act_print_url = ""
        placement_entry = self._placement_act_entry(entries)
        scanner_agents = []
        scanner_ports = []
        scanner_default = {}
        scanner_ports_info = []
        agent_devices_map: dict[str, list[dict]] = {}
        if placement_entry and flow_locked:
            act_print_url = (
                f"/orders/processing/{order_id}/placement/"
                f"?return=/orders/processing/{order_id}/flow/"
            )
        try:
            online_threshold = timezone.now() - timedelta(seconds=30)
            ports_pool = set()
            ports_map: dict[str, list[dict]] = {}
            devices_map: dict[str, list[dict]] = {}
            for agent in DeviceAgent.objects.all().order_by("-last_seen", "-updated_at"):
                is_online = bool(agent.last_seen and agent.last_seen >= online_threshold)
                meta = agent.meta if isinstance(agent.meta, dict) else {}
                ports = meta.get("com_ports") or meta.get("ports")
                devices = meta.get("com_devices") if isinstance(meta.get("com_devices"), list) else []
                agent_devices_map[agent.agent_id] = devices if isinstance(devices, list) else []
                for device in devices:
                    if not isinstance(device, dict):
                        continue
                    port_value = str(device.get("port") or "").strip()
                    if not port_value:
                        continue
                    ports_pool.add(port_value)
                    devices_map.setdefault(port_value.lower(), []).append(device)
                com_config = meta.get("com") if isinstance(meta.get("com"), dict) else {}
                com_status = meta.get("com_status") if isinstance(meta.get("com_status"), dict) else {}
                if isinstance(ports, str):
                    ports = [ports]
                if not isinstance(ports, list):
                    ports = []
                ports = [str(port).strip() for port in ports if str(port).strip()]
                seen_ports = set()
                unique_ports = []
                for port in ports:
                    key = port.lower()
                    if key in seen_ports:
                        continue
                    seen_ports.add(key)
                    unique_ports.append(port)
                ports = unique_ports
                for port in ports:
                    ports_pool.add(port)
                com_port = str(
                    com_status.get("port")
                    or com_config.get("port")
                    or com_config.get("port_name")
                    or ""
                ).strip()
                com_enabled = bool(
                    com_status.get("enabled")
                    if "enabled" in com_status
                    else com_config.get("enabled")
                )
                com_connected = bool(com_status.get("connected"))
                com_error = str(com_status.get("error") or "").strip()
                com_error_type = str(com_status.get("error_type") or "").strip()
                com_error_code = com_status.get("error_code")
                com_last_scan = str(com_status.get("last_scan_at") or "").strip()
                com_last_chunk = str(com_status.get("last_chunk_at") or "").strip()
                com_baud = com_status.get("baud") or com_config.get("baud")
                com_eol = com_status.get("eol") or com_config.get("eol")
                com_idle = com_status.get("idle_ms") or com_config.get("idle_ms")
                scanner_agents.append(
                    {
                        "agent_id": agent.agent_id,
                        "title": agent.name or agent.host or agent.agent_id,
                        "status": "онлайн" if is_online else "нет связи",
                        "is_online": is_online,
                        "host": agent.host,
                        "version": agent.version,
                        "last_seen": agent.last_seen.isoformat() if agent.last_seen else "",
                        "ports": ports,
                        "com_port": com_port,
                        "com_enabled": com_enabled,
                        "com_connected": com_connected,
                        "com_error": com_error,
                        "com_error_type": com_error_type,
                        "com_error_code": com_error_code,
                        "com_last_scan": com_last_scan,
                        "com_last_chunk": com_last_chunk,
                        "com_baud": com_baud,
                        "com_eol": com_eol,
                        "com_idle": com_idle,
                    }
                )
                if com_port:
                    key = com_port.lower()
                    ports_map.setdefault(key, []).append(scanner_agents[-1])
            scanner_ports = sorted(ports_pool)
            settings_payload = load_scanner_settings()
            if isinstance(settings_payload, dict):
                scanner_default = settings_payload.get("default") if isinstance(settings_payload.get("default"), dict) else {}
            scanner_ports_info = []
            for port in scanner_ports:
                key = port.lower()
                owners = ports_map.get(key) or []
                device_items = devices_map.get(key) or []
                if not owners:
                    if device_items:
                        device_lines = []
                        for device in device_items:
                            name = str(device.get("name") or device.get("caption") or device.get("description") or "").strip()
                            status = str(device.get("status") or "").strip()
                            error_code = device.get("error_code")
                            service = str(device.get("service") or "").strip()
                            line_parts = [name] if name else []
                            if status:
                                line_parts.append(status)
                            if error_code is not None and error_code != "":
                                line_parts.append(f"код {error_code}")
                            if service:
                                line_parts.append(f"drv {service}")
                            device_lines.append(" · ".join(line_parts) if line_parts else "устройство")
                        scanner_ports_info.append(
                            {
                                "port": port,
                                "status": " | ".join(device_lines),
                            }
                        )
                    else:
                        scanner_ports_info.append({"port": port, "status": "свободен (по данным агентов)"})
                    continue
                parts = []
                if device_items:
                    for device in device_items:
                        name = str(device.get("name") or device.get("caption") or device.get("description") or "").strip()
                        status = str(device.get("status") or "").strip()
                        error_code = device.get("error_code")
                        service = str(device.get("service") or "").strip()
                        line_parts = [name] if name else []
                        if status:
                            line_parts.append(status)
                        if error_code is not None and error_code != "":
                            line_parts.append(f"код {error_code}")
                        if service:
                            line_parts.append(f"drv {service}")
                        if line_parts:
                            parts.append(" · ".join(line_parts))
                for owner in owners:
                    title = owner.get("title") or owner.get("agent_id") or "агент"
                    state = "нет связи"
                    if owner.get("is_online"):
                        if owner.get("com_connected"):
                            state = "подключен"
                        elif owner.get("com_enabled"):
                            state = "включен"
                        else:
                            state = "выключен"
                    label = f"{title} · {state}"
                    if owner.get("com_error"):
                        label += f" · ошибка: {owner.get('com_error')}"
                        if owner.get("com_error_type"):
                            label += f" ({owner.get('com_error_type')})"
                        if owner.get("com_error_code") is not None:
                            label += f" [{owner.get('com_error_code')}]"
                    parts.append(label)
                scanner_ports_info.append({"port": port, "status": "; ".join(parts)})
        except DatabaseError:
            scanner_agents = []
            scanner_ports = []
            scanner_default = {}
            scanner_ports_info = []
            agent_devices_map = {}
        ctx.update(
            {
                "order_id": order_id,
                "client_label": client_label,
                "client_prefix": client_prefix,
                "status_label": _status_label_from_entry(status_entry) if status_entry else "-",
                "cabinet_url": resolve_cabinet_url(get_request_role(self.request)),
                "goods_type": goods_type,
                "goods_type_label": GOODS_TYPE_LABELS.get(goods_type, ""),
                "cz_required": cz_required,
                "items": display_items,
                "barcode_map": barcode_map,
                "catalog_items": catalog_items,
                "flow_state": flow_state,
                "flow_state_shared": flow_state_shared,
                "flow_locked": flow_locked,
                "act_print_url": act_print_url,
                "scanner_agents": scanner_agents,
                "scanner_ports": scanner_ports,
                "scanner_ports_info": scanner_ports_info,
                "scanner_agent_devices_json": json.dumps(agent_devices_map, ensure_ascii=False),
                "scanner_default": scanner_default,
                "scanner_eols": SCANNER_EOLS,
                "ok": kwargs.get("ok", False),
                "error": kwargs.get("error"),
            }
        )
        return ctx


@login_required
@require_GET
def processing_flow_session(request, order_id: str):
    role = get_request_role(request)
    if role not in {"storekeeper", "processing_head", "processing_worker", "head_manager", "director", "admin", "manager"}:
        return HttpResponseForbidden("Доступ запрещен")
    agent_id = (request.GET.get("agent_id") or "").strip()
    if not agent_id:
        return JsonResponse({"ok": False, "error": "missing_agent_id"}, status=400)
    session = _flow_session_for_request(order_id, agent_id, request, create=True)
    if not session:
        return JsonResponse({"ok": False, "error": "session_not_found"}, status=404)
    flow_state = session.flow_state if isinstance(session.flow_state, dict) else _default_flow_state()
    return JsonResponse(
        {
            "ok": True,
            "session_id": session.id,
            "flow_state": flow_state,
            "updated_at": session.updated_at.isoformat() if session.updated_at else "",
        }
    )


@login_required
@require_GET
def processing_flow_shared(request, order_id: str):
    role = get_request_role(request)
    if role not in {"storekeeper", "processing_head", "processing_worker", "head_manager", "director", "admin", "manager"}:
        return HttpResponseForbidden("Доступ запрещен")
    agent_id = (request.GET.get("agent_id") or "").strip()
    sessions = list(
        ProcessingFlowSession.objects.filter(
            order_id=order_id,
            order_type="processing",
            status=ProcessingFlowSession.STATUS_OPEN,
        )
    )
    if agent_id and request.user.is_authenticated:
        sessions = [
            session
            for session in sessions
            if not (
                session.agent_id == agent_id
                and session.user_id == request.user.id
            )
        ]
    boxes_data, pallets_data = _merge_flow_sessions(sessions) if sessions else ([], [])
    return JsonResponse(
        {
            "ok": True,
            "flow_state": {
                "boxes": boxes_data,
                "pallets": pallets_data,
            },
        }
    )


def processing_flow_box_action(request, order_id: str):
    if request.method != "POST":
        return JsonResponse({"ok": False, "error": "method_not_allowed"}, status=405)
    if not request.user.is_authenticated:
        return HttpResponseForbidden("Доступ запрещен")
    role = get_request_role(request)
    if role != "storekeeper":
        return HttpResponseForbidden("Доступ запрещен")
    try:
        payload = json.loads(request.body.decode("utf-8") or "{}")
    except json.JSONDecodeError:
        payload = request.POST.dict()
    if not isinstance(payload, dict):
        payload = {}
    action = (payload.get("action") or "").lower()
    if action not in {"edit", "delete"}:
        return JsonResponse({"ok": False, "error": "invalid_action"}, status=400)
    box_code = (payload.get("box_code") or "").strip()
    if not box_code:
        return JsonResponse({"ok": False, "error": "missing_box"}, status=400)
    entry = (
        OrderAuditEntry.objects.filter(order_id=order_id, order_type="processing")
        .select_related("agency")
        .order_by("-created_at")
        .first()
    )
    snapshot = {
        "order_id": order_id,
        "box_code": box_code,
        "pallet_code": (payload.get("pallet_code") or "").strip(),
        "pallet_index": _parse_int_value(payload.get("pallet_index")),
        "box_index": _parse_int_value(payload.get("box_index")),
        "total_qty": _parse_int_value(payload.get("total_qty")),
        "items": payload.get("items") if isinstance(payload.get("items"), list) else [],
        "previous_active_box": (payload.get("previous_active_box") or "").strip(),
        "new_active_box": (payload.get("new_active_box") or "").strip(),
    }
    action_label = "Редактирование короба" if action == "edit" else "Удаление короба"
    description = f"{action_label} {box_code} (заявка {order_id})"
    if snapshot["pallet_index"] and snapshot["box_index"]:
        description += f", палета {snapshot['pallet_index']}, короб {snapshot['box_index']}"
    log_staff_overaction(
        "update" if action == "edit" else "delete",
        user=request.user if request.user.is_authenticated else None,
        agency=entry.agency if entry else None,
        description=description,
        snapshot=snapshot,
    )
    return JsonResponse({"ok": True})


@login_required
@require_POST
def processing_flow_marking_scan(request, order_id: str):
    role = get_request_role(request)
    if role not in {"storekeeper", "processing_head", "processing_worker", "head_manager", "director", "admin", "manager"}:
        return JsonResponse({"ok": False, "error": "Доступ запрещен."}, status=403)
    entries = list(
        OrderAuditEntry.objects.filter(order_id=order_id, order_type="processing")
        .select_related("agency")
        .order_by("created_at")
    )
    if not entries:
        return JsonResponse({"ok": False, "error": "Заявка не найдена."}, status=404)
    latest = entries[-1]
    payload = _latest_payload_from_entries(entries)
    marking_each_qty = _parse_qty_value(payload.get("marking_5840_each_qty")) or 0
    if marking_each_qty <= 0:
        return JsonResponse({"ok": False, "error": "ЧЗ не требуется."}, status=400)
    data = _parse_json_body(request)
    if data is None:
        return JsonResponse({"ok": False, "error": "Некорректный JSON."}, status=400)
    code = _normalize_marking_code(data.get("code") or "")
    box_barcode = (data.get("box_barcode") or "").strip()
    agent_id = (data.get("agent_id") or data.get("agentId") or "").strip()
    if not code:
        return JsonResponse({"ok": False, "error": "Код ЧЗ не указан."}, status=400)
    if not box_barcode:
        return JsonResponse({"ok": False, "error": "Не выбран короб."}, status=400)
    if not agent_id:
        return JsonResponse({"ok": False, "error": "Не указан агент."}, status=400)

    session = _flow_session_for_request(order_id, agent_id, request, create=False)
    if not session:
        return JsonResponse({"ok": False, "error": "Контекст агента не найден."}, status=409)

    processed_cards, placed_cards = _processing_card_sets(payload)
    ready_cards = processed_cards - placed_cards if processed_cards else set()
    items = _processing_receiving_items(
        payload,
        latest.agency_id,
        ready_cards if ready_cards else None,
    )
    if not items:
        return JsonResponse({"ok": False, "error": "Нет товаров для размещения."}, status=400)
    allowed_map = {}
    for item in items:
        sku_value = str(item.get("sku_code") or "").strip()
        size_value = str(item.get("size") or "").strip()
        if not sku_value:
            continue
        allowed_map[(sku_value.lower(), size_value.lower())] = item
    if not allowed_map:
        return JsonResponse({"ok": False, "error": "Нет товаров для размещения."}, status=400)

    now = timezone.localtime()
    code_variants = {code}
    if "\x1d" in code:
        code_variants.add(code.replace("\x1d", "_x001D_"))
    with transaction.atomic():
        existing = (
            MarkingCode.objects.select_for_update()
            .filter(code__in=list(code_variants))
            .first()
        )
        if not existing:
            return JsonResponse({"ok": False, "error": "Код ЧЗ не найден."}, status=404)
        if existing.used_at:
            return JsonResponse({"ok": False, "error": "Код уже использован."}, status=409)
        if latest.agency_id and existing.agency_id and existing.agency_id != latest.agency_id:
            return JsonResponse(
                {"ok": False, "error": "Код принадлежит другому клиенту."},
                status=409,
            )
        if existing.order_type and existing.order_type != "processing":
            return JsonResponse(
                {"ok": False, "error": "Код закреплен в другом процессе."},
                status=409,
            )
        if existing.order_id and existing.order_id != order_id:
            return JsonResponse(
                {"ok": False, "error": "Код закреплен за другой заявкой."},
                status=409,
            )

        sku_code = (existing.sku_code or "").strip()
        size = (existing.size or "").strip()
        if not sku_code:
            return JsonResponse({"ok": False, "error": "У кода нет артикула."}, status=409)
        sku_key = sku_code.lower()
        if not size:
            candidate_sizes = {
                str(item.get("size") or "").strip()
                for item in items
                if str(item.get("sku_code") or "").strip().lower() == sku_key
            }
            if len(candidate_sizes) == 1:
                size = candidate_sizes.pop()
        allowed_key = (sku_key, size.lower())
        if allowed_key not in allowed_map:
            return JsonResponse({"ok": False, "error": "Позиция не найдена в заявке."}, status=409)

        allowed_item = allowed_map[allowed_key]
        allowed_qty = _parse_qty_value(allowed_item.get("actual_qty")) or 0
        used_qty = MarkingCode.objects.filter(
            order_type="processing",
            order_id=order_id,
            sku_code=sku_code,
            size=size,
            used_at__isnull=False,
        ).count()
        if allowed_qty and used_qty >= allowed_qty:
            return JsonResponse({"ok": False, "error": "Количество ЧЗ уже закрыто."}, status=409)
        if existing.box_barcode and existing.box_barcode != box_barcode:
            return JsonResponse({"ok": False, "error": "Код закреплен за другим коробом."}, status=409)

        update_fields = []
        if not existing.order_id:
            existing.order_id = order_id
            update_fields.append("order_id")
        if not existing.order_type:
            existing.order_type = "processing"
            update_fields.append("order_type")
        if not existing.size and size:
            existing.size = size
            update_fields.append("size")
        if box_barcode and not existing.box_barcode:
            existing.box_barcode = box_barcode
            update_fields.append("box_barcode")
        if not existing.sku:
            sku_qs = SKU.objects.filter(sku_code=sku_code)
            if latest.agency_id:
                sku = sku_qs.filter(agency_id=latest.agency_id).first() or sku_qs.filter(
                    agency__isnull=True
                ).first()
            else:
                sku = sku_qs.first()
            if sku:
                existing.sku = sku
                update_fields.append("sku")

        existing.used_at = now
        existing.used_by = request.user if request.user.is_authenticated else None
        update_fields.extend(["used_at", "used_by"])
        existing.save(update_fields=update_fields)

    return JsonResponse(
        {
            "ok": True,
            "sku_code": allowed_item.get("sku_code") or sku_code,
            "size": allowed_item.get("size") or size,
            "name": allowed_item.get("name") or "",
            "box_barcode": box_barcode,
        }
    )


@login_required
@require_POST
def processing_assign_packaging(request, order_id: str):
    role = get_request_role(request)
    if role not in {"processing_head", "head_manager", "director", "admin"}:
        return HttpResponseForbidden("Доступ запрещен")
    employee_id = (request.POST.get("assignee_id") or request.POST.get("worker_id") or "").strip()
    if not employee_id:
        return redirect(f"/orders/processing/{order_id}/work/?assign_error=missing_worker")
    assignee = Employee.objects.filter(
        pk=employee_id,
        role="processing_worker",
        is_active=True,
    ).first()
    if not assignee:
        return redirect(f"/orders/processing/{order_id}/work/?assign_error=invalid_worker")

    route = f"/orders/processing/{order_id}/flow/"
    existing = Task.objects.filter(route=route, assigned_to=assignee).exclude(status="done")
    if existing.exists():
        return redirect(f"/orders/processing/{order_id}/work/?assign_error=already_assigned")

    description = f"Раскоробовка товара. Исполнитель: {assignee.full_name}."
    observer = get_employee_for_user(request.user)
    Task.objects.create(
        title=f"Задача на раскоробовку товара по заявке №{order_id}",
        description=description,
        route=route,
        assigned_to=assignee,
        observer=observer,
        created_by=request.user if request.user.is_authenticated else None,
        due_date=timezone.localtime(),
    )
    latest = (
        OrderAuditEntry.objects.filter(order_id=order_id, order_type="processing")
        .select_related("agency")
        .order_by("-created_at")
        .first()
    )
    log_order_action(
        "update",
        order_id=order_id,
        order_type="processing",
        user=request.user if request.user.is_authenticated else None,
        agency=latest.agency if latest else None,
        description=f"Поручена упаковка в короба: {assignee.full_name}",
        payload={
            "packing_assignee_id": assignee.id,
            "packing_assignee": assignee.full_name,
            "packing_assignee_role": assignee.role,
        },
    )
    return redirect(f"/orders/processing/{order_id}/work/?assign=ok")


class ProcessingCardView(RoleRequiredMixin, TemplateView):
    template_name = "processing/processing_card.html"
    allowed_roles = ("storekeeper", "processing_head", "head_manager", "director", "admin")

    def dispatch(self, request, *args, **kwargs):
        order_id = kwargs.get("order_id")
        if not order_id:
            return redirect("/orders/")
        if _client_agency_from_request(request):
            return HttpResponseForbidden("Доступ запрещен")
        latest = (
            OrderAuditEntry.objects.filter(order_id=order_id, order_type="processing")
            .select_related("agency")
            .order_by("-created_at")
            .first()
        )
        if not latest:
            return redirect("/orders/")
        payload = latest.payload or {}
        status_value = (payload.get("status") or payload.get("submit_action") or "").lower()
        status_label = (payload.get("status_label") or "").lower()
        allowed = (
            status_value in {"processing_head", "processing_in_work", "done", "completed", "closed", "finished"}
            or ("передан" in status_label and "обработ" in status_label)
            or "взята" in status_label
            or "выполн" in status_label
        )
        if not allowed:
            return HttpResponseForbidden("Доступ запрещен")
        request._processing_card_payload = payload
        request._processing_card_agency = latest.agency
        request._processing_card_status_label = status_label or payload.get("status_label")
        return super().dispatch(request, *args, **kwargs)

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        order_id = kwargs.get("order_id") or ""
        payload = getattr(self.request, "_processing_card_payload", None) or {}
        agency = getattr(self.request, "_processing_card_agency", None)
        status_label = (getattr(self.request, "_processing_card_status_label", None) or "").strip()
        if not status_label:
            status_value = (payload.get("status") or payload.get("submit_action") or "").strip().lower()
            if status_value == "processing_in_work":
                status_label = "Взята в работу"
            elif status_value == "processing_head":
                status_label = "Передано в обработку"
            elif status_value in {"done", "completed", "closed", "finished"}:
                status_label = "Выполнена"
        return_url = (self.request.GET.get("return") or "").strip()
        if return_url:
            parsed = urlparse(return_url)
            if parsed.scheme or parsed.netloc:
                return_url = ""
            else:
                return_url = parsed.path or ""
                if parsed.query:
                    return_url = f"{return_url}?{parsed.query}"
                if parsed.fragment:
                    return_url = f"{return_url}#{parsed.fragment}"
        if not return_url:
            return_url = f"/orders/processing/{order_id}/work/"
        ctx["order_id"] = order_id
        ctx["return_url"] = return_url
        ctx["cabinet_url"] = resolve_cabinet_url(get_request_role(self.request))
        ctx["status_label"] = status_label or "-"
        ctx["client_label"] = (
            agency.agn_name or agency.fio_agn or str(agency)
        ) if agency else "-"
        printers, printers_meta = load_available_printers_data()
        ctx["available_printers"] = printers
        ctx["available_printers_meta"] = printers_meta
        ctx["label_settings"] = load_label_settings()
        label_sizes = {}
        for entry in LABEL_SIZES:
            key = entry.get("key")
            if not key:
                continue
            label_sizes[key] = {
                "width_mm": entry.get("width_mm", 58),
                "height_mm": entry.get("height_mm", 40),
                "preview_scale": entry.get("preview_scale", 1),
            }
        ctx["label_sizes_json"] = json.dumps(label_sizes, ensure_ascii=True)

        cards = payload.get("cards") or []
        card_id = kwargs.get("card_id") or ""
        ctx["card_id"] = card_id
        article_param = (self.request.GET.get("article") or "").strip()
        selected_card = None
        for card in cards:
            if not isinstance(card, dict):
                continue
            if card_id and str(card.get("id") or "") == card_id:
                selected_card = card
                break
        if not selected_card and article_param:
            for card in cards:
                if not isinstance(card, dict):
                    continue
                if (card.get("article") or "").strip() == article_param:
                    selected_card = card
                    break
        if not selected_card and cards:
            selected_card = cards[0] if isinstance(cards[0], dict) else None

        fallback_rows = payload.get("stock_rows") or payload.get("size_rows") or []
        selected_card = selected_card or {
            "article": payload.get("article") or "",
            "product_name": payload.get("product_name") or "",
            "photo_url": payload.get("product_photo_url") or "",
            "goods_type": payload.get("goods_type") or "",
            "rows": fallback_rows,
        }

        article_value = (selected_card.get("article") or payload.get("article") or "").strip()
        product_name = (selected_card.get("product_name") or payload.get("product_name") or "").strip()
        goods_type = (selected_card.get("goods_type") or payload.get("goods_type") or "").strip().lower()
        goods_type_label = GOODS_TYPE_LABELS.get(goods_type, goods_type) if goods_type else ""
        photo_url = (selected_card.get("photo_url") or selected_card.get("product_photo_url") or payload.get("product_photo_url") or "").strip()
        supplier_value = str(payload.get("supplier") or "").strip()
        if not supplier_value and agency:
            supplier_value = (agency.agn_name or agency.fio_agn or str(agency) or "").strip()
        product_name_value = _normalize_org_name(product_name)
        brand_value = _normalize_org_name(payload.get("brand") or "")
        subject_value = _normalize_org_name(payload.get("subject") or "")
        color_value = _normalize_org_name(payload.get("color") or "")
        composition_value = _normalize_org_name(payload.get("composition") or "")
        country_value = _normalize_org_name(payload.get("made_in") or "")
        supplier_value = _normalize_org_name(supplier_value)
        label_base = {
            "article": article_value,
            "name": product_name_value,
            "brand": brand_value,
            "subject": subject_value,
            "color": color_value,
            "composition": composition_value,
            "supplier": supplier_value,
            "country": country_value,
            "barcode_extra": (payload.get("source_reference") or article_value or ""),
        }
        ctx["label_base_json"] = json.dumps(label_base, ensure_ascii=True)

        card_fields = []

        def add_field(label: str, value: str | None) -> None:
            text = _non_empty_text(value)
            if not text:
                return
            card_fields.append({"label": label, "value": text})

        add_field("Наименование товара", product_name_value)
        add_field("Артикул", article_value)
        add_field("Тип товара", goods_type_label)
        add_field("Поставщик", supplier_value)
        add_field("Бренд", brand_value)
        add_field("Предмет", subject_value)
        add_field("Цвет", color_value)
        add_field("Состав", composition_value)
        add_field("Пол", payload.get("gender"))
        add_field("Сезон", payload.get("season"))
        add_field("Заказ №", payload.get("order_no"))
        add_field("Артикул WB", payload.get("wb_article"))

        rows_source = selected_card.get("rows") or []
        if not rows_source:
            rows_source = fallback_rows
        card_rows = []
        for row in rows_source:
            if not isinstance(row, dict):
                continue
            size_value = row.get("size") or row.get("size_value") or row.get("size_no") or ""
            barcode_value = row.get("barcode") or row.get("barcode_value") or ""
            qty_value = row.get("qty")
            if qty_value in (None, ""):
                qty_value = row.get("recount_qty")
            if qty_value in (None, ""):
                qty_value = row.get("processing_qty")
            if not (size_value or barcode_value or qty_value):
                continue
            card_rows.append(
                {
                    "size": size_value,
                    "barcode": barcode_value,
                    "qty": qty_value if qty_value not in (None, "") else "-",
                }
            )

        label_print_buttons = []
        marking_qty = _parse_qty_value(payload.get("marking_5840_qty"))
        marking_each_qty = _parse_qty_value(payload.get("marking_5840_each_qty"))
        if marking_qty or marking_each_qty:
            card_path_id = card_id or _processing_card_id(selected_card) or article_value or ""
            base_path = f"/orders/processing/{order_id}/card/{card_path_id}/labels/"
            base_params = {}
            if article_param:
                base_params["article"] = article_param
            if return_url:
                base_params["return"] = return_url

            def build_label_print_url(mode: str) -> str:
                params = dict(base_params)
                params["label_print"] = "1"
                params["label_mode"] = mode
                return f"{base_path}?{urlencode(params)}"

            if marking_qty:
                label_print_buttons.append(
                    {
                        "label": "Распечатать этикетки",
                        "mode": "no-cz",
                        "url": build_label_print_url("no-cz"),
                    }
                )
            if marking_each_qty:
                label_print_buttons.append(
                    {
                        "label": "Распечатать этикетки ЧЗ",
                        "mode": "cz",
                        "url": build_label_print_url("cz"),
                    }
                )

        ctx["card"] = {
            "article": article_value,
            "product_name": product_name,
            "photo_url": photo_url,
            "goods_type": goods_type_label,
        }
        ctx["card_fields"] = card_fields
        ctx["card_rows"] = card_rows
        ctx["label_print_buttons"] = label_print_buttons
        printed_cz_total = 0
        if order_id:
            barcodes = [
                str(row.get("barcode") or "").strip()
                for row in card_rows
                if isinstance(row, dict)
            ]
            barcodes = [value for value in barcodes if value]
            qs = MarkingCode.objects.filter(
                order_type="processing",
                order_id=order_id,
                printed_at__isnull=False,
            )
            if agency:
                qs = qs.filter(agency=agency)
            if barcodes:
                qs = qs.filter(barcode__in=barcodes)
            printed_cz_total = qs.count()
        ctx["printed_cz_total"] = printed_cz_total
        card_id_value = _processing_card_id(selected_card) or card_id or article_value
        processed_cards, placed_cards = _processing_card_sets(payload)
        ctx["card_processed"] = bool(card_id_value and card_id_value in processed_cards)
        ctx["card_placed"] = bool(card_id_value and card_id_value in placed_cards)
        ctx["card_processed_at"] = selected_card.get("processed_at") if isinstance(selected_card, dict) else ""
        ctx["card_processed_by"] = selected_card.get("processed_by") if isinstance(selected_card, dict) else ""
        role = get_request_role(self.request)
        ctx["can_finish_card"] = role in {"storekeeper", "processing_head"} and not ctx["card_processed"]
        processing_params = _processing_params_from_payload(payload)
        direction_tables = []
        direction_plan = _parse_json_value(payload.get("direction_plan_json"), {})
        direction_addresses = _parse_json_value(payload.get("direction_addresses_json"), [])
        if isinstance(direction_addresses, dict):
            direction_addresses = (
                direction_addresses.get("directions")
                or direction_addresses.get("addresses")
                or []
            )
        if not isinstance(direction_addresses, list):
            direction_addresses = []
        if not direction_addresses and isinstance(direction_plan, dict):
            plan_dirs = direction_plan.get("directions") or direction_plan.get("addresses") or []
            if isinstance(plan_dirs, list):
                direction_addresses = plan_dirs
        direction_addresses = [
            str(item).strip() for item in direction_addresses if str(item).strip()
        ]
        plan_rows = []
        if isinstance(direction_plan, dict):
            plan_rows = direction_plan.get("rows") or []
        if not isinstance(plan_rows, list):
            plan_rows = []
        filtered_rows = []
        if article_value:
            article_lower = article_value.lower()
            for row in plan_rows:
                if not isinstance(row, dict):
                    continue
                row_article = str(row.get("article") or "").strip().lower()
                if row_article and row_article != article_lower:
                    continue
                filtered_rows.append(row)
        elif product_name:
            product_lower = product_name.lower()
            for row in plan_rows:
                if not isinstance(row, dict):
                    continue
                row_product = str(row.get("product_name") or "").strip().lower()
                if row_product and row_product != product_lower:
                    continue
                filtered_rows.append(row)
        if not filtered_rows and plan_rows:
            filtered_rows = [row for row in plan_rows if isinstance(row, dict)]
        barcode_map = {}
        for row in card_rows:
            size_key = str(row.get("size") or "").strip().lower()
            barcode_value = str(row.get("barcode") or "").strip()
            if size_key and barcode_value and size_key not in barcode_map:
                barcode_map[size_key] = barcode_value

        def parse_dir_qty(value) -> int:
            qty = _parse_qty_value(value)
            return qty if qty is not None else 0

        direction_labels = [_short_city(item) or str(item).strip() for item in direction_addresses]

        for idx, address in enumerate(direction_addresses):
            total = 0
            rows_out = []
            for row in filtered_rows:
                if not isinstance(row, dict):
                    continue
                quantities = row.get("quantities") or []
                if not isinstance(quantities, (list, tuple)):
                    quantities = []
                qty = parse_dir_qty(quantities[idx] if idx < len(quantities) else None)
                if qty <= 0:
                    continue
                size_value = str(row.get("size") or "").strip()
                barcode_value = barcode_map.get(size_value.lower(), "") if size_value else ""
                rows_out.append(
                    {
                        "size": size_value or "-",
                        "barcode": barcode_value or "-",
                        "qty": qty,
                    }
                )
                total += qty
            if total <= 0:
                continue
            direction_tables.append(
                {"address": address, "total": total, "rows": rows_out}
            )
        results_map = {}
        saved_results = payload.get("processing_results") or []
        if isinstance(saved_results, list):
            for item in saved_results:
                if not isinstance(item, dict):
                    continue
                article_key = str(item.get("article") or "").strip().lower()
                size_key = str(item.get("size") or "").strip().lower()
                dest_key = str(item.get("destination") or "").strip().lower()
                results_map[(article_key, size_key, dest_key)] = item

        results_rows = []
        if filtered_rows and direction_labels:
            for row in filtered_rows:
                if not isinstance(row, dict):
                    continue
                size_value = str(row.get("size") or "").strip()
                row_article = str(
                    row.get("article") or row.get("product_name") or article_value or ""
                ).strip()
                barcode_value = barcode_map.get(size_value.lower(), "") if size_value else ""
                quantities = row.get("quantities") or []
                if not isinstance(quantities, (list, tuple)):
                    quantities = []
                for idx, label in enumerate(direction_labels):
                    qty = parse_dir_qty(quantities[idx] if idx < len(quantities) else None)
                    key = (row_article.lower(), size_value.lower(), (label or "").lower())
                    saved = results_map.get(key, {})
                    results_rows.append(
                        {
                            "article": row_article or "-",
                            "size": size_value or "-",
                            "barcode": barcode_value or "-",
                            "received": qty,
                            "destination": label or "-",
                            "processed": saved.get("processed") or "",
                            "defect": saved.get("defect") or "",
                            "shortage": saved.get("shortage") or "",
                            "shipped_qty": saved.get("shipped_qty") or "",
                        }
                    )
        has_direction_distribution = bool(direction_tables)
        if not has_direction_distribution:
            processing_params = [
                row for row in processing_params if row.get("label") != "Распределение по направлениям"
            ]
        ctx["processing_params"] = processing_params
        ctx["results_rows"] = results_rows
        ctx["direction_tables"] = direction_tables
        ctx["has_direction_distribution"] = has_direction_distribution
        role = get_request_role(self.request)
        ctx["can_edit_results"] = role in {
            "processing_head",
            "head_manager",
            "director",
            "admin",
        }
        return ctx

    def render_to_response(self, context, **response_kwargs):
        response = super().render_to_response(context, **response_kwargs)
        response["Cache-Control"] = "no-store, no-cache, must-revalidate, max-age=0"
        response["Pragma"] = "no-cache"
        response["Vary"] = "Cookie"
        return response

    def post(self, request, *args, **kwargs):
        action = (request.POST.get("action") or "").strip().lower()
        if action == "finish_card":
            role = get_request_role(request)
            if role not in {"storekeeper", "processing_head"}:
                return HttpResponseForbidden("Доступ запрещен")
            order_id = kwargs.get("order_id")
            if not order_id:
                return redirect("/orders/")
            entries = list(
                OrderAuditEntry.objects.filter(order_id=order_id, order_type="processing")
                .select_related("agency")
                .order_by("created_at")
            )
            if not entries:
                return redirect("/orders/")
            latest = entries[-1]
            payload = dict(latest.payload or {})
            card_id = str(request.POST.get("card_id") or "").strip()
            cards = payload.get("cards") or []
            target_card = None
            for card in cards:
                if not isinstance(card, dict):
                    continue
                if card_id and _processing_card_id(card) == card_id:
                    target_card = card
                    break
            if not target_card and len(cards) == 1 and isinstance(cards[0], dict):
                target_card = cards[0]
                if not card_id:
                    card_id = _processing_card_id(target_card)
            if target_card:
                target_card["processed_at"] = timezone.localtime().isoformat()
                if request.user and request.user.is_authenticated:
                    target_card["processed_by"] = (
                        request.user.get_full_name().strip()
                        or request.user.username
                        or str(request.user)
                    )
                target_card["processed_done"] = True
            processed_cards = payload.get("processed_cards") or []
            if not isinstance(processed_cards, list):
                processed_cards = []
            if card_id and card_id not in processed_cards:
                processed_cards.append(card_id)
            payload["processed_cards"] = processed_cards
            payload["cards"] = cards
            log_order_action(
                "update",
                order_id=order_id,
                order_type="processing",
                user=request.user if request.user.is_authenticated else None,
                agency=latest.agency if latest else None,
                description="Обработка завершена по карте товара",
                payload=payload,
            )
            return redirect(request.get_full_path())
        if action not in {"save_results", "return_to_processing"}:
            return HttpResponseForbidden("Доступ запрещен")

        role = get_request_role(request)
        order_id = kwargs.get("order_id")
        if not order_id:
            return redirect("/orders/")
        entries = list(
            OrderAuditEntry.objects.filter(order_id=order_id, order_type="processing")
            .select_related("agency")
            .order_by("created_at")
        )
        if not entries:
            return redirect("/orders/")
        latest = entries[-1]
        payload = dict(latest.payload or {})

        def collect_results():
            articles = request.POST.getlist("result_article[]")
            sizes = request.POST.getlist("result_size[]")
            destinations = request.POST.getlist("result_destination[]")
            received_list = request.POST.getlist("result_received[]")
            processed_list = request.POST.getlist("result_processed[]")
            defect_list = request.POST.getlist("result_defect[]")
            shortage_list = request.POST.getlist("result_shortage[]")
            shipped_list = request.POST.getlist("result_shipped_qty[]")
            total_rows = max(
                len(articles),
                len(sizes),
                len(destinations),
                len(received_list),
                len(processed_list),
                len(defect_list),
                len(shortage_list),
                len(shipped_list),
            )
            if total_rows <= 0:
                return [], False
            results = []
            for idx in range(total_rows):
                article = articles[idx].strip() if idx < len(articles) else ""
                size = sizes[idx].strip() if idx < len(sizes) else ""
                destination = destinations[idx].strip() if idx < len(destinations) else ""
                if not any((article, size, destination)):
                    continue
                results.append(
                    {
                        "article": article,
                        "size": size,
                        "destination": destination,
                        "received": received_list[idx].strip() if idx < len(received_list) else "",
                        "processed": processed_list[idx].strip() if idx < len(processed_list) else "",
                        "defect": defect_list[idx].strip() if idx < len(defect_list) else "",
                        "shortage": shortage_list[idx].strip() if idx < len(shortage_list) else "",
                        "shipped_qty": shipped_list[idx].strip() if idx < len(shipped_list) else "",
                    }
                )
            return results, True

        def mark_card_processed(card_id_value: str) -> None:
            if not card_id_value:
                return
            cards = payload.get("cards") or []
            target_card = None
            for card in cards:
                if not isinstance(card, dict):
                    continue
                if _processing_card_id(card) == card_id_value:
                    target_card = card
                    break
            if not target_card and len(cards) == 1 and isinstance(cards[0], dict):
                target_card = cards[0]
            if target_card:
                target_card["processed_at"] = timezone.localtime().isoformat()
                if request.user and request.user.is_authenticated:
                    target_card["processed_by"] = (
                        request.user.get_full_name().strip()
                        or request.user.username
                        or str(request.user)
                    )
                target_card["processed_done"] = True
            processed_cards = payload.get("processed_cards") or []
            if not isinstance(processed_cards, list):
                processed_cards = []
            if card_id_value and card_id_value not in processed_cards:
                processed_cards.append(card_id_value)
            payload["processed_cards"] = processed_cards
            payload["cards"] = cards

        return_to = (request.POST.get("return_to") or request.POST.get("return") or "").strip()

        if action == "save_results":
            if role not in {"processing_head", "head_manager", "director", "admin"}:
                return HttpResponseForbidden("Доступ запрещен")
            results, has_results = collect_results()
            if has_results:
                payload["processing_results"] = results
                payload["processing_results_updated_at"] = timezone.localtime().isoformat()
                log_order_action(
                    "result",
                    order_id=order_id,
                    order_type="processing",
                    user=request.user if request.user.is_authenticated else None,
                    agency=latest.agency if latest else None,
                    description="Результаты обработки",
                    payload=payload,
                )
            if return_to.startswith("/"):
                return redirect(return_to)
            return redirect(request.get_full_path())

        if action == "return_to_processing":
            if role in {"processing_head", "head_manager", "director", "admin"}:
                results, has_results = collect_results()
                if has_results:
                    payload["processing_results"] = results
                    payload["processing_results_updated_at"] = timezone.localtime().isoformat()
                    log_order_action(
                        "result",
                        order_id=order_id,
                        order_type="processing",
                        user=request.user if request.user.is_authenticated else None,
                        agency=latest.agency if latest else None,
                        description="Результаты обработки",
                        payload=payload,
                    )
            if role in {"storekeeper", "processing_head", "head_manager", "director", "admin"}:
                mark_card_processed(str(request.POST.get("card_id") or "").strip())
                log_order_action(
                    "update",
                    order_id=order_id,
                    order_type="processing",
                    user=request.user if request.user.is_authenticated else None,
                    agency=latest.agency if latest else None,
                    description="Карта обработки отмечена как выполненная",
                    payload=payload,
                )
            if return_to.startswith("/"):
                return redirect(return_to)
            return redirect(request.get_full_path())

        return redirect(request.get_full_path())


class ProcessingLabelPrintView(ProcessingCardView):
    template_name = "processing/processing_label_print.html"

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        label_base = {}
        try:
            label_base = json.loads(ctx.get("label_base_json") or "{}")
        except (TypeError, ValueError, json.JSONDecodeError):
            label_base = {}

        card_rows = ctx.get("card_rows") or []
        order_id = kwargs.get("order_id") or ctx.get("order_id") or ""
        agency = getattr(self.request, "_processing_card_agency", None)
        barcodes = [
            str(row.get("barcode") or "").strip()
            for row in card_rows
            if isinstance(row, dict)
        ]
        barcodes = [value for value in barcodes if value]
        codes_map: dict[tuple[str, str], list[str]] = {}
        if order_id and barcodes:
            base_qs = MarkingCode.objects.filter(
                order_type="processing",
                used_at__isnull=True,
                printed_at__isnull=True,
                barcode__in=barcodes,
            )
            if agency:
                base_qs = base_qs.filter(agency=agency)
            reserved_qs = base_qs.filter(order_id=order_id).order_by("created_at")
            free_qs = base_qs.filter(Q(order_id__isnull=True) | Q(order_id="")).order_by(
                "created_at"
            )
            codes_qs = list(reserved_qs.values("barcode", "size", "code")) + list(
                free_qs.values("barcode", "size", "code")
            )
            for entry in codes_qs:
                barcode = str(entry.get("barcode") or "").strip()
                if not barcode:
                    continue
                size_key = str(entry.get("size") or "").strip().lower()
                code = str(entry.get("code") or "").strip()
                if not code:
                    continue
                codes_map.setdefault((barcode, size_key), []).append(code)

        for row in card_rows:
            if not isinstance(row, dict):
                continue
            barcode = str(row.get("barcode") or "").strip()
            size_key = str(row.get("size") or "").strip().lower()
            qty_value = _parse_qty_value(row.get("qty"))
            row["qty_value"] = qty_value if qty_value is not None else 0
            codes = []
            if barcode:
                codes = codes_map.get((barcode, size_key)) or codes_map.get((barcode, "")) or []
            row["cz_codes"] = codes
            row["cz_codes_json"] = json.dumps(codes, ensure_ascii=True)

        default_row = None
        for row in card_rows:
            if isinstance(row, dict) and row.get("barcode"):
                default_row = row
                break
        if default_row is None and card_rows:
            default_row = card_rows[0] if isinstance(card_rows[0], dict) else {}
        if default_row is None:
            default_row = {}

        default_size = str(default_row.get("size") or "").strip()
        default_barcode = str(default_row.get("barcode") or "").strip()

        default_cz = ""
        if isinstance(default_row, dict):
            cz_codes = default_row.get("cz_codes") or []
            if cz_codes:
                default_cz = str(cz_codes[0] or "").strip()

        ctx["label_sample"] = {
            "article": label_base.get("article") or "",
            "name": label_base.get("name") or "",
            "size": default_size,
            "brand": label_base.get("brand") or "",
            "subject": label_base.get("subject") or "",
            "color": label_base.get("color") or "",
            "composition": label_base.get("composition") or "",
            "supplier": label_base.get("supplier") or "",
            "country": label_base.get("country") or "",
            "barcode_extra": label_base.get("barcode_extra") or "",
            "cz_code": default_cz,
        }
        ctx["label_sample_barcode"] = default_barcode
        ctx["label_rows"] = card_rows

        ctx["label_sizes"] = [
            entry for entry in LABEL_SIZES if entry.get("key") in {"item", "item_cz"}
        ]

        agent_status = load_print_agent_status()
        agent_name = str(agent_status.get("agent") or "").strip() or "неизвестно"
        last_seen_raw = agent_status.get("last_seen")
        last_seen_text = "нет данных"
        is_online = False
        if last_seen_raw:
            try:
                last_seen = datetime.fromisoformat(str(last_seen_raw))
                if timezone.is_naive(last_seen):
                    last_seen = timezone.make_aware(last_seen)
                last_seen_text = timezone.localtime(last_seen).strftime("%d.%m.%Y %H:%M:%S")
                is_online = (timezone.now() - last_seen) <= timedelta(seconds=20)
            except (TypeError, ValueError):
                last_seen_text = str(last_seen_raw)

        pending_count = ProcessingPrintJob.objects.filter(
            status=ProcessingPrintJob.STATUS_PENDING,
        ).count()
        printing_count = ProcessingPrintJob.objects.filter(
            status=ProcessingPrintJob.STATUS_PRINTING,
        ).count()
        failed_count = ProcessingPrintJob.objects.filter(
            status=ProcessingPrintJob.STATUS_FAILED,
        ).count()
        last_job = ProcessingPrintJob.objects.order_by("-updated_at").first()
        last_error = ""
        last_job_time = ""
        if last_job:
            last_job_time = timezone.localtime(last_job.updated_at).strftime("%d.%m.%Y %H:%M:%S")
            if last_job.status == ProcessingPrintJob.STATUS_FAILED:
                last_error = last_job.error or "ошибка без описания"

        paused = bool(agent_status.get("paused"))
        if paused:
            print_status = "Печать остановлена"
        elif pending_count:
            print_status = f"В очереди: {pending_count}"
            if not is_online:
                print_status = f"{print_status} (агент не активен)"
        elif last_job and last_job.status == ProcessingPrintJob.STATUS_FAILED:
            print_status = "Ошибка печати"
        elif last_job and last_job.status == ProcessingPrintJob.STATUS_PRINTING:
            print_status = "Печать выполняется"
        else:
            print_status = "Готов к печати"

        agent_line = f"{agent_name} · {last_seen_text}" if last_seen_text else agent_name
        ctx["print_status_line"] = print_status
        ctx["print_agent_line"] = agent_line
        ctx["print_last_error"] = last_error
        ctx["print_last_job_time"] = last_job_time
        ctx["print_queue_pending"] = pending_count
        ctx["print_queue_printing"] = printing_count
        ctx["print_queue_failed"] = failed_count
        ctx["print_paused"] = paused

        order_id = ctx.get("order_id") or kwargs.get("order_id") or ""
        card_id = ctx.get("card_id") or kwargs.get("card_id") or ""
        article_param = (self.request.GET.get("article") or "").strip()
        card_article = (ctx.get("card") or {}).get("article") or ""
        card_path_id = card_id or card_article or ""
        processing_card_url = f"/orders/processing/{order_id}/card/{card_path_id}/"
        if article_param:
            processing_card_url = f"{processing_card_url}?{urlencode({'article': article_param})}"
        ctx["processing_card_url"] = processing_card_url
        return ctx


class ProcessingDetailView(OrdersDetailView):
    order_type = "processing"
    allowed_roles = ("manager", "storekeeper", "head_manager", "director", "admin", "processing_head")

    def dispatch(self, request, *args, **kwargs):
        order_id = kwargs.get("order_id")
        client_agency = _client_agency_from_request(request)
        if client_agency:
            request._client_agency = client_agency
        if order_id and not client_agency:
            latest_entry = (
                OrderAuditEntry.objects.filter(order_id=order_id, order_type=self.order_type)
                .order_by("-created_at")
                .first()
            )
            if latest_entry and _is_draft_payload(latest_entry.payload or {}):
                return HttpResponseForbidden("Доступ запрещен")
        return super().dispatch(request, *args, **kwargs)

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        order_id = kwargs.get("order_id")
        entries_list = list(
            OrderAuditEntry.objects.filter(order_id=order_id, order_type=self.order_type)
            .select_related("user", "agency")
            .order_by("created_at")
        )
        payload = self._payload_from_entries(entries_list)
        stock_rows_raw = payload.get("stock_rows") or []
        size_rows_raw = payload.get("size_rows") or []
        unboxing_rows_raw = payload.get("unboxing_rows") or []

        size_rows = []
        if isinstance(stock_rows_raw, list) and stock_rows_raw:
            for row in stock_rows_raw:
                if not isinstance(row, dict):
                    continue
                size_rows.append(
                    {
                        "article": row.get("article") or "",
                        "size": row.get("size") or "",
                        "barcode": row.get("barcode") or "",
                        "qty": row.get("qty") or "",
                    }
                )
        elif isinstance(size_rows_raw, list):
            for row in size_rows_raw:
                if not isinstance(row, dict):
                    continue
                size_rows.append(
                    {
                        "article": row.get("size_no") or "",
                        "size": row.get("size_value") or "",
                        "barcode": row.get("barcode") or "",
                        "qty": row.get("recount_qty") or "",
                    }
                )

        unboxing_rows = []
        if isinstance(unboxing_rows_raw, list):
            for row in unboxing_rows_raw:
                if not isinstance(row, dict):
                    continue
                unboxing_rows.append(
                    {
                        "date": row.get("date") or "",
                        "box_size": row.get("box_size") or "",
                        "multiple": row.get("multiple") or "",
                        "box_qty": row.get("box_qty") or "",
                        "pallet_qty": row.get("pallet_qty") or "",
                        "storage_zone": row.get("storage_zone") or "",
                    }
                )

        insert_types = payload.get("insert_types") or []
        if isinstance(insert_types, str):
            insert_types = [insert_types] if insert_types else []
        insert_other = (payload.get("insert_other") or "").strip()
        if insert_other:
            insert_types = list(insert_types) + [insert_other]

        marking_stickers = payload.get("marking_stickers") or []
        if isinstance(marking_stickers, str):
            marking_stickers = [marking_stickers] if marking_stickers else []
        marking_sizes = payload.get("marking_sizes") or []
        if isinstance(marking_sizes, str):
            marking_sizes = [marking_sizes] if marking_sizes else []
        processing_params = _processing_params_from_payload(payload)

        def format_pack(prefix: str, title: str):
            needed = _format_payload_value(payload.get(f"{prefix}_needed"))
            parts = []
            type_value = (payload.get(f"{prefix}_type") or "").strip()
            size_value = (payload.get(f"{prefix}_size") or "").strip()
            qty_value = (payload.get(f"{prefix}_qty") or "").strip()
            supply_value = (payload.get(f"{prefix}_supply") or "").strip()
            if type_value:
                parts.append(f"Тип: {type_value}")
            if size_value:
                parts.append(f"Размер: {size_value}")
            if qty_value:
                parts.append(f"Кол-во: {qty_value}")
            if supply_value:
                parts.append(f"Закупка: {supply_value}")
            value = needed
            if parts:
                joiner = ", ".join(parts)
                if value == "-":
                    value = joiner
                else:
                    value = f"{value}; {joiner}"
            return {"label": title, "value": value}

        processing_fields = [
            {"label": "Наименование товара", "value": _format_payload_value(payload.get("product_name"))},
            {"label": "Маркетплейс", "value": _format_payload_value(payload.get("marketplace"))},
            {"label": "Поставщик", "value": _format_payload_value(payload.get("supplier"))},
            {"label": "Бренд", "value": _format_payload_value(payload.get("brand"))},
            {"label": "Предмет", "value": _format_payload_value(payload.get("subject"))},
            {"label": "Артикул", "value": _format_payload_value(payload.get("article"))},
            {"label": "Артикул ВБ", "value": _format_payload_value(payload.get("wb_article"))},
            {"label": "Цвет", "value": _format_payload_value(payload.get("color"))},
            {"label": "Состав", "value": _format_payload_value(payload.get("composition"))},
            {"label": "Пол", "value": _format_payload_value(payload.get("gender"))},
            {"label": "Сезон", "value": _format_payload_value(payload.get("season"))},
            {"label": "Заказ №", "value": _format_payload_value(payload.get("order_no"))},
            {"label": "Приобретение в 1С №", "value": _format_payload_value(payload.get("purchase_1c_no"))},
            {"label": "Приобретение в 1С от", "value": _format_payload_value(payload.get("purchase_1c_date"))},
            {"label": "Менеджер проекта", "value": _format_payload_value(payload.get("project_manager"))},
            {"label": "Склад приемка", "value": _format_payload_value(payload.get("warehouse_receiving"))},
            {"label": "Склад упаковка", "value": _format_payload_value(payload.get("warehouse_packing"))},
            {"label": "Склад раскоробовка", "value": _format_payload_value(payload.get("warehouse_unpacking"))},
            {"label": "Замер в упаковке", "value": _format_payload_value(payload.get("measure_needed"))},
            {"label": "Вес (грамм)", "value": _format_payload_value(payload.get("measure_weight"))},
            {"label": "Ширина (см)", "value": _format_payload_value(payload.get("measure_width"))},
            {"label": "Высота (см)", "value": _format_payload_value(payload.get("measure_height"))},
            {"label": "Глубина (см)", "value": _format_payload_value(payload.get("measure_depth"))},
            {"label": "Проверка на брак (%)", "value": _format_payload_value(payload.get("defect_percent"))},
            {"label": "Проверка на брак (кол-во)", "value": _format_payload_value(payload.get("defect_qty"))},
            {"label": "Обрезание ниток (кол-во)", "value": _format_payload_value(payload.get("trim_threads_qty"))},
            {"label": "Скрепление скотчем (кол-во)", "value": _format_payload_value(payload.get("tape_qty"))},
            {"label": "Удаление бирки", "value": _format_payload_value(payload.get("remove_tag"))},
            {"label": "Удаление бирки (кол-во)", "value": _format_payload_value(payload.get("remove_tag_qty"))},
            {"label": "Скрепление бирки", "value": _format_payload_value(payload.get("attach_tag"))},
            {"label": "Скрепление бирки (кол-во)", "value": _format_payload_value(payload.get("attach_tag_qty"))},
            {"label": "Маркировка", "value": _format_payload_list(marking_stickers)},
            {"label": "Размеры стикеров", "value": _format_payload_list(marking_sizes)},
            {"label": "Информационный", "value": _format_payload_value(payload.get("marking_info"))},
            {"label": "Сборка набора", "value": _format_payload_value(payload.get("set_build"))},
            {"label": "Кол-во ед. в наборе", "value": _format_payload_value(payload.get("set_qty"))},
            {"label": "Доп. вложение", "value": _format_payload_value(payload.get("insert_needed"))},
            {"label": "Типы вложений", "value": _format_payload_list(insert_types)},
            {"label": "Вытянуть из мешка и наклеить ЧЗ", "value": _format_payload_value(payload.get("pull_from_bag"))},
            format_pack("bubble_wrap", "Упаковка в бабл пленку"),
            format_pack("bag_replace", "Замена пакета"),
            format_pack("box_replace", "Замена гофрокороба"),
            format_pack("shrink_wrap", "Термоусадочная упаковка"),
            {"label": "Кол-во оптовых мест", "value": _format_payload_value(payload.get("wholesale_places_qty"))},
            {"label": "Счет №", "value": _format_payload_value(payload.get("invoice_no"))},
            {"label": "Дата выставления", "value": _format_payload_value(payload.get("invoice_date"))},
            {"label": "Дата оплаты", "value": _format_payload_value(payload.get("payment_date"))},
            {"label": "Бухгалтер", "value": _format_payload_value(payload.get("accountant"))},
            {"label": "В архив (дата)", "value": _format_payload_value(payload.get("archive_date"))},
            {"label": "Исполнитель", "value": _format_payload_value(payload.get("executor_name"))},
            {"label": "Дата начала", "value": _format_payload_value(payload.get("start_date"))},
            {"label": "Дата окончания", "value": _format_payload_value(payload.get("end_date"))},
            {"label": "Дата приема заказа", "value": _format_payload_value(payload.get("receive_date"))},
            {"label": "Ответственный", "value": _format_payload_value(payload.get("responsible_name"))},
            {"label": "Комментарий", "value": _format_payload_value(payload.get("comments"))},
        ]

        ctx["processing_fields"] = processing_fields
        ctx["processing_size_rows"] = size_rows
        ctx["processing_unboxing_rows"] = unboxing_rows
        ctx["processing_params"] = processing_params
        ctx["processing_meta"] = {
            "product_name": _format_payload_value(payload.get("product_name")),
            "order_no": _format_payload_value(payload.get("order_no")),
            "supplier": _format_payload_value(payload.get("supplier")),
            "brand": _format_payload_value(payload.get("brand")),
            "subject": _format_payload_value(payload.get("subject")),
        }
        ctx["items"] = []
        ctx["can_edit_order"] = False
        ctx["can_send_to_warehouse"] = False
        ctx["can_create_receiving_act"] = False
        ctx["has_receiving_act"] = False
        ctx["has_placement_act"] = False
        ctx["act_label"] = ""
        ctx["placement_act_label"] = ""
        ctx["can_send_act_to_client"] = False
        status_payload = entries_list[-1].payload if entries_list else {}
        status_value = (status_payload.get("status") or status_payload.get("submit_action") or "").lower()
        status_label = (status_payload.get("status_label") or "").lower()
        if status_value == "processing_in_work" or "взята" in status_label:
            Task.objects.filter(
                route=f"/orders/processing/{order_id}/",
                assigned_to__role="processing_head",
            ).update(status="in_progress")
        is_done = (
            status_value in {"done", "completed", "closed", "finished", "processing_head", "processing_in_work"}
            or "выполн" in status_label
            or "утверж" in status_label
            or "передан" in status_label
            or "взята" in status_label
        )
        is_waiting = status_value in {"sent_unconfirmed", "send", "submitted"} or "подтверждени" in status_label
        role = get_request_role(self.request)
        can_manage = role in {"manager", "head_manager", "director", "admin"}
        client_view = bool(ctx.get("client_view"))
        is_ready_for_work = (
            status_value == "processing_head"
            or ("передан" in status_label and "обработ" in status_label)
        )
        ctx["can_approve_processing"] = bool(can_manage and not client_view and is_waiting and not is_done)
        ctx["can_edit_processing"] = bool(can_manage and not client_view and not is_done)
        ctx["can_take_processing"] = bool(
            not client_view
            and role in {"storekeeper", "processing_head"}
            and is_ready_for_work
        )
        packers = []
        if order_id:
            packing_tasks = (
                Task.objects.filter(
                    route=f"/orders/processing/{order_id}/flow/",
                    assigned_to__role="processing_worker",
                )
                .exclude(status="done")
                .select_related("assigned_to")
                .order_by("-created_at")
            )
            seen = set()
            for task in packing_tasks:
                assignee = task.assigned_to
                if not assignee:
                    continue
                label = assignee.full_name or str(assignee)
                if label in seen:
                    continue
                seen.add(label)
                packers.append(label)
        ctx["processing_packers"] = packers
        ctx["processing_packers_label"] = ", ".join(packers)
        ctx["processing_work_url"] = f"/orders/processing/{order_id}/work/"
        if ctx["can_edit_processing"]:
            agency = ctx.get("agency")
            if agency and getattr(agency, "id", None):
                ctx["processing_edit_url"] = (
                    f"/orders/processing/?order={order_id}&agency={agency.id}&edit=1"
                )
            else:
                ctx["processing_edit_url"] = f"/orders/processing/?order={order_id}&edit=1"
        return ctx

    def post(self, request, *args, **kwargs):
        order_id = kwargs.get("order_id")
        if not order_id:
            return redirect("/orders/")
        action = (request.POST.get("action") or "").strip().lower()
        if action == "take_processing":
            role = get_request_role(request)
            if role not in {"storekeeper", "processing_head"}:
                return HttpResponseForbidden("Доступ запрещен")
            entries = list(
                OrderAuditEntry.objects.filter(order_id=order_id, order_type=self.order_type)
                .select_related("agency")
                .order_by("created_at")
            )
            if not entries:
                return redirect("/orders/")
            latest = entries[-1]
            status_payload = latest.payload or {}
            status_value = (status_payload.get("status") or status_payload.get("submit_action") or "").lower()
            status_label = (status_payload.get("status_label") or "").lower()
            if status_value == "processing_in_work" or "взята" in status_label:
                return redirect(f"/orders/processing/{order_id}/work/")
            payload = dict(self._payload_from_entries(entries))
            payload["status"] = "processing_in_work"
            payload["status_label"] = "Взята в работу"
            payload["work_started_at"] = timezone.localtime().isoformat()
            log_order_action(
                "status",
                order_id=order_id,
                order_type=self.order_type,
                user=request.user if request.user.is_authenticated else None,
                agency=latest.agency if latest else None,
                description="Заявка на обработку принята в работу",
                payload=payload,
            )
            Task.objects.filter(
                route=f"/orders/processing/{order_id}/",
                assigned_to__role="processing_head",
            ).update(status="in_progress")
            return redirect(f"/orders/processing/{order_id}/work/")
        if action == "approve_processing":
            role = get_request_role(request)
            if role not in {"manager", "head_manager", "director", "admin"}:
                return HttpResponseForbidden("Доступ запрещен")
            entries = list(
                OrderAuditEntry.objects.filter(order_id=order_id, order_type=self.order_type)
                .select_related("agency")
                .order_by("created_at")
            )
            if not entries:
                return redirect("/orders/")
            latest = entries[-1]
            status_payload = latest.payload or {}
            status_value = (status_payload.get("status") or status_payload.get("submit_action") or "").lower()
            status_label = (status_payload.get("status_label") or "").lower()
            if status_value in {"done", "completed", "closed", "finished"} or "выполн" in status_label:
                return redirect(f"/orders/processing/{order_id}/")
            payload = dict(self._payload_from_entries(entries))
            payload["status"] = "processing_head"
            payload["status_label"] = "Передано в обработку"
            payload["approved_at"] = timezone.localtime().isoformat()
            log_order_action(
                "status",
                order_id=order_id,
                order_type=self.order_type,
                user=request.user if request.user.is_authenticated else None,
                agency=latest.agency if latest else None,
                description="Заявка на обработку утверждена менеджером и передана в обработку",
                payload=payload,
            )
            Task.objects.filter(
                route=f"/orders/processing/{order_id}/",
                assigned_to__role="manager",
            ).exclude(status="done").update(status="done")
            _create_processing_head_task(
                order_id,
                latest.agency if latest else None,
                request,
                timezone.localtime(),
            )
            return redirect(f"/orders/processing/{order_id}/")
        comment = (request.POST.get("comment") or "").strip()
        if comment:
            latest = (
                OrderAuditEntry.objects.filter(order_id=order_id, order_type=self.order_type)
                .select_related("agency")
                .order_by("-created_at")
                .first()
            )
            log_order_action(
                "comment",
                order_id=order_id,
                order_type=self.order_type,
                user=request.user if request.user.is_authenticated else None,
                agency=latest.agency if latest else None,
                description=comment,
                payload={"comment": comment},
            )
        return redirect(f"/orders/processing/{order_id}/")


@login_required
@require_POST
def processing_marking_availability(request):
    data = _parse_json_body(request)
    if data is None:
        return HttpResponseBadRequest("Invalid JSON")
    items = data.get("items") or []
    if not isinstance(items, list):
        return HttpResponseBadRequest("Invalid items")
    order_id = str(data.get("order_id") or "").strip()
    client_agency = _client_agency_from_request(request)
    agency = client_agency
    if not agency:
        role = get_request_role(request)
        if role not in {"manager", "storekeeper", "head_manager", "director", "admin"}:
            return HttpResponseForbidden("Доступ запрещен")
        agency_id = data.get("agency_id") or request.GET.get("client") or request.GET.get("agency")
        agency = Agency.objects.filter(pk=agency_id).first() if agency_id else None
    if not agency:
        return HttpResponseBadRequest("Клиент не выбран")
    required: dict[str, int] = {}
    required_total = 0
    missing_barcodes = 0
    for item in items:
        if not isinstance(item, dict):
            continue
        qty = _parse_qty_value(item.get("qty"))
        if not qty or qty <= 0:
            continue
        barcode = str(item.get("barcode") or "").strip()
        if not barcode:
            missing_barcodes += qty
            continue
        required_total += qty
        required[barcode] = required.get(barcode, 0) + qty
    available_map = _marking_available_by_barcode(agency, order_id)
    free_map = _marking_free_by_barcode(agency)
    details = []
    covered_total = 0
    missing_total = 0
    free_total = 0
    for barcode, need in required.items():
        have = available_map.get(barcode, 0)
        free_count = free_map.get(barcode, 0)
        covered = min(need, have)
        missing = max(need - have, 0)
        covered_total += covered
        missing_total += missing
        free_total += free_count
        details.append({"barcode": barcode, "required": need, "available": have, "missing": missing})
    return JsonResponse(
        {
            "ok": True,
            "required": required_total,
            "available": covered_total,
            "missing": missing_total,
            "free": free_total,
            "missing_barcodes": missing_barcodes,
            "items": details,
        }
    )


@login_required
@require_POST
def processing_marking_import(request):
    client_agency = _client_agency_from_request(request)
    agency = client_agency
    if not agency:
        role = get_request_role(request)
        if role not in {"manager", "storekeeper", "head_manager", "director", "admin"}:
            return HttpResponseForbidden("Доступ запрещен")
        agency_id = request.POST.get("agency_id") or request.GET.get("client") or request.GET.get("agency")
        agency = Agency.objects.filter(pk=agency_id).first() if agency_id else None
    if not agency:
        return HttpResponseBadRequest("Клиент не выбран")
    cz_file = request.FILES.get("marking_cz_file") or request.FILES.get("file")
    if not cz_file:
        return JsonResponse({"ok": False, "error": "Файл не выбран."}, status=400)
    cards_payload = _parse_json_value(request.POST.get("cards_json"), [])
    payload = {}
    if cards_payload:
        payload["cards"] = cards_payload
    order_id = str(request.POST.get("order_id") or "").strip()
    ok, import_result = _import_marking_codes(cz_file, payload, order_id, agency, request.user)
    if not ok:
        message = import_result.get("error") or "Ошибка импорта ЧЗ."
        return JsonResponse({"ok": False, "error": message}, status=400)
    return JsonResponse({"ok": True, **import_result})



@login_required
@require_POST
def enqueue_processing_print_job(request):
    data = _parse_json_body(request)
    if data is None:
        return HttpResponseBadRequest("Invalid JSON")
    barcode = str(data.get("barcode") or "").strip()
    if not barcode:
        return JsonResponse({"ok": False, "error": "Barcode is required"}, status=400)
    label_png_base64 = str(data.get("label_png_base64") or "").strip()
    if not label_png_base64:
        return JsonResponse({"ok": False, "error": "Label image is required"}, status=400)
    try:
        label_width_mm = int(data.get("label_width_mm") or 58)
    except (TypeError, ValueError):
        label_width_mm = 58
    try:
        label_height_mm = int(data.get("label_height_mm") or 40)
    except (TypeError, ValueError):
        label_height_mm = 40
    job = ProcessingPrintJob.objects.create(
        order_id=str(data.get("order_id") or "").strip(),
        card_id=str(data.get("card_id") or "").strip(),
        article=str(data.get("article") or "").strip(),
        barcode=barcode,
        size=str(data.get("size") or "").strip(),
        printer_name=str(data.get("printer_name") or "").strip(),
        label_png_base64=label_png_base64,
        label_width_mm=label_width_mm,
        label_height_mm=label_height_mm,
        requested_by=request.user.username if request.user.is_authenticated else "",
    )
    return JsonResponse({"ok": True, "job_id": job.id})


@require_GET
def processing_print_jobs_next(request):
    ok, response = _check_print_agent_token(request)
    if not ok:
        return response
    agent_name = (request.GET.get("agent") or request.headers.get("X-Print-Agent") or "").strip()
    save_print_agent_status(agent_name)
    if load_print_agent_status().get("paused"):
        return JsonResponse({"ok": True, "has_job": False, "paused": True})
    with transaction.atomic():
        job = (
            ProcessingPrintJob.objects.select_for_update()
            .filter(status=ProcessingPrintJob.STATUS_PENDING)
            .order_by("created_at")
            .first()
        )
        if not job:
            return JsonResponse({"ok": True, "has_job": False})
        job.status = ProcessingPrintJob.STATUS_PRINTING
        if agent_name:
            job.agent = agent_name
        job.save(update_fields=["status", "agent", "updated_at"])
    return JsonResponse({"ok": True, "has_job": True, "job": _serialize_print_job(job)})


@csrf_exempt
@require_POST
def processing_print_jobs_complete(request):
    ok, response = _check_print_agent_token(request)
    if not ok:
        return response
    data = _parse_json_body(request)
    if data is None:
        data = request.POST
    job_id = data.get("job_id") or data.get("id")
    if not job_id:
        return JsonResponse({"ok": False, "error": "job_id is required"}, status=400)
    job = ProcessingPrintJob.objects.filter(pk=job_id).first()
    if not job:
        return JsonResponse({"ok": False, "error": "Job not found"}, status=404)
    status_value = str(data.get("status") or "").strip().lower()
    error_text = str(data.get("error") or "").strip()
    if status_value not in {
        ProcessingPrintJob.STATUS_PRINTED,
        ProcessingPrintJob.STATUS_FAILED,
        ProcessingPrintJob.STATUS_PENDING,
    }:
        status_value = ProcessingPrintJob.STATUS_PRINTED
    job.status = status_value
    job.error = error_text
    job.save(update_fields=["status", "error", "updated_at"])
    return JsonResponse({"ok": True})


@login_required
@require_POST
def processing_print_jobs_pause(request):
    ok, response = _require_print_admin(request)
    if not ok:
        return response
    data = _parse_json_body(request) or {}
    printer = str(data.get("printer") or request.POST.get("printer") or "").strip()
    username = request.user.get_full_name().strip() if request.user.is_authenticated else ""
    if not username:
        username = request.user.username if request.user.is_authenticated else ""
    set_print_agent_pause(True, by=username)
    if printer:
        _enqueue_agent_command("printer.pause", {"printer": printer})
    return JsonResponse({"ok": True, "paused": True, "counts": _print_queue_counts()})


@login_required
@require_POST
def processing_print_jobs_resume(request):
    ok, response = _require_print_admin(request)
    if not ok:
        return response
    data = _parse_json_body(request) or {}
    printer = str(data.get("printer") or request.POST.get("printer") or "").strip()
    username = request.user.get_full_name().strip() if request.user.is_authenticated else ""
    if not username:
        username = request.user.username if request.user.is_authenticated else ""
    set_print_agent_pause(False, by=username)
    if printer:
        _enqueue_agent_command("printer.resume", {"printer": printer})
    return JsonResponse({"ok": True, "paused": False, "counts": _print_queue_counts()})


@login_required
@require_POST
def processing_print_jobs_clear(request):
    ok, response = _require_print_admin(request)
    if not ok:
        return response
    data = _parse_json_body(request) or {}
    printer = str(data.get("printer") or request.POST.get("printer") or "").strip()
    scope = str(data.get("scope") or request.POST.get("scope") or "pending").strip().lower()
    if scope not in {"pending", "failed", "all"}:
        return JsonResponse({"ok": False, "error": "invalid_scope"}, status=400)
    qs = ProcessingPrintJob.objects.all()
    if scope == "pending":
        qs = qs.filter(status=ProcessingPrintJob.STATUS_PENDING)
    elif scope == "failed":
        qs = qs.filter(status=ProcessingPrintJob.STATUS_FAILED)
    else:
        qs = qs.filter(status__in=[ProcessingPrintJob.STATUS_PENDING, ProcessingPrintJob.STATUS_FAILED])
    deleted_count = qs.count()
    qs.delete()
    if printer:
        _enqueue_agent_command("printer.clear", {"printer": printer})
    return JsonResponse({"ok": True, "deleted": deleted_count, "counts": _print_queue_counts()})


@login_required
@require_POST
def processing_print_jobs_reset(request):
    ok, response = _require_print_admin(request)
    if not ok:
        return response
    data = _parse_json_body(request) or {}
    printer = str(data.get("printer") or request.POST.get("printer") or "").strip()
    mode = str(data.get("mode") or request.POST.get("mode") or "pending").strip().lower()
    if mode not in {"pending", "failed"}:
        return JsonResponse({"ok": False, "error": "invalid_mode"}, status=400)
    qs = ProcessingPrintJob.objects.filter(status=ProcessingPrintJob.STATUS_PRINTING)
    if mode == "failed":
        updated = qs.update(
            status=ProcessingPrintJob.STATUS_FAILED,
            error="Сброшено вручную",
        )
    else:
        updated = qs.update(
            status=ProcessingPrintJob.STATUS_PENDING,
            error="",
            agent="",
        )
    if printer:
        _enqueue_agent_command("printer.clear", {"printer": printer})
    return JsonResponse({"ok": True, "updated": updated, "counts": _print_queue_counts()})


@login_required
@require_GET
def download_processing_print_agent_cmd(request):
    token = _get_print_agent_token()
    if not token:
        return HttpResponseBadRequest("PRINT_AGENT_TOKEN not set")
    server_url = request.build_absolute_uri("/").rstrip("/")
    script_url = f"{server_url}/orders/processing/print-agent/script/?token={token}"
    content = (
        "@echo off\r\n"
        f"set SERVER_URL={server_url}\r\n"
        f"set TOKEN={token}\r\n"
        "set AGENT_DIR=%~dp0\r\n"
        "powershell -ExecutionPolicy Bypass -Command "
        f"\"Invoke-WebRequest -Uri '{script_url}' -OutFile '%AGENT_DIR%\\print_agent.ps1'\""
        "\r\n"
        "powershell -ExecutionPolicy Bypass -File \"%AGENT_DIR%print_agent.ps1\" "
        "-ServerUrl \"%SERVER_URL%\" -Token \"%TOKEN%\"\r\n"
        "pause\r\n"
    )
    response = HttpResponse(content, content_type="application/octet-stream")
    response["Content-Disposition"] = "attachment; filename=print_agent_start.cmd"
    return response


@login_required
@require_GET
def download_processing_print_agent_install_cmd(request):
    token = _get_print_agent_token()
    if not token:
        return HttpResponseBadRequest("PRINT_AGENT_TOKEN not set")
    server_url = request.build_absolute_uri("/").rstrip("/")
    script_url = f"{server_url}/orders/processing/print-agent/script/?token={token}"
    content = (
        "@echo off\r\n"
        "setlocal\r\n"
        f"set SERVER_URL={server_url}\r\n"
        f"set TOKEN={token}\r\n"
        "set AGENT_DIR=%APPDATA%\\FullboxPrintAgent\r\n"
        "if not exist \"%AGENT_DIR%\" mkdir \"%AGENT_DIR%\"\r\n"
        "powershell -ExecutionPolicy Bypass -Command "
        f"\"Invoke-WebRequest -Uri '{script_url}' -OutFile '%AGENT_DIR%\\print_agent.ps1'\""
        "\r\n"
        "set STARTUP_DIR=%APPDATA%\\Microsoft\\Windows\\Start Menu\\Programs\\Startup\r\n"
        "set VBS_FILE=%STARTUP_DIR%\\FullboxPrintAgent.vbs\r\n"
        "> \"%VBS_FILE%\" echo Set WshShell = CreateObject(\"WScript.Shell\")\r\n"
        ">> \"%VBS_FILE%\" echo WshShell.Run \"powershell -WindowStyle Hidden -ExecutionPolicy Bypass -File \"\"%AGENT_DIR%\\print_agent.ps1\"\" -ServerUrl \"\"%SERVER_URL%\"\" -Token \"\"%TOKEN%\"\"\", 0\r\n"
        ">> \"%VBS_FILE%\" echo Set WshShell = Nothing\r\n"
        "wscript \"%VBS_FILE%\"\r\n"
        "echo Autostart installed.\r\n"
        "pause\r\n"
    )
    response = HttpResponse(content, content_type="application/octet-stream")
    response["Content-Disposition"] = "attachment; filename=print_agent_install.cmd"
    return response


@login_required
@require_GET
def download_processing_printer_sync_package(request):
    path = settings.BASE_DIR.parent / "sync_printers.ps1"
    if not path.exists():
        return HttpResponseBadRequest("sync_printers.ps1 not found")
    ps1_content = path.read_text(encoding="utf-8")
    encoded = base64.b64encode(ps1_content.encode("utf-16le")).decode("ascii")
    cmd_content = (
        "@echo off\r\n"
        f"powershell -NoProfile -ExecutionPolicy Bypass -EncodedCommand {encoded}\r\n"
        "pause\r\n"
    )
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("printer_sync.ps1", ps1_content)
        archive.writestr("run_printer_sync.cmd", cmd_content)
        archive.writestr(
            "README.txt",
            "Запускайте run_printer_sync.cmd. Файл printer_sync.ps1 вручную не запускать.\r\n",
        )
    buffer.seek(0)
    response = HttpResponse(buffer.getvalue(), content_type="application/zip")
    response["Content-Disposition"] = "attachment; filename=printer_sync_package.zip"
    return response

@require_GET
def download_processing_print_agent_script(request):
    ok, response = _check_print_agent_token(request)
    if not ok:
        return response
    path = settings.BASE_DIR.parent / "print_agent.ps1"
    if not path.exists():
        return HttpResponseBadRequest("print_agent.ps1 not found")
    content = path.read_text(encoding="utf-8")
    response = HttpResponse(content, content_type="text/plain")
    response["Content-Disposition"] = "attachment; filename=print_agent.ps1"
    return response

